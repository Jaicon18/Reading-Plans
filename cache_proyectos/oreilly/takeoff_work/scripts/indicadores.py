import openpyxl, json, sys, re
def leer(p):
    wb = openpyxl.load_workbook(p, data_only=True)
    ms = {str(r[0]): r[2] for r in wb["Material Summary"].iter_rows(min_row=4, values_only=True) if r and r[0] and isinstance(r[2], (int, float))}
    lf = sf = area = 0.0
    for r in wb["Partition Takeoff"].iter_rows(min_row=5, values_only=True):
        if r and r[1] and str(r[1]).startswith("TOTAL - PARTITIONS"):
            lf, sf = float(r[2] or 0), float(r[7] or 0)
        elif r and r[0] and r[1] and isinstance(r[2], (int, float)) and isinstance(r[3], (int, float)) and not str(r[0]).startswith("TYPE"):
            area += float(r[2]) * float(r[3])
    return ms, lf, sf, area
def s(ms, rx, unit_rx=None):
    return sum(v for k, v in ms.items() if re.search(rx, k, re.I))
def ind(p):
    ms, lf, sf, area = leer(p)
    studs = s(ms, r"metal stud \(LF of material\)|stud \(all gauges\)|Z-furring channel \(LF")
    track = s(ms, r"^(Bottom|Top).*track")
    board = s(ms, r"gypsum board")
    scr = s(ms, r"drywall screws|S-12 drywall screws")
    return {"Partition LF (grand total)": lf, "Wall area LF x H (SF)": round(area, 1), "Partition board SF": sf,
            "Board SF per LF of wall": sf / lf, "Stud + furring LF per SF of wall area": studs / area,
            "Bottom + top track LF per LF of wall": track / lf, "Total gypsum board SF (Material Summary)": board,
            "Drywall screws per SF of board": scr / board, "Joint compound lb per SF of board": s(ms, r"joint compound") / board,
            "Joint tape LF per SF of board": s(ms, r"joint tape") / board, "Bracing LF per LF of wall": s(ms, r"^Bracing") / lf,
            "Average wall height (ft)": area / lf}
bn = ind(sys.argv[1]); orl = ind(sys.argv[2])
rows = []
for k in bn:
    a, b = bn[k], orl[k]
    d = (b - a) / a * 100 if a else 0
    rows.append((k, a, b, d))
print(json.dumps(rows))
