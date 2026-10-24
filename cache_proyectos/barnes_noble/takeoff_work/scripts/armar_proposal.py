"""Proposal B&N desde "Proposal Template.xlsx": solo la hoja PROPOSAL DW, todo en inglés, alcance y cantidades del
take-off, exclusiones, clarifications con los RFIs y PRECIOS EN BLANCO. Quita todo texto de otros clientes.

Uso: python3 scripts/armar_proposal.py <Proposal Template.xlsx> <takeoff.json> <familias_out.json> <out.xlsx>
"""
import json, re, sys, datetime
import openpyxl
from openpyxl.styles import Alignment, Font

tpl, tk, fam, out = sys.argv[1:5]
T = json.load(open(tk))
F = json.load(open(fam))
fa = F["familias"]
wb = openpyxl.load_workbook(tpl)
for ws in list(wb.worksheets):          # las demás hojas traen costos y textos de otros proyectos
    if ws.title != "PROPOSAL DW":
        wb.remove(ws)
ws = wb["PROPOSAL DW"]
ws.title = "PROPOSAL"
ws.sheet_state = "visible"

# celdas de otros clientes / precios / cálculos sueltos: se vacían todas fuera del texto que se reescribe abajo
for row in ws.iter_rows():
    for c in row:
        if c.coordinate in ("C3",):       # membrete del proponente (se conserva)
            continue
        if c.column > 11 or (c.value is not None and c.row >= 5):
            try:
                c.value = None
            except AttributeError:        # celda fusionada (no la primera)
                pass

WRAP = Alignment(wrap_text=True, vertical="top")
def put(ref, v, wrap=True, bold=False):
    ws[ref].value = v
    if wrap:
        ws[ref].alignment = WRAP
    if bold:
        ws[ref].font = Font(bold=True, name=ws[ref].font.name, size=ws[ref].font.size)

def q(k):
    return fa.get(k, 0) or 0

TBC = "[TO BE COMPLETED]"
put("C5", "PROPOSAL SUBMITTED TO:"); put("G5", "PER PLAN & SPECS:")
put("C6", TBC); put("G6", "YES - bid set listed below")
put("C7", "PROJECT NAME:"); put("G7", "ARCHITECT:")
put("C8", "BARNES & NOBLE – COLORADO SPRINGS (Store 3566, Project 33288)")
put("G8", "WJCA - William J. Commer, Architect")
put("C9", "PROJECT ADDRESS:"); put("G9", "DATE OF PLANS:")
put("C10", "3173 CINEMA POINT, COLORADO SPRINGS, CO 80922")
put("G10", "07-23-2025 (BD comments revision, title block A-101)")
put("C11", "DATE OF PROPOSAL:"); put("G11", "ADDENDUM:")
put("C12", datetime.datetime(2026, 10, 24)); ws["C12"].number_format = "mm-dd-yyyy"
put("G12", "None received")
put("C13", "PHONE NUMBER:"); put("G13", "E-MAIL:")
put("C14", TBC); put("G14", TBC)
put("C15", "Scope of Work:", bold=True)
put("C16", "We propose to supply the material, labor, and equipment necessary to complete the work in accordance "
           "with the bid set and the scope listed below.")
put("C17", "(Please take special note of the scope description; we are not required to perform any items or tasks "
           "not specifically listed)")

put("C19", "1.-"); put("D19", "Division of work: METAL FRAMING, DRYWALL & DRYWALL FINISH (CSI Div. 05 40 00 / 09 21 16 / 09 22 16).", bold=True)
tl = [f"• Partition types {', '.join(T['type_order'])} per partition schedule A-410 and Cafe types on A-221 (plans A-101, A-221).",
      f"• Steel studs: 3-5/8\" ≈ {q('3-5/8\" stud (todos los gauges)'):,.0f} LF, 6\" ≈ {q('6\" stud (todos los gauges)'):,.0f} LF, "
      f"8\" ≈ {q('8\" stud (todos los gauges)'):,.0f} LF (net, before waste).",
      f"• Bottom track {q('bottom track (todos)'):,.0f} LF; top / deep-leg deflection track {q('top / deflection track (todos)'):,.0f} LF; "
      f"ceiling / soffit track {q('track de ceilings / soffits'):,.0f} LF.",
      f"• Lateral bracing C.R. channel {q('bracing C.R. channel'):,.0f} LF (5/A-410); type 2 > 22'-0\" with 14 ga studs + diagonal bracing per note 3/A-410 (see RFI-04).",
      f"• 5/8\" gypsum board walls and ceilings ≈ {q('5/8\" gyp board regular (walls + ceilings)'):,.0f} SF; 1/2\" cement board {q('cement board'):,.0f} SF (tile walls).",
      f"• FRT plywood substrate ≈ {q('FRT plywood (SF)'):,.0f} SF (types I, 3, 4 and electrical boards).",
      f"• Suspended GWB ceilings (5/A-121): furring channel {q('furring channel'):,.0f} LF, carrying channel {q('carrying channel'):,.0f} LF, "
      f"hanger wire {F.get('hanger_pcs', {}).get('total_pcs', 0)} pcs; framed soffits 2/A-221, 3/A-221.",
      "• Column surrounds C25 (9 ea, 24\"x24\" GWB over 3-5/8\" framing to 8'-0\" A.F.F.).",
      f"• Metal trim (corner bead / L-bead) ≈ {q('metal trim (corner bead, casing, shadow molding)'):,.0f} LF; tape, joint compound and Level 4 finish.",
      "• Quantities are from the attached take-off (Barnes_Noble_ColoradoSprings_Takeoff_FINAL.xlsx); waste not included in the figures above."]
put("C21", "\n".join(tl))
for r in list(range(19, 44)) + [46, 69, 70, 75, 76]:
    ws.row_dimensions[r].hidden = False
    ws.row_dimensions[r].height = 15
ws["C21"].font = Font(name="Calibri", size=10)
put("C44", "1a. Metal framing\n1.1. Interior partitions, exterior wall lining types 1 / 1A / 1B / 1C (framing only), Cafe types, "
           "column surrounds, soffits and ceiling framing per A-101, A-121, A-221, A-410.\n1.2. Deflection top track at deck per 2/A-410.")
put("C45", "1b. Drywall & drywall finish Level 4\n1.1. 5/8\" gypsum board per partition type and face; board to deck where there is "
           "no ceiling (2/A-410).\n1.2. Industry standard touch-ups & punch-list for turnover (see section 6).")
put("C46", "1c. Gypsum board ceilings and soffits\n1.1. Suspended GWB ceilings per A-121 (5/A-121, 4/A-121) and Cafe soffit 2/A-221.")
put("C47", "2. Accessories\n2.1. FRT plywood substrate, 1/2\" cement board at tile walls, metal trim, fasteners, tape and joint compound.")
put("C48", "Exclusions:", bold=True)
ex = ["• Gyp bd on exterior types 1/1A/1B/1C (by future tenant)", "• Acoustical sealant", "• Insulation",
      "• FRP", "• Moisture resistant board (RFI-11)", "• FRT blocking (RFI-23)", "• Door headers (RFI-09)",
      "• Control joints (RFI-24)", "• Demising wall patch (RFI-25)", "• Vestibule ceiling box (RFI-16)",
      "• Painting", "• Acoustical ceilings", "• Demolition", "• Door frames install", "• Millwork / casework",
      "• Permits & bonds", "• Temporary utilities", "• Dumpster", "• Prevailing wages", "• Shop drawings"]
for m in [m for m in ws.merged_cells.ranges if 49 <= m.min_row <= 55]:
    ws.unmerge_cells(str(m))
ws.merge_cells("C49:J55")
lin = ["      ".join(ex[i:i + 3]) for i in range(0, len(ex), 3)]
put("C49", "\n".join(lin))
ws["C49"].font = Font(name="Calibri", size=10)
for r in range(49, 56):
    ws.row_dimensions[r].height = 16
put("C56", "• Anything not specifically listed in the scope above.")
put("C57", "NOTE: PROPOSAL HOLDS GOOD ONLY FOR THE ITEMS AND QUANTITIES IN THE ATTACHED TAKE-OFF")
put("C59", "3.-"); put("D59", "Work Standard:", bold=True)
put("C60", "All work is to be completed in a workmanlike manner according to standard practices. Workers will remain on "
           "the job until completion of the project. The worksite will be cleaned daily and upon project completion. "
           "All agreements are contingent upon delays beyond our control.")
put("C61", "4.-"); put("D61", "Limited Warranty:", bold=True)
put("C62", "The completed work shall carry a 1-year warranty on labor and materials. This warranty is in lieu of all "
           "other warranties, expressed or implied, and excludes consequential or incidental damages caused by accident "
           "or abuse, temperature changes, settlement, moisture, or cracks caused by expansion and/or contraction.")
put("C63", "5.-"); put("D63", "Insurance:", bold=True)
put("C64", "All employees are covered under blanket workman's comp, with liability and umbrella insurance policies.")
put("C65", "6.-"); put("D65", "Touch-Ups:", bold=True)
put("C66", "Drywall & finishes: included on all work performed, limited to 1 hr per every 1,000 SF for quality issues; "
           "trade damage repairs tracked on work ticket. Finish levels per USG / GA-214.")
put("C67", "7.-"); put("D67", "Others:", bold=True)
put("C68", "7.1. Work sequenced according to industry standards; out of sequence work and schedule delays require added cost.\n"
           "7.2. Work days Monday – Saturday, 7:00 am - 5:00 pm (no overtime or night work included).\n"
           "7.3. T&M rates: [TO BE COMPLETED].\n7.4. Price valid for ninety (90) days, unless otherwise noted.")
pr = [("C69", "8.- Metal framing:"), ("C71", "9.- Drywall & drywall finish Level 4:"),
      ("C73", "9.a-"), ("D73", "Gypsum board ceilings & soffits:"),
      ("C75", "9.b-"), ("D75", "Accessories (FRT plywood, cement board, trim):")]
for ref, t in pr:
    put(ref, t, wrap=False, bold=True)
put("D83", "Total all trades:"); put("D85", "Taxes:"); put("D87", "Total including tax:")
for r in (70, 72, 74, 76, 84, 86, 88):
    ws[f"I{r}"].value = None               # PRECIOS EN BLANCO (los llena el estimador)
put("C88", "[PRICE IN WORDS - TO BE COMPLETED]")
ws.merge_cells("C89:J89")
put("C89", "**NOTE: PRICES LEFT BLANK FOR THE ESTIMATOR. SEE CLARIFICATIONS (RFIs) BELOW.")
put("C91", "Acceptance of proposal:", bold=True)
put("C92", "The prices, specifications and conditions are satisfactory and are hereby accepted. You are authorized to do "
           "the work as specified.")
put("G93", "Signature:"); put("G95", "Date of acceptance:")

for rr in range(77, 83):
    ws.row_dimensions[rr].hidden = True
r = 98
ws.merge_cells(f"C{r}:J{r}")
put(f"C{r}", "CLARIFICATIONS (open RFIs - pricing is based on the assumption written in each one):", bold=True)
for x in T["rfis"]:
    r += 1
    ws.merge_cells(f"C{r}:J{r}")
    txt = re.sub(r"\s*\((?:rule|R)\s?[\d.]+[a-z]?\)", "", f"{x['id']}: {x['pregunta']}  Delta: {x['delta']}")
    put(f"C{r}", txt)
    ws.row_dimensions[r].height = max(30, 15 * (1 + len(x['pregunta'] + x['delta']) // 110))
ws.print_area = f"A1:K{r}"
wb.save(out)
print("OK", out, "RFIs", len(T["rfis"]))
