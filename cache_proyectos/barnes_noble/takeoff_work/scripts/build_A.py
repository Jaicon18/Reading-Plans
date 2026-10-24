"""Lectura A: normaliza el caché (PARTITIONS_STEP1_rev1 + walls_coords) al esquema del canon y completa
alturas, caras, aberturas. Salida: fragmentos en 2_lectura/A/."""
import json, os, math
BN = '/home/user/Reading-Plans/cache_proyectos/barnes_noble'
OUT = BN + '/takeoff_work/etapas/2_lectura/A'
wc = {w['wall_id']: w for w in json.load(open(BN + '/walls_coords.json'))['walls']}
inv = json.load(open(BN + '/takeoff_work/etapas/1_inventario/inventario.json'))
# tag_ids (inventario PTnnnn) -> wall de la caché (asignación por cercanía, verificada con imagen)
TAGS = {'EXT-01': ['PT0001'], 'EXT-03': ['PT0002', 'PT0003'], 'EXT-02': ['PT0004', 'PT0005'], 'EXT-07': ['PT0006'],
        'EXT-18': ['PT0008'], 'EXT-17': ['PT0009'], 'EXT-04': ['PT0010'], 'EXT-12': ['PT0011'], 'EXT-05': ['PT0012'],
        'EXT-11': ['PT0013'], 'EXT-09': ['PT0014'], 'EXT-10': ['PT0015'], 'EXT-06': ['PT0016'],
        'EXT-16': ['PT0017'], 'EXT-15': ['PT0018'], 'EXT-14': ['PT0019'], 'EXT-13': ['PT0020'],
        'W2-01': ['PT0021'], 'W2-03': ['PT0022', 'PT0023'], 'W2-04': ['PT0024', 'PT0025', 'PT0026'],
        'W2-09': ['PT0027'], 'W2-05': ['PT0028'], 'W2-06': ['PT0029', 'PT0030'], 'W2-07': ['PT0031'],
        'W5-01': ['PT0032'], 'W3-01': ['PT0033'], 'W3-04': ['PT0034'], 'W3-02': ['PT0035'], 'W3-03': ['PT0036'],
        'W4-01': ['PT0037'],
        'C-H1-N': ['PT0038'], 'C-H1-W': ['PT0039'], 'C-H2a': ['PT0040'], 'C-H2b': ['PT0041'], 'C-H3a': ['PT0042'],
        'C-H3b': ['PT0043'], 'C-E2': ['PT0044'], 'C-E1-P': ['PT0045'], 'C-E1-S': ['PT0046'], 'C-F': ['PT0047'],
        'C-I': ['PT0051'], 'C-A-1': ['PT0052'], 'C-A-2': ['PT0053'], 'C-A-3': ['PT0054'], 'C-B': ['PT0055'],
        'C-D': ['PT0056'], 'C-C': ['PT0057']}
TYPE = {k: w['type'] for k, w in wc.items()}
DET = {'1A': '1A/A-410', '1': '1/A-410', '1B': '1B/A-410', '1C': '1C/A-410', '2': '2/A-410', '5': '5/A-410',
       '3': '3/A-410', '4': '4/A-410', 'H1': 'G/A-221', 'H2a': 'G/A-221', 'H2b': 'G/A-221', 'H3a': 'H/A-221',
       'H3b': 'H/A-221', 'E1': 'E/A-221', 'E2': 'E/A-221', 'F': 'F/A-221', 'I': 'I/A-221', 'A': 'A/A-221',
       'B': 'B/A-221', 'C': 'C/A-221', 'D': 'D/A-221'}
STUD = {'1A': '3-5/8"', '1': '3-5/8"', '1C': '3-5/8"', '1B': '8"', '2': '3-5/8"', '3': '3-5/8"', '4': '3-5/8"',
        '5': '6"', 'H1': '6"', 'H2a': '3-5/8"', 'H2b': '3-5/8"', 'H3a': '6"', 'H3b': '6"', 'E1': '6"', 'E2': '6"',
        'F': '3-5/8"', 'I': '6"', 'A': '3-5/8"', 'B': '3-5/8"', 'C': '3-5/8"', 'D': '3-5/8"'}
D25 = (25.0, 'B.O. DECK ABOVE EL. 25\'-0" V.I.F. - 1/A-400A, 1/A-400B')
D257 = (25.583, 'B.O. DECK ABOVE EL. 25\'-7" V.I.F. - 2/A-400A (east facade)')
D24 = (24.0, 'B.O. DECK ABOVE EL. 24\'-0" V.I.F. - 2/A-400B (west facade)')
HEIGHT = {'EXT-01': D25, 'EXT-02': D25, 'EXT-03': D24, 'EXT-04': D24, 'EXT-05': D24, 'EXT-06': D25, 'EXT-07': D25,
          'EXT-09': D25, 'EXT-10': D25, 'EXT-11': D257, 'EXT-12': D257, 'EXT-17': D257, 'EXT-18': D257,
          'EXT-13': D257, 'EXT-14': D257, 'EXT-15': D257, 'EXT-16': D257}
for k in wc:
    t = TYPE[k]
    if t in ('2', '5', 'E1', 'E2', 'I'):
        HEIGHT[k] = D25
    elif t == '3':
        HEIGHT[k] = (8.0, '8\'-0" A.F.F. T.O. WALL - 3/A-410')
    elif t == '4':
        HEIGHT[k] = (7.5, '7\'-6" A.F.F. T.O. WALL - 4/A-410')
    elif t in ('A', 'B', 'C'):
        HEIGHT[k] = (2.708, '2\'-8 1/2" B.O. SLAB TO T.O. STUD - %s' % DET[t])
    elif t == 'D':
        HEIGHT[k] = (1.0, '1\'-0" B.O. SLAB TO T.O. STUD - D/A-221')
    elif t in ('H1', 'H2a', 'H2b', 'H3a', 'H3b', 'F'):
        HEIGHT[k] = (25.0, 'NOT PRINTED (%s break line at top; A-225 finish plan has no T.O. stud): loaded to deck '
                     '25\'-0" (1/A-400B) per rule 53.3 / 60.11 - RFI-06' % DET[t])
HSTAT = {k: ('PENDING' if TYPE[k] in ('H1', 'H2a', 'H2b', 'H3a', 'H3b', 'F') else 'CONFIRMED') for k in wc}
json.dump({'TAGS': TAGS, 'HEIGHT': HEIGHT, 'HSTAT': HSTAT, 'STUD': STUD, 'DET': DET}, open('/tmp/claude-0/-home-user-Reading-Plans/b9c208e2-64c4-5dc6-9d1d-2546bd4c85dc/scratchpad/lecturaA/_tablas.json', 'w'), indent=1)
print('ok', len(wc))

# ---------------------------------------------------------------- walls (segmentos con band)
import openpyxl
xl = openpyxl.load_workbook(BN + '/PARTITIONS_STEP1_rev1.xlsx')['Walls']
cache = {}
for r in xl.iter_rows(min_row=4, values_only=True):
    if r[0] and r[1] and not str(r[0]).startswith(('SUBTOTAL', 'GRAND')):
        cache[r[0]] = dict(zip(['id', 'type', 'sheet', 'tags', 'location', 'detail', 'framing', 'lf', 'height',
                                'hsrc', 'status', 'notes'], r))
caras = json.load(open('/tmp/claude-0/-home-user-Reading-Plans/b9c208e2-64c4-5dc6-9d1d-2546bd4c85dc/scratchpad/lecturaA/caras.json'))
HALF = {'A-101': {'2': 1.92, '5': 2.64, '3': 1.92, '4': 1.92, '1': 2.28, '1A': 2.28, '1B': 4.74, '1C': 3.0},
        'A-221': {'H1': 10.86, 'H2a': 7.32, 'H2b': 7.32, 'H3a': 10.86, 'H3b': 10.86, 'E1': 10.14, 'E2': 10.14,
                  'F': 6.36, 'I': 10.08, 'A': 9.0, 'B': 11.0, 'C': 9.0, 'D': 9.0}}
SCALE = {'A-101': 9.0, 'A-221': 36.0}

def band_de(wid, seg, t, sheet):
    (x0, y0), (x1, y1) = seg[0], seg[-1]
    h = HALF[sheet][t]
    if abs(y0 - y1) < 0.5 and len(seg) == 2:
        return {'ori': 'H', 'c0': round(y0 - h, 2), 'c1': round(y0 + h, 2), 'lo': round(min(x0, x1), 2),
                'hi': round(max(x0, x1), 2), 'sheet': sheet}
    if abs(x0 - x1) < 0.5 and len(seg) == 2:
        return {'ori': 'V', 'c0': round(x0 - h, 2), 'c1': round(x0 + h, 2), 'lo': round(min(y0, y1), 2),
                'hi': round(max(y0, y1), 2), 'sheet': sheet}
    return {'ori': 'D', 'p0': [round(x0, 2), round(y0, 2)], 'p1': [round(x1, 2), round(y1, 2)], 'sheet': sheet}

walls = {'A-101': [], 'A-221': []}
for wid, w in wc.items():
    t, sheet = w['type'], w['sheet']
    c = cache[wid]
    segs = w['points_pt']
    if w['lf'] == 0:
        continue                       # EXT-08: contado una vez como E1/E2 (RFI-05) -> tag descartado con owner
    page = 6 if sheet == 'A-101' else 15
    base = {'tag': t, 'element': 'wall', 'sheet': sheet, 'page': page, 'location': c['location'], 'source': DET[t],
            'detail': DET[t], 'detail_source': DET[t], 'stud_size': STUD[t], 'stud_source': DET[t], 'oc_in': 16,
            'oc_source': DET[t], 'height_ft': HEIGHT[wid][0], 'height_source': HEIGHT[wid][1],
            'height_status': HSTAT[wid], 'length_source': c['detail'] if False else sheet,
            'status': c['status'].strip('[]'), 'img': 'img_cache/' + ('B&N_A-101_walls_by_type.png' if page == 6 else 'B&N_A-221_cafe_walls_by_type.png')}
    if len(segs) == 1 and len(segs[0]) > 2 and segs[0][0] == segs[0][-1]:      # box cerrado -> 4 tramos rectos
        segs = [[segs[0][i], segs[0][i + 1]] for i in range(len(segs[0]) - 1)]
    if len(segs) == 1 and len(segs[0]) > 2:                                   # arco -> 6 cuerdas D
        p = segs[0]; n = 6; k = [round(i * (len(p) - 1) / n) for i in range(n + 1)]
        for i in range(n):
            q0, q1 = p[k[i]], p[k[i + 1]]
            L = math.dist(q0, q1) / SCALE[sheet]
            walls[sheet].append(dict(base, id=f"{wid}.{i + 1}", tag_ids=TAGS.get(wid, []) if i == 0 else [],
                                     length_ft=round(L, 2),
                                     length_op=f"arc R32'-9\" (A-101) by Bezier vector, chord {i + 1} of {n} = {round(L, 2)} ft (cache arc {w['lf']} ft)",
                                     band={'ori': 'D', 'p0': [round(q0[0], 2), round(q0[1], 2)],
                                           'p1': [round(q1[0], 2), round(q1[1], 2)], 'sheet': sheet}))
        continue
    tot = 0.0
    for i, s in enumerate(segs):
        L = math.dist(s[0], s[-1]) / SCALE[sheet]
        tot += L
    for i, s in enumerate(segs):
        L = math.dist(s[0], s[-1]) / SCALE[sheet]
        sid = wid if len(segs) == 1 else f"{wid}.{i + 1}"
        tid = TAGS.get(wid, []) if (len(segs) == 1 or i == 0) else []
        lf = w['lf'] if len(segs) == 1 else round(L, 2)
        walls[sheet].append(dict(base, id=sid, tag_ids=tid, length_ft=lf,
                                 length_op=f"vector centerline {round(L, 2)} ft" + ('' if len(segs) == 1 else f" (segment {i + 1} of {wid})"),
                                 band=band_de(sid, s, t, sheet)))
for sh, L in walls.items():
    print(sh, len(L), round(sum(w['length_ft'] for w in L), 2))
json.dump(walls, open(OUT + '/_walls_tmp.json', 'w'), indent=1)

# ---------------------------------------------------------------- caras (board por face)
RCP = 'A-121'
CEIL = {  # room -> (board height ft, source, MR)
    'Sales Area 101': (None, 'OPEN TO DECK (A-121) - board to deck, 2/A-410 "EXTEND GYPSUM BOARD TO DECK IN LOCATIONS WITH NO CEILING"', False),
    'Storage 112': (None, 'OPEN TO DECK (A-121) - board to deck, 2/A-410', False),
    'Stock Room 111': (None, 'OPEN TO DECK (A-121) - board to deck, 2/A-410', False),
    'Janitor Closet 107': (None, 'OPEN TO DECK (A-121) - board to deck, 2/A-410; greenboard per A-410 partition note 7', True),
    'Corridor 110': (9.5, 'ACT-3 9\'-0" (A-121) + 6" above suspended ceiling (A-410 partition note 3, 2/A-410) = 9.5', False),
    'Break Room 106': (9.5, 'ACT-3 9\'-0" (A-121) + 6" (A-410 partition note 3, 2/A-410) = 9.5', False),
    "Manager's Office 105": (9.5, 'ACT-3 9\'-0" (A-121) + 6" (A-410 partition note 3, 2/A-410) = 9.5', False),
    "Women's Restroom 109": (9.0, 'GWB 9\'-0" (A-121); W.R. gyp bd on wet side to B.O. gyp ceiling (2/A-410, 5/A-410)', True),
    "Men's Restroom 108": (9.0, 'GWB 9\'-0" (A-121); W.R. gyp bd on wet side to B.O. gyp ceiling (2/A-410, 5/A-410)', True),
    'Cafe Seating 102': (None, 'full height of partial wall (3/A-410, 4/A-410)', False),
    'Cafe Workroom 104': (10.5, 'ACT-1 10\'-0" (A-121) + 6" (A-410 partition note 3) = 10.5', False),
    'Cafe 103': (None, 'OPEN TO DECK (A-121, Cafe) - board to deck', False),
    'Grab & Go': (7.583, 'GWB 7\'-1" (A-121) + 6" (A-410 partition note 3) = 7.583', False),
}
def face(side, room, wall_h, length=None, src=None, board=None, layers=1):
    h, hs, mr = CEIL[room]
    hh = wall_h if h is None else min(h, wall_h)
    # MR: el paquete carga MR solo detrás de toilet / lavatory (D1 / 59.9); en este set los fixtures van contra el
    # exterior EXT-07 (gyp by future tenant). Las caras de wet room van regular + RFI-11 (A-410 nota 7, 2/A-410).
    f = {'side': side, 'room': room, 'board': board or '5/8 regular', 'layers': layers,
         'height_ft': round(hh, 3), 'height_source': hs if h is not None else hs + f' = wall height {wall_h}'}
    if mr and not board:
        f['wet_room_note'] = 'A-410 partition note 7 / 2/A-410: W.R. board in wet rooms - not loaded (fixtures back onto exterior EXT-07) - RFI-11'
    if length is not None:
        f['length_ft'] = round(length, 2); f['length_source'] = src or 'A-101 (vector, split at room boundary)'
    return f
S = 'Sales Area 101'
FACES = {
 'W2-01': [('A', 'Storage 112'), ('B', S)], 'W2-02': [('A', S), ('B', 'Storage 112')],
 'W2-03': [('A', 'Stock Room 111'), ('B', S, 81.12), ('B', 'Corridor 110', 13.12)],
 'W2-04': [('A', S), ('B', 'Corridor 110', 29.04), ('B', 'Break Room 106', 19.03), ('B', "Manager's Office 105", 8.69)],
 'W2-05': [('A', 'Corridor 110'), ('B', 'Corridor 110', 5.41), ('B', "Women's Restroom 109", 7.51)],
 'W2-06.1': [('A', 'Corridor 110'), ('B', "Women's Restroom 109")],
 'W2-06.2': [('A', "Women's Restroom 109"), ('B', 'Corridor 110')],
 'W2-06.3': [('A', 'Corridor 110'), ('B', "Women's Restroom 109", 'r:18.05/36.4'), ('B', "Men's Restroom 108", 'r:18.35/36.4')],
 'W2-06.4': [('A', 'Corridor 110'), ('B', "Men's Restroom 108")],
 'W2-06.5': [('A', 'Corridor 110'), ('B', "Men's Restroom 108", 'r:61.95/101.6'), ('B', 'Janitor Closet 107', 'r:39.65/101.6')],
 'W2-07': [('A', "Men's Restroom 108"), ('B', 'Janitor Closet 107')],
 'W2-08': [('A', 'Corridor 110', 5.41), ('A', 'Janitor Closet 107', 7.51), ('B', 'Break Room 106')],
 'W2-09': [('A', 'Break Room 106'), ('B', "Manager's Office 105")],
 'W5-01': [('A', "Women's Restroom 109"), ('B', "Men's Restroom 108")],
}
for k in ('W3-01', 'W3-02', 'W3-03', 'W3-04', 'W3-05', 'W3-06', 'W4-02', 'W4-03'):
    FACES[k] = [('A', 'Cafe Seating 102' if k.startswith('W3') else S), ('B', S)]
CEIL['Sales Area 101 (partial wall)'] = CEIL['Cafe Seating 102']
for k in ('W4-01', 'W4-04'):
    for i in range(1, 5):
        FACES[f'{k}.{i}'] = [('A', S)]
for i in range(1, 7):
    FACES[f'W4-02.{i}'] = [('A', S), ('B', S)]; FACES[f'W4-03.{i}'] = [('A', S), ('B', S)]
NONE_EXT = ('5/8" GYP BD BY FUTURE TENANT (%s): gyp bd EXCLUDED from our scope (estimator decision) - RFI-01; '
            'framing only')
for sh in walls:
    for w in walls[sh]:
        wid, t = w['id'], w['tag']
        if t in ('1', '1A', '1C', '1B'):
            w['faces'] = []
            w['faces_none_reason'] = NONE_EXT % DET[t]
            continue
        spec = FACES.get(wid)
        if spec is None:
            continue
        fl = []
        L = w['length_ft']
        porside = {}
        for s_ in spec:
            porside.setdefault(s_[0], []).append(s_)
        for side, lst in porside.items():
            if len(lst) == 1:
                fl.append(face(side, lst[0][1], w['height_ft']))
                continue
            fijos = 0.0; vals = []
            for s_ in lst:
                v = s_[2] if len(s_) > 2 else None
                if isinstance(v, str) and v.startswith('r:'):
                    a, b = v[2:].split('/'); v = round(L * float(a) / float(b), 2)
                vals.append(v); fijos += v or 0.0
            vals = [v if v is not None else round(L - fijos, 2) for v in vals]
            vals[-1] = round(L - sum(vals[:-1]), 2)
            for s_, v in zip(lst, vals):
                fl.append(face(side, s_[1], w['height_ft'], v))
        if t in ('3', '4'):
            for f in fl:
                f['height_ft'] = w['height_ft']; f['height_source'] = DET[t] + ' T.O. WALL (board full height)'
        w['faces'] = fl
json.dump(walls, open(OUT + '/_walls_tmp.json', 'w'), indent=1)
print('faces ok')

# ---------------------------------------------------------------- Cafe (A-221): caras y FRT plywood
CEIL['Cafe 103'] = (None, 'OPEN TO DECK (A-121 Cafe tag) - board to wall height', False)
def bw(w): return w['band']
strips = []
def strip(w, side, room, h, hsrc, mat='frt_5_8', label='5/8" FRT PLYWOOD', length=None):
    strips.append({'ref': f"{w['id']}-{side}-FRT", 'wall': w['id'], 'lf': round(length or w['length_ft'], 2),
                   'width_ft': round(h, 3), 'material': mat, 'uso': 'substrate', 'label': label,
                   'source': w['detail'], 'room': room,
                   'op': f"{round(length or w['length_ft'], 2)} LF x {round(h, 3)} ft ({hsrc})"})
CAFE_RFI = 'height not printed (A-221 sections break at top; A-225 finish plan; A-226A elevations give finish tops only) - loaded to wall height per rule 53.3 - RFI-06'
for w in walls['A-221']:
    wid, t, H = w['id'], w['tag'], w['height_ft']
    if t in ('H1',):
        if wid == 'C-H1-N':
            b = bw(w); gg = round((b['hi'] - (1112.1 - 10.86)) / 36.0, 2)
            w['faces'] = [face('A', 'Sales Area 101', H, round(w['length_ft'] - gg, 2)),
                          face('A', 'Grab & Go', H, gg, 'A-221 (vector, x of H3b face)'),
                          face('B', 'Cafe Workroom 104', H)]
            strip(w, 'A', 'Sales Area 101', H, CAFE_RFI, length=w['length_ft'] - gg)
            strip(w, 'B', 'Cafe Workroom 104', 10.5, 'to ACT-1 10\'-0" + 6" (assumed equal to GWB; G/A-221 H2a note) - RFI-06')
        else:
            w['faces'] = [face('A', "Manager's Office 105", H), face('B', 'Cafe Workroom 104', H)]
            strip(w, 'B', 'Cafe Workroom 104', 10.5, 'to ACT-1 10\'-0" + 6" - RFI-06')
    elif t == 'H2a':
        w['faces'] = [face('A', 'Cafe Workroom 104', H), face('B', 'Cafe 103', H)]
        strip(w, 'A', 'Cafe Workroom 104', 10.5, 'G/A-221 "H2a = 5/8" GWB, 5/8" FRT PLYWOOD" to ACT-1 + 6" - RFI-06')
        strip(w, 'B', 'Cafe 103', H, CAFE_RFI)
    elif t == 'H2b':
        w['faces'] = []; w['faces_none_reason'] = 'G/A-221 "H2b = 5/8" FRT BOTH SIDES" - plywood only (plywood_strips)'
        strip(w, 'A', 'Cafe Workroom 104', 10.5, 'G/A-221 FRT both sides, to ACT-1 + 6" - RFI-06')
        strip(w, 'B', 'Cafe Workroom 104', 10.5, 'G/A-221 FRT both sides, to ACT-1 + 6" - RFI-06')
    elif t in ('H3a', 'H3b'):
        w['faces'] = [face('B', 'Grab & Go', H)]
        if t == 'H3a':
            strip(w, 'A', 'Sales Area 101', H, 'H/A-221 "H3a = PLYWOOD TO TOP OF WD-1": WD-1 top not printed - loaded to wall height (53.3) - RFI-06')
        else:
            strip(w, 'A', 'Sales Area 101', 8.0, 'H/A-221 "H3b = PLYWOOD UP TO 8\'-0"" [CONFIRMED]')
    elif t == 'E2':
        w['faces'] = []; w['faces_none_reason'] = 'E/A-221 "E2-5/8" FRT PLYWOOD" only (plywood_strips); other side = void at existing exterior wall'
        strip(w, 'A', 'Cafe Workroom 104', H, CAFE_RFI)
    elif t == 'E1':
        w['faces'] = [face('A', 'Cafe 103', H, board='1/2 cement board')]
        w['faces'][0]['height_source'] = 'E/A-221 "E1-1/2" CEMENT BOARD"; ' + CAFE_RFI
    elif t == 'F':
        w['faces'] = []; w['faces_none_reason'] = 'F/A-221 "5/8" FRT PLYWOOD" one side only (plywood_strips)'
        strip(w, 'A', 'Cafe 103', H, 'F/A-221 to bottom of nested stud bracing (elevation not printed) - RFI-06')
    elif t == 'I':
        w['faces'] = [face('A', 'Cafe 103', H)]
        w['faces'][0]['height_source'] = 'I/A-221 "5/8" GWB FROM TOP OF WD-1 TO DECK": WD-1 top not printed - GWB loaded full height (53.3) - RFI-08'
    elif t in ('A', 'B', 'C', 'D'):
        w['faces'] = [face('A', 'Cafe 103', H, board='1/2 cement board'), face('B', 'Cafe 103', H, board='1/2 cement board')]
        for f in w['faces']:
            f['height_source'] = f"{DET[t]} 1/2\" cementitious board both sides, B.O. slab to T.O. stud {H} ft"
        if t in ('A', 'B'):
            strip(w, 'A', 'Cafe 103 (customer side)', H, f'{DET[t]} "3/4" FRT PLYWOOD (SUBSTRATE)" to T.O. stud', mat='frt_3_4', label='3/4" FRT PLYWOOD')
json.dump(walls, open(OUT + '/_walls_tmp.json', 'w'), indent=1)
json.dump(strips, open(OUT + '/_strips_tmp.json', 'w'), indent=1)
print('cafe ok', len(strips))


# ---------------------------------------------------------------- column surrounds C25 (keynote A-101) y soffit sides
KN = {k['id']: k for k in json.load(open(OUT + '/img_cache/keynotes_p6.json'))}
for i, kid in enumerate([26, 27, 28, 29, 30, 64, 65, 66, 67]):
    k = KN[kid]
    walls['A-101'].append({'id': f'CW-{i + 1:02d}', 'tag': 'C25', 'element': 'column_wrap', 'sheet': 'A-101', 'page': 6,
        'location': 'Sales Area 101 - existing column, new 24"x24" GWB surround (keynote C25)',
        'source': 'A-101', 'detail': 'C25/A-101', 'detail_source': 'A-101 construction plan key note C25',
        'stud_size': '3-5/8"', 'stud_source': 'A-101 key note C25 "OVER 3 5/8" METAL FRAMING"', 'oc_in': 16,
        'oc_source': 'default 16" O.C. (not printed) - RFI-17',
        'height_ft': 8.0, 'height_source': 'A-101 key note C25 "FROM SLAB TO 8\'-0" AFF"', 'height_status': 'CONFIRMED',
        'length_ft': 8.0, 'length_source': 'A-101 key note C25 24"x24" = 4 faces x 2\'-0"', 'length_op': '4 x 2.0 = 8.0',
        'faces_ft': [2.0, 2.0, 2.0, 2.0], 'fos_chain': {'value': [2.0], 'source': 'A-101 key note C25 "24"x24" GWB SURROUND"'},
        'faces': [{'side': 'A', 'room': 'Sales Area 101', 'board': '5/8 regular', 'layers': 1, 'height_ft': 8.0,
                   'height_source': 'A-101 key note C25 to 8\'-0" AFF', 'length_ft': 8.0, 'length_source': 'A-101 key note C25 (4 x 2\'-0")'}],
        'tag_ids': [], 'status': 'CONFIRMED', 'img': 'img_cache/keynotes_contactos.png',
        'band': {'ori': 'BOX', 'c0': round(k['cx'] - 9, 2), 'c1': round(k['cx'] + 9, 2), 'lo': round(k['cy'] - 9, 2),
                 'hi': round(k['cy'] + 9, 2), 'sheet': 'A-101'}})
SS = [('SS-N', 'outer', 'H', 331.6, 1248.24, 2067.96, 25.0 - 11.0, 'Cafe 11\'-0" zone N open edge to Sales'),
      ('SS-E', 'outer', 'V', 2067.96, 331.6, 695.76, 25.0 - 11.0, 'Cafe 11\'-0" zone E open edge to Sales'),
      ('SS-IN-N', 'inner', 'H', 470.65, 1231.8, 1951.8, 25.0 - 11.5, 'Cafe 11\'-6" band N edge (6" step)'),
      ('SS-IN-E', 'inner', 'V', 1951.8, 470.65, 695.76, 25.0 - 11.5, 'Cafe 11\'-6" band E edge (6" step)')]
for sid, line, ori, c, lo, hi, h, loc in SS:
    L = round((hi - lo) / 36.0, 2)
    fh = h if line == 'outer' else 0.5
    walls['A-221'].append({'id': sid, 'tag': '', 'element': 'soffit_side', 'soffit_of': 'Cafe 103 - 11\'-0" GWB zone',
        'drop_line': line, 'sheet': 'A-221', 'page': 15, 'location': loc, 'source': '2/A-221', 'detail': '2/A-221',
        'detail_source': '2/A-221 "3 5/8" @ 16" O.C. METAL STUD UP TO DECK"', 'stud_size': '3-5/8"', 'stud_source': '2/A-221',
        'oc_in': 16, 'oc_source': '2/A-221', 'height_ft': h,
        'height_source': ('B.O. deck 25\'-0" (1/A-400B) - soffit ' + ('11\'-0"' if line == 'outer' else '11\'-6"') + ' (A-121); studs up to deck per 2/A-221'),
        'height_status': 'DERIVED', 'length_ft': L, 'length_source': 'A-221 (vector, dashed soffit edge)',
        'length_op': f'({hi} - {lo}) / 36 = {L}', 'tag_ids': [], 'status': 'DERIVED', 'img': 'img_ceilings/A121_cafe_GWB_zones.png',
        'faces': [{'side': 'A', 'room': 'Cafe 103', 'board': '5/8 regular', 'layers': 1, 'height_ft': fh,
                   'height_source': ('2/A-221 5/8" GWB on the open face to deck (T.O. WD-1 split not printed, full height per 53.3) - RFI-14' if line == 'outer' else '2/A-221 6" GWB return 11\'-6" - 11\'-0" = 0.5')}],
        'band': ({'ori': 'H', 'c0': c - 2.6, 'c1': c + 2.6, 'lo': lo, 'hi': hi, 'sheet': 'A-221'} if ori == 'H'
                 else {'ori': 'V', 'c0': c - 2.6, 'c1': c + 2.6, 'lo': lo, 'hi': hi, 'sheet': 'A-221'})})

# ---------------------------------------------------------------- nodos (L / T / free) desde las bands
def ends(w):
    b = w['band']
    if b['ori'] == 'H':
        e = (b['c0'] + b['c1']) / 2; return [(b['lo'], e), (b['hi'], e)]
    if b['ori'] == 'V':
        e = (b['c0'] + b['c1']) / 2; return [(e, b['lo']), (e, b['hi'])]
    if b['ori'] == 'D':
        return [tuple(b['p0']), tuple(b['p1'])]
    return []
def on_body(p, v, tol):
    b = v['band']
    if b['ori'] == 'H':
        e = (b['c0'] + b['c1']) / 2
        return abs(p[1] - e) <= tol and b['lo'] + tol < p[0] < b['hi'] - tol
    if b['ori'] == 'V':
        e = (b['c0'] + b['c1']) / 2
        return abs(p[0] - e) <= tol and b['lo'] + tol < p[1] < b['hi'] - tol
    if b['ori'] == 'D':
        (x0, y0), (x1, y1) = b['p0'], b['p1']; dx, dy = x1 - x0, y1 - y0; L2 = dx * dx + dy * dy
        t = ((p[0] - x0) * dx + (p[1] - y0) * dy) / L2
        if not (0.02 < t < 0.98): return False
        return math.hypot(p[0] - x0 - t * dx, p[1] - y0 - t * dy) <= tol
    return False
nodes = []
for sh in walls:
    tol = 1.0 * SCALE[sh]
    W = [w for w in walls[sh] if w['band']['ori'] != 'BOX' and w.get('element') != 'column_wrap']
    usados = set()
    for w in W:
        for i, p in enumerate(ends(w)):
            if (w['id'], i) in usados: continue
            hecho = False
            for v in W:                                   # L: extremo con extremo de otro wall
                if v is w: continue
                for j, q in enumerate(ends(v)):
                    if (v['id'], j) in usados: continue
                    if math.dist(p, q) <= tol:
                        same = (w['band']['ori'] == v['band']['ori'] and w['band']['ori'] in 'HV')
                        kind = 'split' if same and w['id'].split('.')[0] == v['id'].split('.')[0] else 'L'
                        n = {'id': f"N{len(nodes) + 1}", 'kind': kind, 'walls': [w['id'], v['id']],
                             'source': f"{sh} (vector: ends within {tol / SCALE[sh]:.1f} ft)"}
                        if kind == 'split': n['why'] = 'one wall drawn in collinear segments (cache polyline)'
                        nodes.append(n); usados |= {(w['id'], i), (v['id'], j)}; hecho = True; break
                if hecho: break
            if hecho: continue
            for v in W + [x for x in walls[sh] if x['band']['ori'] == 'BOX']:
                if v is w: continue
                if v['band']['ori'] == 'BOX':
                    b = v['band']
                    if b['c0'] - tol <= p[0] <= b['c1'] + tol and b['lo'] - tol <= p[1] <= b['hi'] + tol:
                        nodes.append({'id': f"N{len(nodes) + 1}", 'kind': 'T', 'branch': w['id'], 'through': v['id'],
                                      'source': f"{sh} (vector)"}); usados.add((w['id'], i)); hecho = True; break
                    continue
                if on_body(p, v, tol):
                    nodes.append({'id': f"N{len(nodes) + 1}", 'kind': 'T', 'branch': w['id'], 'through': v['id'],
                                  'source': f"{sh} (vector: end on the body of {v['id']})"})
                    usados.add((w['id'], i)); hecho = True; break
            if not hecho:
                nodes.append({'id': f"N{len(nodes) + 1}", 'kind': 'free', 'wall': w['id'],
                              'source': f"{sh} (vector: free end)"}); usados.add((w['id'], i))
    for v in walls[sh]:
        if v['band']['ori'] == 'BOX' and v.get('element') != 'column_wrap':
            nodes.append({'id': f"N{len(nodes) + 1}", 'kind': 'free', 'wall': v['id'], 'source': 'closed box (one end)'})
            nodes.append({'id': f"N{len(nodes) + 1}", 'kind': 'free', 'wall': v['id'], 'source': 'closed box (other end)'})
from collections import Counter
print('nodes', Counter(n['kind'] for n in nodes))
json.dump(nodes, open(OUT + '/_nodes_tmp.json', 'w'), indent=1)

# ---------------------------------------------------------------- cambios de type sin nodo dibujado (R75.5 / R73.10)
TCR = {
 'EXT-01': {'alternativas': ['1A/1 change at W2-03 centerline x=498.5 pt (cache reading; printed 15\'-7" to face of W2-03 supports it)',
                             '1A/1 change at the edge of the adjacent tenant building x=800.4 pt (A-101 north facade, EXTERIOR vs ADJACENT TENANT)'],
            'delta_lf': round((800.4 - 498.5) / 9.0, 2), 'source': 'A-101 north facade (no node, no dimension at the 1A/1 change) - RFI-02'},
 'C-E2': {'alternativas': ['E2/E1 change at H2a centerline x=1223.65 pt (taken)', 'at the Workroom face of H2a x=1216.33 pt',
                           'at the Cafe face of H2a x=1230.97 pt'],
          'delta_lf': round((1230.97 - 1216.33) / 36.0 / 2, 2), 'source': '1/A-221 (no node drawn at the E2/E1 change) - RFI-05'},
}
for sh in walls:
    for w in walls[sh]:
        if w['id'] in TCR:
            w['type_change_rfi'] = TCR[w['id']]

# ---------------------------------------------------------------- metal trim (corner bead)
trim = []
libres = Counter(n['wall'] for n in nodes if n['kind'] == 'free')
for t in ('3', '4'):
    ws = [w for w in walls['A-101'] if w['tag'] == t]
    top2 = sum(w['length_ft'] for w in ws if len(w.get('faces') or []) == 2)
    top1 = sum(w['length_ft'] for w in ws if len(w.get('faces') or []) == 1)
    h = ws[0]['height_ft']
    nfree = sum(libres[w['id']] for w in ws)
    lf = 2 * top2 + top1 + nfree * 2 * h
    trim.append({'ref': f'Corner bead - type {t} partial height walls (top edges + free ends)', 'kind': 'corner_bead',
                 'lf': round(lf, 2), 'source': DET[t],
                 'op': f"top edges 2 x {round(top2, 2)} + 1 x {round(top1, 2)} (cap box outer face) + {nfree} free end(s) x 2 corners x {h} ft = {round(lf, 2)}"})
cb = [w for w in walls['A-101'] if w['tag'] == '4' and len(w.get('faces') or []) == 1]
trim.append({'ref': 'Corner bead - type 4 cap boxes outside vertical corners', 'kind': 'corner_bead', 'lf_each': 7.5,
             'qty': 8, 'lf': 60.0, 'source': '4/A-410',
             'op': '2 cap boxes x 4 outside corners x 7\'-6" (T.O. wall 4/A-410) = 60.0'})
trim.append({'ref': 'Corner bead - corridor recess (drinking fountain alcove, W2-06) outside corners', 'kind': 'corner_bead',
             'lf_each': 9.5, 'qty': 2, 'lf': 19.0, 'source': 'A-101',
             'op': '2 outside corners x 9.5 ft (ACT-3 9\'-0" + 6", corridor face) = 19.0'})


# ---------------------------------------------------------------- ceilings GWB (A-121; medidos en ceilings_measured.json)
CM = {c['room']: c for c in json.load(open(OUT + '/datos/ceilings_measured.json'))}
def ceil(key, room, system, drop, detail=None, **kw):
    c = CM[key]
    d = {'room': room, 'system': system, 'area_sf': c['area_sf'], 'perimeter_ft': c['perimeter_ft'], 'ns_ft': c['ns_ft'],
         'ew_ft': c['ew_ft'], 'height_ft': c['height_ft'], 'source': 'A-121', 'area_op': c.get('area_op'),
         'img': 'img_ceilings/' + os.path.basename(c['img'].split(';')[0].strip())}
    if system in ('furring', 'stud_joist'):
        d['hanger_drop_ft'] = {'value': round(25.0 - c['height_ft'], 3),
                               'source': f"B.O. deck 25'-0\" (1/A-400B) - ceiling {c['height_ft']} (A-121) - RFI-03"}
    if detail:
        d['detail'] = detail
    d.update(kw)
    return d
ceilings = [
 ceil("Women's Restroom", "Women's Restroom 109", 'furring', None, '5/A-121'),
 ceil("Men's Restroom", "Men's Restroom 108", 'furring', None, '5/A-121'),
 ceil('Cafe - 11\'-6" GWB band (recess)', 'Cafe 103 - 11\'-6" GWB band', 'furring', None, '5/A-121'),
 ceil('Cafe - 9\'-0" GWB soffit band (storefront)', 'Cafe 103 - 9\'-0" GWB band', 'furring', None, '5/A-121'),
 ceil('Cafe Seating Area - GWB cloud', 'Cafe Seating Area 102 - GWB cloud', 'furring', None, '5/A-121',
      system_note='4/A-121 USG drywall suspension grid is drawn for this cloud (callout on cloud N edge); package carries the 5/A-121 hat channel + 1-1/2" channel system - grid alternative RFI-15'),
 ceil('Cafe - 11\'-0" GWB zone (outer L, at counter)', 'Cafe 103 - 11\'-0" GWB zone', 'framed_soffit', None, '2/A-221',
      system_callout={'detail': '2/A-221', 'cuts_this': True, 'cut_on': ['A-221'], 'source': '2/A-221 Soffit Section - Cafe Sales Area (cuts the 11\'-0" zone on 1/A-221)'},
      drop_lines={'value': 2, 'source': '2/A-221: stud lines up to deck at the outer open edge (N + E) and at the 11\'-6" step (inner)'}),
]

# ---------------------------------------------------------------- aberturas
HSRC = 'A-500 door schedule + 1/A-500 typical hollow metal frame (no header drawn); no header schedule in the set (A-000 index, A-410, A-500)'
def hrfi(ro, pcs=2):
    return (f"Door header not called: if a 2-piece track header is required, +({ro} + 1.333) x {pcs} = "
            f"{round((ro + 16 / 12) * pcs, 2)} LF - RFI-09")
OPN = {  # wall -> [(ref, ro, head, xy, kind, size_src)]
 'W2-01': [('Door 14', 3.0, 7.0, (1898.8, 607.2), 'door', 'A-500 door 14 3\'-0" x 7\'-0"; vector gap 3.00 ft (A-101)')],
 'W2-03': [('Door 11', 4.0, 7.0, (498.5, 1193.5), 'door', 'A-500 door 11 4\'-0" x 7\'-0"; vector gap 4.00 ft')],
 'W2-04': [('Door 4', 3.0, 7.0, (534.7, 1290.45), 'door', 'A-500 door 4 3\'-0" x 7\'-0"; vector gap 3.00 ft')],
 'W2-08': [('Door 5', 3.0, 7.0, (761.7, 1310.3), 'door', 'A-500 door 5 3\'-0" x 7\'-0"; vector gap 3.00 ft')],
 'W2-09': [('Door 6', 3.0, 7.0, (932.95, 1310.3), 'door', 'A-500 door 6 3\'-0" x 7\'-0"; vector gap 3.00 ft')],
 'W2-06.1': [('Door 9', 3.33, 7.0, (579.1, 1339.2), 'door', 'A-500 door 9 3\'-0" x 7\'-0"; vector gap 3.33 ft (frame)')],
 'W2-06.5': [('Door 8', 3.33, 7.0, (700.6, 1339.2), 'door', 'A-500 door 8 3\'-0" x 7\'-0"; vector gap 3.33 ft (frame)'),
             ('Door 7', 3.0, 7.0, (741.9, 1339.2), 'door', 'A-500 door 7 3\'-0" x 7\'-0"; vector gap 3.00 ft')],
 'EXT-03': [('Door 13', 6.33, 7.0, (353.82, 871.7), 'door', 'A-500 door 13 EXISTING V.I.F.; furring gap 6.33 ft (A-101 vector) - RFI-10')],
 'EXT-07': [('Door 10', 3.49, 7.0, (535.7, 1410.84), 'door', 'A-500 door 10 EXISTING V.I.F.; furring gap 3.49 ft (A-101 vector) - RFI-10')],
 'EXT-04': [('Door 12', 10.0, 10.0, (300.8, 1175.3), 'door', 'A-500 door 12 type D EXISTING V.I.F. (overhead); gap 10.00 ft along the angled wall (A-101 vector); head not printed - RFI-10')],
 'C-H2a': [('Door 3', 4.0, 8.0, (1223.65, 587.8), 'door', 'A-500 door 3 4\'-0" x 8\'-0"; vector gap 4.00 ft (A-221)')],
}
items = []
for sh in walls:
    for w in walls[sh]:
        for ref, ro, head, xy, kind, ssrc in OPN.get(w['id'], []):
            w.setdefault('openings', []).append({'ref': ref, 'kind': kind, 'ro_w_ft': ro, 'head_ft': head,
                                                  'source': ssrc, 'xy_pt': list(xy), 'header': 'not_called',
                                                  'header_source': HSRC, 'header_rfi': hrfi(ro),
                                                  'location': 'interior' if not w['tag'].startswith('1') else 'exterior',
                                                  **({'substrate_none_reason': 'Existing exterior door in existing CMU (A-500 EXISTING V.I.F.); no head / jamb detail for the furring at this opening - no substrate printed',
                                                      'jamb_return': False, 'jamb_return_source': 'A-500 (existing door, EXIST. frame; no head / jamb detail for the furring) - gyp bd by future tenant (1/A-410): no jamb return in our scope'} if w['tag'].startswith('1') else {})})
            items.append({'ref': ref, 'wall': w['id'], 'source': ssrc})
opening_schedule = {'sources': ['A-500', 'A-101', 'A-221', 'A-210', 'M-102'],
                    'sources_none': {'S': 'No structural sheets in the bid set (A-000 drawing index: Architectural, Electrical, Mechanical, Plumbing only)'},
                    'items': items + [{'ref': 'Door 1', 'source': 'A-500'}, {'ref': 'Door 2', 'source': 'A-500'}],
                    'none': [{'ref': 'Door 1', 'why': 'Vestibule exterior storefront entrance (EXISTING V.I.F., type A) - storefront system, not in a partition (A-101 C1, A-210)'},
                             {'ref': 'Door 2', 'why': 'Vestibule interior storefront doors (alum/glass, type A) - storefront system (A-101 C1/C3, A-210), not in a partition'}]}

# ---------------------------------------------------------------- tags descartados
descartes = [
 {'id': 'PT0007', 'motivo': 'Type 1 tag on the Cafe south line (A-101): the same line is drawn on enlarged plan 1/A-221 as 6" stud walls E2/E1 (E/A-221 "6" METAL STUDS @16" O.C. TO UNDERSIDE OF DECK", void at existing wall). Enlarged plan governs; counted ONCE as E1/E2 (estimator decision) - RFI-05.', 'owner': 'C-E2', 'imagen': '1_inventario/recortes/PT0007_1_A-101_p6.png'},
 {'id': 'PT0048', 'motivo': 'I tag (SIM.) points to the exposed end face of a wall (finish callout at wall end, I/A-221), not a wall line.', 'imagen': '1_inventario/recortes/PT0048_I_A-221_p15.png'},
 {'id': 'PT0049', 'motivo': 'I tag (SIM.) points to the exposed end face of a wall (finish callout at wall end, I/A-221), not a wall line.', 'imagen': '1_inventario/recortes/PT0049_I_A-221_p15.png'},
 {'id': 'PT0050', 'motivo': 'I tag at the Workroom door pier: marks the inside face finish of the E1 pier (1\'-0" x 2\'-8" HOLD), not a separate wall line - pier counted as E1 (C-E1-P).', 'owner': 'C-E1-P.1', 'imagen': '1_inventario/recortes/PT0050_I_A-221_p15.png'},
]
frag_walls = {sh: {'walls': walls[sh]} for sh in walls}
frag_walls['A-101']['tags_descartados'] = [d for d in descartes if d['id'] == 'PT0007']
frag_walls['A-221']['tags_descartados'] = [d for d in descartes if d['id'] != 'PT0007']
general = {
 'project': 'Barnes & Noble - Colorado Springs, CO (3173 Cinema Point) - bid set 250918',
 'constants': {'plan_scale_pt_per_ft': {'value': 9.0, 'source': 'A-101 1/8"=1\'-0" (grid C-D 30\'-0" CL/CL = 270 pt)'},
               'plan_scale_by_sheet': {'A-221': {'value': 36.0, 'source': '1/A-221 1/2"=1\'-0" (printed 28\'-5 7/8" = 1025.6 pt)'}},
               'carrying_direction': {'value': 'E-W', 'source': 'not printed (5/A-121 N.T.S., no plan arrow on A-121) - assumed E-W, RFI-15 with N-S delta'},
               'furring_oc_in': {'value': 24, 'source': '5/A-121 "7/8" HAT CHANNEL FRAMING @ 24" O.C. MAX"'},
               'carrying_oc_in': {'value': 48, 'source': '5/A-121 "1-1/2" CHANNEL @ 48" O.C. MAX"'},
               'hanger_oc_in': {'value': 48, 'source': '5/A-121 "HANGING WIRE TO STRUCTURE ABOVE AT 48" O.C. MAX"'},
               'bracing_spacing_in': {'value': 48, 'source': '5/A-410 lateral bracing CR 1-1/2" 16 ga @ 4\'-0" O.C. full wall height'}},
 'details': {d: {'layers_per_face': 1, 'source': d} for d in sorted(set(DET.values()) | {'C25/A-101', '2/A-221'})},
 'opening_schedule': opening_schedule,
 'plywood_strips': strips,
 'nodes': nodes,
 'trim_items': trim,
 'ceilings': ceilings,
 'deflection_clips': {'none_reason': 'Deflection by deep leg track / slotted deflection at the top of full height walls (3/A-410 Top Track Detail, 2/A-410 "SLOTTED DEFLECTION", A-410 partition note 7 "PROVIDE DEFLECTION TRACK AS REQUIRED"); no deflection clips called', 'source': '3/A-410'},
 'accesorios_rfi': [
   {'familia': 'stud reinforcement MEP', 'motivo': 'MEP sweep not registrable: mep.registro M-102 (page 48) vs A-101 (page 6) gave 20 / 14 matching wall segments (< 25 required); the HVAC sheet is not used for geometric crossings (mep.py rule). Ducts crossing the full-height type 2 / 5 / E walls must be counted from M-102 by the estimator.',
    'delta': 'per duct crossing a stud wall: + MAX(1, ROUNDUP(duct width / 16")) stud reinforcement plates x stud rows (6/A-410 penetration details); plus 1 duct sleeve each', 'source': 'M-102'}],
 'accesorios_none': [
   {'familia': 'extra framing ATM', 'why': 'No ATM / equipment recess in the set (A-101, A-500, A-221 reviewed)', 'source': 'A-101'},
   {'familia': 'gypsum board jamb return', 'why': 'No opening with a gypsum board return: doors in hollow metal frames (1/A-500); cased opening 2/A-500 is wood jamb + B-1 trim; storefronts are Div 08', 'source': '1/A-500'},
   {'familia': 'access panel', 'why': 'No access panel shown in GWB ceilings or walls (A-121 ceiling key notes 1-11, M-102 and P sheets reviewed)', 'source': 'A-121'},
   {'familia': 'Level 5 finish', 'why': 'No LEVEL 5 finish in A-510 room finish schedule / finish notes (paint finishes PT-1..PT-8 only)', 'source': 'A-510'}],

 'tile_base_none_reason': 'A-510 room finish schedule: ceramic cove tile base in Cafe 103, Men\'s 108, Women\'s 109; no cement board backer drawn behind the base on any partition detail (2/A-410, 5/A-410); cement board is loaded only where E/A-221 and A, B, C, D/A-221 draw it (wall faces) - backer for restroom base / wall tile (A-200, TL-2) = RFI-12 with delta',
 'frp_none_reason': 'FRP-1 (A-510: Cafe Work Room 104 walls, Janitor 107 walls, Cafe 103 below countertop) - no FRP height / extent printed for partitions (A-225 finish plan dashed FRP-1 zones only) - not quantified, RFI-13 with delta',
 'edge_blocking_none_reason': 'No stud blocking at gypsum board edges in the set: details 2/A-410, 3/A-410, 4/A-410, 5/A-410 (lateral bracing CR @ 4\'-0"), 1/A-500, 2/A-500 reviewed - none draws blocking at horizontal board joints (bridging per A-410 partition note 4 is carried as bracing C.R. channel)',
 'deck_breaks_none_reason': 'No stepped deck breaks: deck slopes continuously, B.O. DECK 24\'-0" (2/A-400B, west) to 25\'-7" (2/A-400A, east), 25\'-0" at 1/A-400A and 1/A-400B; interior walls to deck loaded at 25\'-0" - RFI-03',
}
for sh, fr in frag_walls.items():
    json.dump(fr, open(OUT + f'/{sh}.json', 'w'), indent=1, ensure_ascii=False)
json.dump(general, open(OUT + '/general.json', 'w'), indent=1, ensure_ascii=False)
for f in ('_walls_tmp.json', '_strips_tmp.json', '_nodes_tmp.json'):
    os.remove(OUT + '/' + f)
print('fragments written')
