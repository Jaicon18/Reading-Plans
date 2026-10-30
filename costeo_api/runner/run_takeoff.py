#!/usr/bin/env python3
"""run_takeoff.py — runner de take-off por API con la cadena COMPLETA del skill (R80e.37: nunca saltar la cadena).

El modelo SOLO LEE (canon A y canon B independientes + resolución dirigida de diferencias). Los programas calculan:
inventario_tags.py → [lectura A | lectura B] → comparar.py → [resolución] → overlay de decisiones → canon.py
→ tools (inventario, framed_openings, alcance_filtro, markups) → plantilla_excel.py armar (compuerta reglas.py
adentro; si falla, el runner se DETIENE) → reglas.py compuerta (reporte; nunca stub).

Modos
  --dry-run        no llama a la API: usa --canon-ref como lectura A y B (prueba gratis de toda la cadena).
  (sin --dry-run)  llama `claude -p` con ANTHROPIC_API_KEY tomado de $TAKEOFF_API_KEY (el runner nunca imprime
                   ni guarda la clave), --model sonnet, --max-budget-usd por paso, --output-format json,
                   --allowedTools limitadas (nunca bypassPermissions). Tope total --budget (default 20 USD).

Uso (B&N, dry-run):
  python3 costeo_api/runner/run_takeoff.py --dry-run \
      --pdf <bid_set.pdf> --skill <dir del skill v80s> --project-inputs cache_proyectos/barnes_noble/takeoff_work \
      --evidence-root cache_proyectos/barnes_noble --overlay costeo_api/runner/decisiones_BN.json \
      --canon-ref cache_proyectos/barnes_noble/takeoff_work/etapas/4_canon/takeoff_canon.json \
      --compare-with cache_proyectos/barnes_noble/FINAL_FOR_REVIEW_2026-10-24_rev13/Take-off/<rev13>.xlsx \
      --work <dir de trabajo nuevo>
"""
import argparse
import csv
import json
import os
import shutil
import subprocess
import sys
import time

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import overlay  # noqa: E402
import reconcile  # noqa: E402

# Archivos de entrada del proyecto que la lectura (o el estimador) produce además del canon. En dry-run se copian
# del proyecto de referencia; en una corrida real los escribe el paso de lectura (prompts/lectura.md).
PROJECT_INPUTS = ["etapas/4_canon/fo_in.json", "etapas/4_canon/alcance_in.json", "etapas/4_canon/existing_excluded.json",
                  "etapas/4_canon/evidencias_reglas.json", "renglones.json", "takeoff.json",
                  "etapas/PROCESO_CORRIDA.md", "etapas/5_auditoria/AUDIT_COMPLETENESS.md"]   # evidencia de proceso de la lectura
# Herramientas de PROCESO previas al canon (las corre la lectura: arranque, índice, conteo de tags, test ciego, alcances,
# registro). En dry-run se re-ejecutan desde la bitácora de la corrida de referencia, con las rutas reescritas.
PROCESO = ("arranque.py", "indice_set.py", "contar_tags.py", "test_ciego.py")
ALLOWED = "Read,Write,Glob,Grep,Bash(python3:*)"          # nunca bypassPermissions
STEP_BUDGET = {"lectura_A": 6.0, "lectura_B": 6.0, "resolucion": 2.5, "cobertura": 1.5}


class Alto(Exception):
    """El runner se detiene y reporta (R80e.37)."""


def log(msg):
    print(time.strftime("%H:%M:%S"), msg, flush=True)


def run(cmd, cwd, env=None, check=True):
    p = subprocess.run(cmd, cwd=cwd, env=env, capture_output=True, text=True)
    if check and p.returncode != 0:
        raise Alto(f"falló: {' '.join(map(str, cmd[:3]))} ... rc={p.returncode}\n{(p.stdout + p.stderr)[-1500:]}")
    return p


class Runner:
    def __init__(self, a):
        self.a = a
        self.root = os.path.abspath(a.work)
        self.k = os.path.join(self.root, "takeoff_work")        # misma profundidad que el proyecto (rutas ../../)
        self.t = os.path.join(os.path.abspath(a.skill), "tools")
        self.gasto = 0.0
        self.pasos = []
        self.env = dict(os.environ, BITACORA=os.path.join(self.k, "bitacora.jsonl"))

    # ---------------------------------------------------------------- utilidades
    def tool(self, *args, check=True):
        return run([sys.executable, os.path.join(self.t, "bitacora.py"), "run", os.path.join(self.t, args[0]), *args[1:]],
                   self.k, self.env, check)

    def costo(self, paso, usd=0.0, inp=0, out=0, cr=0, cw=0, seg=0.0, modelo="-"):
        self.gasto += usd
        self.pasos.append((paso, round(usd, 4)))
        if self.a.dry_run:                      # dry-run: no ensucia costos.csv (no hubo llamada)
            return
        f = os.path.join(AQUI, "..", "costos.csv")
        nuevo = not os.path.exists(f)
        with open(f, "a", newline="") as fh:
            w = csv.writer(fh)
            if nuevo:
                w.writerow(["pasada", "paso", "modelo", "usd", "input", "output", "cache_read", "cache_write", "segundos"])
            w.writerow(["runner-dry" if self.a.dry_run else "runner", paso, modelo, round(usd, 6), inp, out, cr, cw, round(seg)])

    def claude(self, paso, prompt_file, extra):
        """Una llamada `claude -p`. Nunca en dry-run. Corta si el tope total se pasaría."""
        if self.a.dry_run:
            raise Alto("BUG: claude() llamado en dry-run")
        tope = STEP_BUDGET.get(paso.split(":")[0], 2.0)
        if self.gasto + tope > self.a.budget:
            raise Alto(f"tope total {self.a.budget} USD: gastado {self.gasto:.2f}, el paso {paso} pide hasta {tope}")
        key = os.environ.get("TAKEOFF_API_KEY")
        if not key:
            raise Alto("falta $TAKEOFF_API_KEY")
        skill_md = open(os.path.join(self.a.skill, "SKILL.md"), encoding="utf-8").read()
        prompt = open(prompt_file, encoding="utf-8").read().format(**extra)
        env = dict(self.env, ANTHROPIC_API_KEY=key)
        t0 = time.time()
        p = run(["claude", "-p", prompt, "--model", "sonnet", "--max-budget-usd", str(tope), "--output-format", "json",
                 "--allowedTools", ALLOWED, "--append-system-prompt", skill_md], self.k, env)
        r = json.loads(p.stdout or "{}")
        u = r.get("usage") or {}
        self.costo(paso, float(r.get("total_cost_usd") or 0), u.get("input_tokens", 0), u.get("output_tokens", 0),
                   u.get("cache_read_input_tokens", 0), u.get("cache_creation_input_tokens", 0), time.time() - t0, "sonnet")
        return r

    # ---------------------------------------------------------------- pasos
    def preparar(self):
        if os.path.exists(self.root) and os.listdir(self.root):
            raise Alto(f"--work {self.root} no está vacío (no se pisa nada)")
        for d in ("1_inventario", "2_lectura", "3_revision", "4_canon/out", "5_auditoria", "6_markups", "7_imagenes", "8_excel"):
            os.makedirs(os.path.join(self.k, "etapas", d), exist_ok=True)
        # imágenes de evidencia con su ruta relativa (el canon las cita como ../../<carpeta>/img/...)
        ev = os.path.abspath(self.a.evidence_root)
        for n in sorted(os.listdir(ev)):
            src = os.path.join(ev, n)
            if os.path.isdir(src) and n != "takeoff_work":
                os.symlink(src, os.path.join(self.root, n))
        # imágenes propias del proyecto (4_canon/img, etc.) que el canon cita dentro de takeoff_work
        pin = os.path.abspath(self.a.project_inputs)
        for sub in ("etapas/4_canon/img", "etapas/2_lectura", "etapas/1_inventario", "etapas/3_revision", "etapas/5_auditoria/img",
                    "etapas/0_mapa"):
            s = os.path.join(pin, sub)
            if os.path.isdir(s):
                dst = os.path.join(self.k, sub)
                shutil.rmtree(dst, ignore_errors=True)
                shutil.copytree(s, dst, symlinks=True)
        for f in PROJECT_INPUTS:
            s = os.path.join(pin, f)
            if os.path.exists(s):
                os.makedirs(os.path.dirname(os.path.join(self.k, f)), exist_ok=True)
                shutil.copy2(s, os.path.join(self.k, f))
        log("1/11 workdir listo (evidencia enlazada con su ruta relativa)")

    def proceso(self):
        """Arranque, alcances y registro: SIEMPRE por bitacora.py run en esta corrida (la compuerta lo exige)."""
        pdf = os.path.abspath(self.a.pdf)
        pin = os.path.abspath(self.a.project_inputs)
        if self.a.dry_run:
            for l in open(os.path.join(pin, "bitacora.jsonl")):
                d = json.loads(l)
                if d.get("tool") in PROCESO:
                    args = [str(x).replace(pin, self.k) for x in d.get("args") or []]
                    for x in args:                      # carpetas de salida que la herramienta espera
                        if x.startswith(self.k) and os.path.splitext(x)[1]:
                            os.makedirs(os.path.dirname(x), exist_ok=True)
                    self.tool(d["tool"], *args)
        else:
            self.tool("arranque.py", pdf)
        self.tool("alcances.py", "escanear", pdf, "escaneo.json")
        self.tool("alcances.py", "renglones", "escaneo.json", "renglones_alcances.json")
        self.tool("registro.py", "crear", "registro.json", "renglones.json")
        if self.a.dry_run:
            # el registro verificado (marcas con su evidencia) es producto de la lectura de referencia
            shutil.copy2(os.path.join(pin, "registro.json"), os.path.join(self.k, "registro.json"))
        self.costo("proceso")
        log("2/11 proceso OK: arranque, alcances escanear, registro crear" + (" (+ herramientas de la lectura de referencia)" if self.a.dry_run else ""))

    def inventario_tags(self):
        out = os.path.join(self.k, "etapas", "1_inventario", "runner")
        self.tool("inventario_tags.py", os.path.abspath(self.a.pdf), out, "--sin-recortes")
        self.costo("inventario_tags")
        log("3/11 inventario_tags.py OK")
        return os.path.join(out, "inventario.json")

    def lecturas(self, inv):
        A = os.path.join(self.k, "etapas", "2_lectura", "takeoff_canon_A.json")
        B = os.path.join(self.k, "etapas", "2_lectura", "takeoff_canon_B.json")
        if self.a.dry_run:
            shutil.copy2(self.a.canon_ref, A)
            shutil.copy2(self.a.canon_ref, B)
            self.costo("lectura_A")
            self.costo("lectura_B")
        else:
            pr = os.path.join(AQUI, "prompts", "lectura.md")
            for lado, out in (("A", A), ("B", B)):      # independientes: cada una en su propio proceso, sin ver la otra
                self.claude(f"lectura_{lado}", pr, {"pdf": os.path.abspath(self.a.pdf), "inventario": inv, "out": out,
                                                     "lado": lado})
                if not os.path.exists(out):
                    raise Alto(f"la lectura {lado} no escribió {out}")
        log("4/11 lecturas A y B " + ("(dry-run: canon de referencia)" if self.a.dry_run else "OK"))
        return A, B

    def comparar(self, A, B, inv):
        out = os.path.join(self.k, "etapas", "3_revision", "runner")
        self.tool("comparar.py", A, B, inv, out, check=False)      # rc 1 = hay diferencias / alertas (se reporta)
        rv = json.load(open(os.path.join(out, "revision.json")))
        n = sum(len(rv.get(k) or []) for k in ("difieren", "solo_A", "solo_B"))
        self.costo("comparar")
        log(f"5/11 comparar.py: {n} diferencias")
        return n, out

    def resolver(self, A, n, rev_dir):
        final = os.path.join(self.k, "etapas", "2_lectura", "takeoff_canon_AB.json")
        if n == 0 or self.a.dry_run:
            shutil.copy2(A, final)
            self.costo("resolucion")
        else:                                                     # SOLO las diferencias, con su recorte
            self.claude("resolucion", os.path.join(AQUI, "prompts", "resolucion.md"),
                        {"pdf": os.path.abspath(self.a.pdf), "revision": os.path.join(rev_dir, "revision.json"),
                         "canon_A": A, "out": final})
            if not os.path.exists(final):
                raise Alto("la resolución no escribió el canon")
        log("6/11 resolución " + ("(sin diferencias)" if n == 0 else "OK"))
        return final

    def decisiones(self, canon_ab):
        ov = json.load(open(self.a.overlay)) if self.a.overlay else {}
        self.sin_objeto, self.emparejados = [], None
        if self.a.ref_base:                 # reconciliación de ids por geometría (solo para aplicar decisiones)
            c, mapa, self.sin_objeto = reconcile.aplicar(json.load(open(canon_ab)), json.load(open(self.a.ref_base)), ov)
            self.emparejados = len(mapa)
            log(f"7/11 reconciliación: {len(mapa)} walls emparejados con la referencia; {len(self.sin_objeto)} decisión(es) sin objeto")
        else:
            c = overlay.aplicar_canon(json.load(open(canon_ab)), ov)
        t = overlay.aplicar_takeoff(json.load(open(os.path.join(self.k, "takeoff.json"))), ov)
        dst = os.path.join(self.k, "etapas", "4_canon", "takeoff_canon.json")
        json.dump(c, open(dst, "w"), indent=1, ensure_ascii=False)
        json.dump(t, open(os.path.join(self.k, "takeoff.json"), "w"), indent=1, ensure_ascii=False)
        log("7/11 overlay de decisiones del estimador fusionado (idempotente)")
        return dst

    def calcular(self, canon):
        k = self.k
        c4 = os.path.join(k, "etapas", "4_canon")
        self.tool("canon.py", canon, os.path.join(c4, "out"), "--fo", os.path.join(c4, "fo_in.json"))
        self.tool("inventario.py", canon, os.path.join(c4, "out", "inventario_out.json"))
        self.tool("framed_openings.py", os.path.join(c4, "fo_in.json"), os.path.join(c4, "out", "framed_openings_out.json"))
        self.tool("alcance_filtro.py", os.path.join(c4, "alcance_in.json"), os.path.join(c4, "out", "alcance_out.json"))
        if self.a.sealer:                       # sello closer.py + evidencia planmark.py (script del proyecto, textos de exclusión propios)
            dst = os.path.join(k, "scripts", os.path.basename(self.a.sealer))
            os.makedirs(os.path.dirname(dst), exist_ok=True)
            shutil.copy2(self.a.sealer, dst)
            run([sys.executable, os.path.join(self.t, "bitacora.py"), "run", dst, self.t, os.path.abspath(self.a.pdf)], k, self.env)
        self.tool("vigas_sobre_walls.py", os.path.abspath(self.a.pdf), "NONE", self.a.plan_sheet, canon, "--json",
                  os.path.join(c4, "vigas_out.json"))
        self.tool("markups_bluebeam.py", os.path.abspath(self.a.pdf), canon, os.path.join(k, "etapas", "6_markups", "markups.pdf"),
                  "--json", os.path.join(k, "etapas", "6_markups", "markups.json"))
        log("8/11 canon.py + tools OK (los números salen de los programas)")

    def excel(self):
        x = os.path.join(self.k, "etapas", "8_excel", "Takeoff_runner.xlsx")
        self.tool("plantilla_excel.py", "validar", "takeoff.json")
        p = self.tool("plantilla_excel.py", "armar", "takeoff.json", x, "--familias", "etapas/4_canon/out/familias_out.json",
                      "--trabajo", "etapas", check=False)
        if p.returncode != 0 or not os.path.exists(x):
            raise Alto("plantilla_excel armar NO armó (compuerta reglas.py dentro): ver etapas/REGLAS_FALLAS.md\n"
                       + (p.stdout + p.stderr)[-800:])
        self.tool("registro.py", "entregable", "registro.json")
        self.tool("alcances.py", "compuerta", "registro.json", "escaneo.json", x)
        r = self.tool("reglas.py", "compuerta", "etapas", "--excel", x, check=False)
        fallas = [l for l in r.stdout.splitlines() if l.startswith("FALLA")]
        # La única falla aceptada antes de entregar es la del back-check (paso humano / agente posterior, R80e.30);
        # cualquier otra detiene el runner.
        otras = [l for l in fallas if "back-check" not in l.lower() and "BACK-CHECK" not in l]
        if otras:
            raise Alto("reglas.py compuerta con fallas:\n" + "\n".join(otras[:20]))
        log(f"9/11 Excel armado; reglas.py: {len(fallas)} falla(s) = back-check pendiente" if fallas else "9/11 Excel armado; reglas.py OK")
        return x, len(fallas)

    def comparar_excel(self, x):
        if not self.a.compare_with:
            return None
        import openpyxl
        def load(p):
            wb = openpyxl.load_workbook(p, data_only=True)
            return {s: [list(r) for r in wb[s].iter_rows(values_only=True)] for s in wb.sheetnames if not s.startswith("_")}
        a, b = load(self.a.compare_with), load(x)
        res = {}
        for s in a:
            A, B = a[s], b.get(s, [])
            n = 0
            for i in range(max(len(A), len(B))):
                ra = A[i] if i < len(A) else []
                rb = B[i] if i < len(B) else []
                n += sum(1 for j in range(max(len(ra), len(rb)))
                         if (ra[j] if j < len(ra) else None) != (rb[j] if j < len(rb) else None))
            res[s] = n
        ms = lambda t: {r[0]: r[2] for r in t["Material Summary"][3:] if r and r[0]}
        ma, mb = ms(a), ms(b)
        res["material_summary_rows"] = (len(ma), len(mb))
        res["material_summary_net_diffs"] = [(k, ma.get(k), mb.get(k)) for k in sorted(set(ma) | set(mb), key=str)
                                            if ma.get(k) != mb.get(k)]
        log(f"10/11 comparación contra la referencia: {res}")
        return res

    def cobertura(self, x):
        """Paso 12 reducido: la lectura lista detalles / notas del trade con su destino; cobertura.py verifica."""
        lista = os.path.join(self.k, "etapas", "5_auditoria", "cobertura_lista.json")
        if self.a.dry_run:
            if not self.a.cobertura_lista:
                log("11/11 cobertura: sin --cobertura-lista en dry-run (omitida)")
                return None
            shutil.copy2(self.a.cobertura_lista, lista)
            self.costo("cobertura")
        else:
            self.claude("cobertura", os.path.join(AQUI, "prompts", "cobertura.md"),
                        {"pdf": os.path.abspath(self.a.pdf), "out": lista, "excel": x})
            if not os.path.exists(lista):
                raise Alto("la cobertura no escribió la lista")
        out = os.path.join(self.k, "etapas", "5_auditoria", "cobertura")
        p = run([sys.executable, os.path.join(AQUI, "cobertura.py"), lista,
                 os.path.join(self.k, "etapas", "4_canon", "takeoff_canon.json"), os.path.join(self.k, "takeoff.json"), out], self.k)
        s = json.loads(p.stdout.strip().splitlines()[-1])
        log(f"11/11 cobertura: {s['covered']}/{s['items']} ({s['coverage_pct']} %), {s['gaps']} huecos -> etapas/5_auditoria/cobertura.md")
        return s

    def correr(self):
        self.preparar()
        self.proceso()
        inv = self.inventario_tags()
        A, B = self.lecturas(inv)
        n, rev = self.comparar(A, B, inv)
        canon_ab = self.resolver(A, n, rev)
        canon = self.decisiones(canon_ab)
        self.calcular(canon)
        x, nf = self.excel()
        cmp_ = self.comparar_excel(x)
        cob = self.cobertura(x)
        out = {"reconciliacion": {"emparejados": self.emparejados, "sin_objeto": self.sin_objeto},
               "cobertura": cob, "excel": x, "gasto_usd": round(self.gasto, 4), "pasos": self.pasos, "diferencias_AB": n,
               "reglas_fallas_backcheck": nf, "comparacion": cmp_, "dry_run": self.a.dry_run}
        json.dump(out, open(os.path.join(self.root, "runner_resultado.json"), "w"), indent=1, ensure_ascii=False, default=str)
        return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pdf", required=True)
    ap.add_argument("--skill", required=True)
    ap.add_argument("--project-inputs", required=True)
    ap.add_argument("--evidence-root", required=True)
    ap.add_argument("--overlay")
    ap.add_argument("--canon-ref")
    ap.add_argument("--compare-with")
    ap.add_argument("--work", required=True)
    ap.add_argument("--budget", type=float, default=20.0)
    ap.add_argument("--plan-sheet", default="A-101", help="hoja de planta para vigas_sobre_walls.py")
    ap.add_argument("--ref-base", help="canon de referencia SIN decisiones (ids del estimador) para reconciliar ids")
    ap.add_argument("--cobertura-lista", help="dry-run: lista de detalles / notas del trade con destino (paso 12)")
    ap.add_argument("--sealer", help="script del proyecto que sella con closer.py y marca con planmark.py")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    if a.dry_run and not a.canon_ref:
        ap.error("--dry-run exige --canon-ref")
    try:
        r = Runner(a).correr()
    except Alto as e:
        print("RUNNER DETENIDO:", e)
        return 1
    print(json.dumps(r, indent=1, ensure_ascii=False, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main())
