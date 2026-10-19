#!/usr/bin/env python3
"""reglas.py — v72b: COMPUERTA DE REGLAS INVIOLABLES. Ninguna regla se interpreta ni se omite: reglas.py manda.

Por qué existe (orden del estimador, 2026-10-19: "Las reglas están hechas para respetarse, conviértelas a
inviolables"): con el mismo skill una corrida aplicaba una regla y otra no (3/8" + ballistic en walls ballistic,
anchors solo en exterior walls, ductos que no se enmarcan, PVC trim sin offsets < 6", LF del type K = hatch medido).
Desde v72b TODA regla del texto (SKILL.md + referencias/, catálogo en `reglas_catalogo.py`) tiene tipo
(`reglas_clasificacion.py`) y esta compuerta la hace cumplir:
  (a) verificable con datos → chequeo automático sobre el canon, las salidas de las tools, la bitácora o el Excel;
  (b) requiere juicio de plano → evidencia registrada POR CASO: N/HOJA (u HOJA en una medida de planta) + imagen
      que EXISTE + valor + regla citada; falla si falta o si el valor contradice la regla. La evidencia sale del
      campo del canon que la regla nombra o de `evidencias_reglas.json` [{regla, caso, source, img, valor}];
      un caso que no aplica se declara "NO APLICA: <motivo>" con su fuente;
  (c) redundante / obsoleta → fusionada en otra regla (que sí se revisa) o anulada por una regla posterior.
Una regla sin clasificar, un (b) sin evidencia, un (a) que falla: la compuerta sale rc 1. No hay flag de override.

Quién la llama (no hay forma de entregar Excel sin ella):
  - `plantilla_excel.py armar` la corre EN PROCESO antes de escribir el libro y no escribe nada si no da OK; el
    libro lleva el sello de esta corrida (hoja oculta `_reglas`);
  - `etapas.py compuerta` la corre (Etapa 5) y `bitacora.py compuerta` la corre en modo final con el Excel
    (sello vigente + `plantilla_excel.verificar_excel`).

Uso:
  python3 reglas.py compuerta <work_dir> [--excel <final.xlsx>] [--out <dir>]
        work_dir = la carpeta de etapas ($TAKEOFF_WORK/etapas) o la del proyecto: se busca
        4_canon/takeoff_canon.json (o takeoff_canon.json), las salidas en out/ (familias_out.json …), la
        bitácora y evidencias_reglas.json. Escribe reglas_out.json + REGLAS_FALLAS.md (en --out o en work_dir).
        rc 0 = OK; rc 1 = al menos una falla.
  python3 reglas.py faltantes <work_dir>    lista la evidencia (b) que falta, como plantilla de evidencias_reglas.json
  python3 reglas.py tabla [md]              tabla regla → tipo → chequeo → test (REGLAS_INVIOLABLES.md)
"""
import glob
import hashlib
import json
import math
import os
import re
import shutil
import sys
import tempfile

AQUI = os.path.dirname(os.path.abspath(__file__))
if AQUI not in sys.path:
    sys.path.insert(0, AQUI)

VERSION = "v73"
NHOJA = re.compile(r"\b[0-9A-Z]{1,3}/[A-Z]{1,3}-?\d{1,3}\.\d{1,3}")
HOJA = re.compile(r"\b[A-Z]{1,3}-?\d{1,3}\.\d{1,3}")
NO_APLICA = re.compile(r"^\s*NO APLICA\s*:\s*\S.{10,}", re.I | re.S)
A_OJO = re.compile(r"a ojo|eyeball|by eye|~\s?\d|approx\.?\s?\d|aprox\.?\s?\d|estimated? (by|from) (the )?(image|png)|"
                   r"scaled from (the )?(image|png|raster)|escalad[oa] (de|del) (la )?(imagen|png)", re.I)
ESTADOS = ("CONFIRMADO", "DERIVADO", "PENDIENTE", "CRITERIO")
TOL = 0.05


class ReglasNoOK(Exception):
    """La compuerta no dio OK: NO SE ARMA / NO ENTREGABLE."""


# ----------------------------------------------------------------------------------------------- contexto
def _sha_obj(o):
    return hashlib.sha256(json.dumps(o, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def sha_archivo(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def _primero(*rutas):
    for p in rutas:
        if p and os.path.isfile(p):
            return os.path.abspath(p)
    return None


def buscar_canon(d):
    return _primero(os.path.join(d, "4_canon", "takeoff_canon.json"), os.path.join(d, "takeoff_canon.json"),
                    os.path.join(d, "etapas", "4_canon", "takeoff_canon.json"))


def buscar_out(d, canon_p):
    for c in (os.path.join(d, "out"), os.path.join(d, "..", "out"), os.path.join(os.path.dirname(canon_p or d), "out"),
              os.path.join(os.path.dirname(canon_p or d), "..", "out"), os.path.join(d, "4_canon")):
        if os.path.isfile(os.path.join(c, "familias_out.json")):
            return os.path.abspath(c)
    return None


class Ctx:
    def __init__(self, d, excel=None):
        self.d = os.path.abspath(d)
        self.excel = os.path.abspath(excel) if excel else None
        self.canon_p = buscar_canon(self.d)
        self.canon = None
        self.err_canon = None
        if self.canon_p:
            try:
                self.canon = json.load(open(self.canon_p, encoding="utf-8"))
            except (OSError, ValueError) as e:
                self.err_canon = f"{type(e).__name__}: {e}"
        self.out_dir = buscar_out(self.d, self.canon_p)
        self.salidas = {}
        if self.out_dir:
            for f in ("familias_out.json", "accesorios_out.json", "board_out.json", "studs_out.json",
                      "corridas_out.json", "headers_in.json", "framed_openings_out.json", "vigas_out.json",
                      "contar_tags_out.json"):
                p = os.path.join(self.out_dir, f)
                if os.path.isfile(p):
                    try:
                        self.salidas[f] = json.load(open(p, encoding="utf-8"))
                    except ValueError:
                        self.salidas[f] = None
        cd = os.path.dirname(self.canon_p) if self.canon_p else self.d
        if "vigas_out.json" not in self.salidas:
            pv = _primero(os.path.join(cd, "vigas_out.json"), os.path.join(self.d, "4_canon", "vigas_out.json"))
            if pv:
                self.salidas["vigas_out.json"] = json.load(open(pv, encoding="utf-8"))
        self.fo_p = _primero(os.path.join(cd, "fo_in.json"), os.path.join(self.out_dir or cd, "fo_in.json"),
                             os.path.join(self.d, "fo_in.json"))
        self.ev_p = _primero(os.path.join(cd, "evidencias_reglas.json"), os.path.join(self.d, "evidencias_reglas.json"))
        self.evidencias = []
        self.err_ev = None
        if self.ev_p:
            try:
                ev = json.load(open(self.ev_p, encoding="utf-8"))
                self.evidencias = ev.get("evidencias") if isinstance(ev, dict) else ev
            except ValueError as e:
                self.err_ev = str(e)
        if self.canon and isinstance(self.canon.get("evidencias_reglas"), list):
            self.evidencias = list(self.evidencias or []) + self.canon["evidencias_reglas"]
        self.bit_p = _primero(os.environ.get("BITACORA"), os.path.join(self.d, "bitacora.jsonl"),
                              os.path.join(self.d, "..", "bitacora.jsonl"), os.path.join(cd, "..", "bitacora.jsonl"))
        self.fresco = None            # salidas recalculadas desde el canon (vigencia)
        self.err_fresco = None

    # --- recorridos del canon
    def walls(self, element=None):
        ws = (self.canon or {}).get("walls") or []
        return [w for w in ws if isinstance(w, dict) and (element is None or w.get("element", "wall") in
                                                            ((element,) if isinstance(element, str) else element))]

    def openings(self):
        out = []
        for w in self.walls():
            for o in w.get("openings") or []:
                if isinstance(o, dict):
                    out.append((w.get("id"), o))
        for o in (self.canon or {}).get("openings_extra") or []:
            if isinstance(o, dict):
                out.append((o.get("wall"), o))
        return out

    def resolver_img(self, p):
        """La imagen existe: ruta absoluta, relativa a la carpeta de trabajo / del canon, o el mismo nombre debajo."""
        if not p or not isinstance(p, str):
            return None
        cands = [p, os.path.join(self.d, p)]
        if self.canon_p:
            cands.append(os.path.join(os.path.dirname(self.canon_p), p))
        for c in cands:
            if os.path.isfile(c):
                return c
        base = os.path.basename(p)
        if base:
            for root in (self.d, os.path.dirname(self.canon_p or self.d)):
                hits = glob.glob(os.path.join(root, "**", base), recursive=True)
                if hits:
                    return hits[0]
        return None


# ----------------------------------------------------------------------------------------------- utilidades
def _f(x, dflt=None):
    try:
        if isinstance(x, dict):
            x = x.get("value")
        return float(x)
    except (TypeError, ValueError):
        return dflt


def _txt(*xs):
    return " ".join(json.dumps(x, ensure_ascii=False) if isinstance(x, (dict, list)) else str(x)
                    for x in xs if x is not None)


def es_exterior(w):
    t = str(w.get("tag") or "").upper()
    return t.startswith("EXT") or w.get("location") == "exterior" or w.get("exterior") is True


def es_k(w, canon):
    import chequeos_v72
    return chequeos_v72.es_k(w, chequeos_v72._k_details(canon or {}))


def F(lista, regla, elem, valor, esperado):
    lista.append({"regla": regla, "elemento": str(elem), "valor": str(valor)[:300], "esperado": str(esperado)[:400]})


# =============================================================================================== chequeos (a)
# Cada chequeo: fn(ctx) -> [(elemento, valor, esperado)]. Las reglas que lo usan están en reglas_clasificacion.
CHEQUEOS = {}


def chequeo(nombre, desc):
    def deco(fn):
        CHEQUEOS[nombre] = (fn, desc)
        return fn
    return deco


def _lista(*a):
    return [a]


@chequeo("A_CANON", "existe 4_canon/takeoff_canon.json, es JSON válido y canon.py lo procesa sin error")
def a_canon(ctx):
    if not ctx.canon_p:
        return _lista("takeoff_canon.json", "no está", "4_canon/takeoff_canon.json (Etapa 3 → 4)")
    if ctx.err_canon:
        return _lista(ctx.canon_p, ctx.err_canon, "JSON válido")
    recalcular(ctx)
    if ctx.err_fresco:
        return _lista("canon.py", ctx.err_fresco, "canon.py corre sin error sobre el canon conciliado")
    return []


def recalcular(ctx):
    """Corre canon.py (todas las tools de familias) sobre el canon en una carpeta temporal: la verdad de las salidas."""
    if ctx.fresco is not None or ctx.err_fresco or ctx.canon is None:
        return
    import contextlib
    import io
    import canon as C
    tmp = tempfile.mkdtemp(prefix="reglas_")
    try:
        fo = json.load(open(ctx.fo_p, encoding="utf-8")) if ctx.fo_p else None
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            out = C.calcular(json.loads(json.dumps(ctx.canon)), tmp, fo)
        ctx.fresco = {"familias_out.json": out}
        for f in ("accesorios_out.json", "board_out.json", "studs_out.json", "corridas_out.json", "headers_in.json"):
            p = os.path.join(tmp, f)
            if os.path.isfile(p):
                ctx.fresco[f] = json.load(open(p, encoding="utf-8"))
    except Exception as e:                                        # noqa: BLE001 — cualquier error de canon es falla
        ctx.err_fresco = f"{type(e).__name__}: {e}"
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _sal(ctx, f):
    """Salida vigente: la recalculada si existe (manda), si no la de disco."""
    recalcular(ctx)
    return (ctx.fresco or {}).get(f) or ctx.salidas.get(f) or {}


@chequeo("A_VIGENCIA", "las salidas en disco (familias_out.json …) son las que canon.py da HOY sobre el canon vigente")
def a_vigencia(ctx):
    out = []
    if ctx.canon is None:
        return out
    if not ctx.out_dir or "familias_out.json" not in ctx.salidas:
        return _lista("out/familias_out.json", "no está", "canon.py takeoff_canon.json out/ (Etapa 4) antes del Excel")
    recalcular(ctx)
    if not ctx.fresco:
        return out
    disco = (ctx.salidas.get("familias_out.json") or {}).get("familias") or {}
    hoy = ctx.fresco["familias_out.json"].get("familias") or {}
    for k in sorted(set(disco) | set(hoy)):
        a, b = _f(disco.get(k), 0.0), _f(hoy.get(k), 0.0)
        if abs(a - b) > max(0.01, 1e-6 * abs(b)):
            out.append((f"familias_out.json '{k}'", a, f"{b} (canon.py hoy): salida vencida o editada a mano"))
    return out


@chequeo("A_ETAPAS", "etapas.py compuerta (inventario, lecturas A/B, revisión vigente, tags con wall o alerta)")
def a_etapas(ctx):
    import etapas
    raiz = ctx.d if os.path.isdir(os.path.join(ctx.d, "1_inventario")) or os.path.isdir(os.path.join(ctx.d, "2_lectura")) \
        else (os.path.join(ctx.d, "etapas") if os.path.isdir(os.path.join(ctx.d, "etapas")) else ctx.d)
    fallas, _ = etapas.compuerta(raiz, _desde_reglas=True)
    return [("etapas", x, "etapas.py compuerta sin fallas") for x in fallas]


def _revisar_mod(m, mod, ctx):
    if mod == "chequeos_v72":
        return m.revisar(ctx.canon, ctx.salidas.get("vigas_out.json"))
    if mod == "chequeos_v73":
        return m.revisar(ctx.canon, ctx.salidas.get("familias_out.json"))
    return m.revisar(ctx.canon)


def _de_modulo(mod, prefijo):
    def fn(ctx):
        if ctx.canon is None:
            return []
        m = __import__(mod)
        r = _revisar_mod(m, mod, ctx)
        return [(prefijo, x, "sin falla") for x in r["fallas"] if True]
    return fn


def _filtro(mod, patron, desc, nombre):
    rx = re.compile(patron)

    def fn(ctx):
        if ctx.canon is None:
            return []
        m = __import__(mod)
        r = _revisar_mod(m, mod, ctx)
        return [(x.split(" ")[1] if len(x.split(" ")) > 1 else mod, x, desc) for x in r["fallas"] if rx.search(x)]
    CHEQUEOS[nombre] = (fn, desc)


_filtro("esquinas", r".", "R70.2 / R70.3: cada esquina con un dueño, ningún hueco colineal sin explicar", "A_ESQUINAS")
for _r, _d in (("R71.1", "flecha norte con fachadas"), ("R71.2", "box contado una vez"),
               ("R71.3", "accesorios de bulkhead sobre los vanos"), ("R71.5", "ningún tag descartado por otra tarea")):
    _filtro("chequeos_v71", r"^" + re.escape(_r) + r"\b", _d, "A_" + _r.replace(".", "_"))
for _r, _d in (("R72.1", "type K por hatch medido; sin segmento sin tag tipado K"),
               ("R72.2", "diagonal = promedio de sus dos caras"), ("R72.3", "parte sin tag de un box con part_of"),
               ("R72.4", "vigas_sobre_walls manda"), ("R72.5", "MEP de mep.cruces; FD una vez"),
               ("R72.6", "faces[].side A | B | END"), ("R72.10", "FRP con jogs medidos"),
               ("R72.11", "ATM con alternativas")):
    _filtro("chequeos_v72", r"^" + re.escape(_r) + r"\b", _d, "A_" + _r.replace(".", "_"))
# v73 (arbitraje #21): consistencia con el plano, no solo forma (cota impresa del type K, contorno sin esquinas dobles,
# MEP solo de mep.cruces)
for _r, _d in (("R73.1", "type K: la cota impresa junto al tramo manda (dim_scan de cotas_banda.py)"),
               ("R73.2", "contorno ballistic sin esquinas contadas dos veces (bands de walls K y jogs)"),
               ("R73.3", "plywood_strips con uso (renglón de la plantilla)"),
               ("R73.5", "stud reinforcement MEP = solo cruces de ducto de mep.cruces; lo demás a RFI"),
               ("R73.6", "deflection clips declarados (nota general o none_reason con hoja)"),
               ("R73.7", "segmento sin tag con espesor de otro type no entra al type por defecto"),
               ("R73.8", "la operación escrita cuadra (aritmética del *_op = su valor)")):
    _filtro("chequeos_v73", r"^" + re.escape(_r) + r"\b", _d, "A_" + _r.replace(".", "_"))


@chequeo("A_VIGAS_OUT", "vigas_out.json (salida de vigas_sobre_walls.py) está para revisar R72.4 / E26")
def a_vigas_out(ctx):
    if ctx.canon is not None and "vigas_out.json" not in ctx.salidas:
        return _lista("vigas_out.json", "no está", "vigas_sobre_walls.py → out/vigas_out.json (R72.4 manda la tool)")
    return []


@chequeo("A_ESTADO_DIM", "#0: cada largo / alto medido con estado CONFIRMADO|DERIVADO|PENDIENTE; CONFIRMADO cita la "
                         "hoja; DERIVADO trae su operación o su fuente; nunca 'a ojo' / '~N'")
def a_estado_dim(ctx):
    out = []
    for w in ctx.walls():
        for k in ("length", "height"):
            if w.get(f"{k}_ft") is None:
                continue
            st = str(w.get(f"{k}_status") or w.get("status") or "").upper()
            src = _txt(w.get(f"{k}_source"), w.get(f"{k}_op"))
            if st not in ESTADOS:
                out.append((f"{w.get('id')}.{k}", st or "sin estado", "CONFIRMADO | DERIVADO | PENDIENTE (#0)"))
                continue
            if st == "CONFIRMADO" and not HOJA.search(_txt(w.get(f"{k}_source"))):
                out.append((f"{w.get('id')}.{k}", _txt(w.get(f"{k}_source"))[:80] or "sin fuente",
                            "[CONFIRMADO] cita la dimension impresa con su hoja (#0, 42)"))
            if st == "DERIVADO" and not src.strip():
                out.append((f"{w.get('id')}.{k}", "DERIVADO sin operación ni fuente", "[DERIVADO] con su operación (#0)"))
            if A_OJO.search(src):
                out.append((f"{w.get('id')}.{k}", A_OJO.search(src).group(0), "nunca escalar ni estimar a ojo (#0, 59.6)"))
    return out


@chequeo("A_FUENTES", "todo número medido lleva source (walls, aberturas, ceilings, faces)")
def a_fuentes(ctx):
    out = []
    for w in ctx.walls():
        if w.get("length_ft") is not None and not (w.get("length_source") or w.get("source")):
            out.append((w.get("id"), "length sin source", "source N/HOJA (§4)"))
        if w.get("height_ft") is not None and not (w.get("height_source") or w.get("source")):
            out.append((w.get("id"), "height sin source", "height_source (§4)"))
        for i, fc in enumerate(w.get("faces") or []):
            if isinstance(fc, dict) and fc.get("height_ft") is not None and not (fc.get("height_source") or
                                                                                w.get("height_source")):
                out.append((f"{w.get('id')}.faces[{i}]", "height sin source", "height_source por face (§4)"))
    for wid, o in ctx.openings():
        if not (o.get("source") or o.get("ro_source") or o.get("ro_w_source")):
            out.append((o.get("ref"), "abertura sin source", "source del R.O. (§4, §8.4)"))
    for c in (ctx.canon or {}).get("ceilings") or []:
        if isinstance(c, dict) and not c.get("source"):
            out.append((c.get("id") or c.get("room"), "ceiling sin source", "source (§8.4)"))
    return out


TIPOS_SIN_DETALLE = ("ceiling_transition", "soffit_side", "soffit_drop", "soffit_bottom", "joist", "fascia_box",
                     "canopy_soffit", "storefront_cripples", "door_bulkhead", "vestibule_head", "grille_head",
                     "roof_step", "parapet")


@chequeo("A_DETALLE", "DETALLES: cada wall cita el detalle de su tipo (N/HOJA) y la cita existe en canon.details o en "
                      "su cadena; un wall sin detalle declara por qué")
def a_detalle(ctx):
    out = []
    dets = set(((ctx.canon or {}).get("details") or {}).keys())
    for w in ctx.walls():
        if w.get("element", "wall") in TIPOS_SIN_DETALLE and (w.get("source") or w.get("detail")):
            continue
        d = str(w.get("detail") or "")
        if not d:
            if not (w.get("detail_none_reason") or w.get("part_of")):
                out.append((w.get("id"), "sin detail", "N/HOJA del detalle del tipo (cadena tag → schedule → detalle)"))
            continue
        if not NHOJA.search(d):
            out.append((w.get("id"), d, "detail = N/HOJA (cita legible por máquina, 34)"))
        elif dets and NHOJA.search(d).group(0) not in dets and not w.get("detail_source"):
            out.append((w.get("id"), d, f"el detalle citado está en canon.details (cadena abierta) o detail_source"))
    return out


@chequeo("A_EJES", "EJES: un wall exterior / parapet = un eje / fachada (nunca dos en un número)")
def a_ejes(ctx):
    out = []
    multi = re.compile(r"[,+&/]| and | y ", re.I)
    for w in ctx.walls(("wall", "parapet")):
        if not (es_exterior(w) or w.get("element") == "parapet"):
            continue
        for k in ("facade", "line"):
            v = w.get(k)
            if isinstance(v, str) and multi.search(v.strip()):
                out.append((w.get("id"), f"{k} = {v}", "un eje / una fachada por wall (E28.4, regla EJES)"))
        b = w.get("band") or {}
        if isinstance(b, dict) and str(b.get("ori") or "").upper() not in ("H", "V", "D", "BOX"):
            out.append((w.get("id"), f"band.ori {b.get('ori')}", "H | V | D | BOX"))
    return out


def _rfis(canon):
    out = []
    for k in ("rfis", "accesorios_rfi"):
        for r in (canon or {}).get(k) or []:
            if isinstance(r, dict):
                out.append((k, r))
    return out


@chequeo("A_RFI_DELTA", "AMBIGÜEDAD = RFI con delta: todo RFI del canon trae su delta (valor o alternativas)")
def a_rfi_delta(ctx):
    out = []
    for k, r in _rfis(ctx.canon):
        d = r.get("delta") if r.get("delta") is not None else r.get("delta_txt") or r.get("alternativas") or \
            r.get("deltas") or r.get("delta_lf") or r.get("delta_sf")
        if d in (None, "", [], {}):
            out.append((r.get("id") or r.get("familia") or k, "RFI sin delta", "delta de cada alternativa (§1.7, E23)"))
    return out


@chequeo("A_TAG_IDS", "E28 / R71.12: cada wall tipado lleva tag_ids del inventario (o default_type_source / part_of)")
def a_tag_ids(ctx):
    out = []
    for w in ctx.walls("wall"):
        t = str(w.get("tag") or "").upper()
        if not t or es_exterior(w):
            continue
        if not [x for x in w.get("tag_ids") or [] if x] and not (w.get("default_type_source") or w.get("part_of")
                                                                or w.get("untagged_reason")):
            out.append((w.get("id"), f"{t} sin tag_ids", "tag_ids del inventario, default_type_source o "
                                                          "untagged_reason (segmento sin tag = return del vecino)"))
    return out


@chequeo("A_DOOR_SILL", "Doors: solo header, nunca sill (E18, 53.7, regla inviolable de doors)")
def a_door_sill(ctx):
    out = []
    for wid, o in ctx.openings():
        if str(o.get("kind")) == "door" and (o.get("with_sill") is True or (_f(o.get("sill_pieces"), 0) or 0) > 0):
            out.append((o.get("ref"), f"door con sill (with_sill {o.get('with_sill')}, sill_pieces "
                                      f"{o.get('sill_pieces')})", "doors sin sill: banda de abajo = 0"))
    return out


DUCTO = re.compile(r"\b(duct|ducto|ductwork|fire ?damper|FD-?\d|damper|louver)\b", re.I)


@chequeo("A_DUCTOS", "E12: ductos / FD en walls NO se enmarcan (kind duct|fire_damper, NOT_FRAMED, 0 header / sill / "
                     "jambs; sí plate de reinforcement A2)")
def a_ductos(ctx):
    out = []
    for wid, o in ctx.openings():
        k = str(o.get("kind"))
        ref = _txt(o.get("ref"), o.get("description"))
        if k in ("duct", "fire_damper"):
            if (_f(o.get("pieces"), 0) or 0) > 0 or o.get("with_sill") is True or \
                    str(o.get("header_type") or "").upper() not in ("", "NOT_FRAMED"):
                out.append((o.get("ref"), f"{k} con framing (pieces {o.get('pieces')}, with_sill {o.get('with_sill')},"
                                          f" header_type {o.get('header_type')})", "NOT_FRAMED: 0 header, 0 sill (E12)"))
        elif DUCTO.search(ref) and k not in ("door", "window", "storefront", "atm"):
            out.append((o.get("ref"), f"kind '{k}' en una abertura de ducto / FD", "kind duct | fire_damper → NOT_FRAMED "
                                                                                   "(E12; no se cambia el kind)"))
    for e in (ctx.canon or {}).get("extra_framing") or []:
        if isinstance(e, dict) and (str(e.get("kind")) == "mep" or DUCTO.search(_txt(e.get("ref")))):
            out.append((e.get("ref"), "extra_framing de ducto / MEP", "error (E12, §9 v68): ductos no se enmarcan"))
    return out


@chequeo("A_HEADER_LF", "E1 / E2 / 53.6: header LF = (R.O. + 1'-4\") × piezas × qty (recalculado contra la familia)")
def a_header_lf(ctx):
    hi = _sal(ctx, "headers_in.json")
    fam = (_sal(ctx, "familias_out.json").get("familias") or {}).get("header / box beam")
    if fam is None:
        return []
    tot = 0.0
    for o in hi.get("openings") or []:
        k = str(o.get("kind"))
        if k in ("duct", "fire_damper") or str(o.get("header") or "") in ("steel_by_others", "none", "fits_module"):
            continue
        if str(o.get("estado") or "") == "NOT_FRAMED":
            continue
        ro, pz, q = _f(o.get("ro_w_ft")), _f(o.get("pieces"), 0), _f(o.get("qty"), 1) or 1
        if ro is None:
            continue
        tot += (ro + 16 / 12.0) * pz * q * (_f(o.get("beams"), 1) or 1)
    if abs(tot - _f(fam, 0)) > max(0.5, 0.005 * tot):
        return _lista("header / box beam", _f(fam, 0), f"Σ (R.O. + 1'-4\") × piezas × qty = {round(tot, 2)} (E1, E2)")
    return []


@chequeo("A_BOX4", "E19 / E20 / R71.10: todo box = 4 piezas; 'OR STEEL FRAMING AS REQUIRED' = metal box; nunca ASSUMED")
def a_box4(ctx):
    out = []
    for wid, o in ctx.openings():
        ht = str(o.get("header_type") or "").upper()
        if "BOX" in ht and "NOT_FRAMED" not in ht and (_f(o.get("pieces"), 0) or 0) != 4 and \
                str(o.get("header") or "") != "steel_by_others":
            out.append((o.get("ref"), f"{o.get('header_type')} pieces {o.get('pieces')}", "box = 4 piezas (E20)"))
        if ("OR STEEL" in ht or o.get("or_steel")) and str(o.get("header") or "") == "steel_by_others":
            out.append((o.get("ref"), "OR STEEL → steel_by_others", "metal stud box beam (E19)"))
        if re.search(r"ASSUM|SUPUEST", _txt(o.get("header_type"), o.get("pieces_source")), re.I):
            out.append((o.get("ref"), "header ASSUMED", "0 LF + RFI con delta (R71.10)"))
    return out


@chequeo("A_RUN_ENDS", "E22 / R71.11: CONT. BOX BEAM declara run_ends (jamb_to_jamb | wall_to_wall) con fuente")
def a_run_ends(ctx):
    out = []
    for wid, o in ctx.openings():
        if o.get("continuous_run") or "CONT" in str(o.get("header_type") or "").upper():
            if o.get("run_ends") not in ("jamb_to_jamb", "wall_to_wall") or not o.get("run_ends_source"):
                out.append((o.get("ref"), f"run_ends {o.get('run_ends')}", "run_ends jamb_to_jamb | wall_to_wall + "
                                                                            "run_ends_source (R71.11)"))
    return out


@chequeo("A_STOREFRONT", "E16 / D10: storefront declara structure (+ fuente) y sus dos jambs con remate")
def a_storefront(ctx):
    out = []
    for wid, o in ctx.openings():
        if str(o.get("kind")) != "storefront":
            continue
        if not o.get("structure") or not o.get("structure_source"):
            out.append((o.get("ref"), "storefront sin structure / structure_source", "structure + fuente (E16)"))
        js = o.get("jambs")
        if js is not None and (not isinstance(js, list) or len(js) != 2 or
                               not all(isinstance(j, dict) and j.get("abuts") and j.get("source") for j in js)):
            out.append((o.get("ref"), f"jambs {str(js)[:60]}", "jambs [{abuts, source}] × 2 (D10)"))
    return out


@chequeo("A_SILL", "E18 / 53.7: R.O. > 4'-0\" con sill = DOUBLE sill track; toda abertura con sill (no doors) lleva sill")
def a_sill(ctx):
    out = []
    hi = _sal(ctx, "headers_in.json")
    for o in hi.get("openings") or []:
        if o.get("with_sill") and (_f(o.get("ro_w_ft"), 0) or 0) > 4.0 and (_f(o.get("sill_pieces"), 2) or 0) < 2:
            out.append((o.get("ref"), f"sill_pieces {o.get('sill_pieces')}", "double sill en R.O. > 4'-0\" (E18)"))
    for wid, o in ctx.openings():
        if str(o.get("kind")) in ("window", "atm") and o.get("with_sill") is False and \
                (_f(o.get("sill_ft"), 0) or 0) > 0 and not o.get("with_sill_source"):
            out.append((o.get("ref"), f"{o.get('kind')} con sill_ft {o.get('sill_ft')} y with_sill false",
                        "with_sill true (sill track, 53.7) o with_sill_source"))
    return out


@chequeo("A_TRACK_RUNS", "Toda wall lleva bottom Y top track (LF ÷ largo = 2 × rows); bottom sin descuento de doors (E10)")
def a_track_runs(ctx):
    out = []
    co = _sal(ctx, "corridas_out.json")
    for r in co.get("renglones") or []:
        if r.get("element") not in ("wall", "parapet", "chase", "pier", "column_wrap"):
            continue
        b, t = _f(r.get("bottom_lf"), 0), _f(r.get("top_lf"), 0)
        if b <= 0 or t <= 0:
            out.append((r.get("id"), f"bottom {b} / top {t}", "bottom Y top track > 0 (2 runs mínimo)"))
    if co.get("doors_descontados") or co.get("bottom_deduct_doors"):
        out.append(("corridas", "bottom track con doors descontados", "E10: bottom track en doors NO se descuenta"))
    return out


@chequeo("A_BALLISTIC", "E5 / E6 / R70.5 / D2: todo wall ballistic (type K / detalle k_panel) tiene su face customer y "
                        "board_caras le da ballistic panel Y 3/8\" gyp sobre el tope del panel")
def a_ballistic(ctx):
    out = []
    if ctx.canon is None:
        return out
    dets = (ctx.canon.get("details") or {})
    bo = _sal(ctx, "board_out.json")
    ops = {}
    for c in bo.get("caras") or []:
        ops.setdefault(c.get("wall"), []).append(str(c.get("operacion") or ""))
    for w in ctx.walls("wall"):
        if not es_k(w, ctx.canon):
            continue
        if not [x for x in w.get("tag_ids") or [] if x]:
            continue                                   # R72.1 lo rechaza aparte (segmento sin tag)
        d = dets.get(str(w.get("detail") or "")) or {}
        kp = d.get("k_panel") if isinstance(d, dict) else None
        if not isinstance(kp, dict):
            out.append((w.get("id"), f"detail {w.get('detail')} sin k_panel", "el detalle del type K declara k_panel "
                                                                               "{panel_to_ft, source}"))
            continue
        kc = [f for f in w.get("faces") or [] if isinstance(f, dict) and f.get("k_customer")]
        if len(kc) != 1:
            out.append((w.get("id"), f"{len(kc)} faces k_customer", "una face customer (ballistic + 3/8\") por wall K"))
            continue
        txt = " ".join(ops.get(w.get("id"), []))
        if "ballistic" not in txt:
            out.append((w.get("id"), "board_caras sin ballistic", "ballistic panel = L × min(panel_to, H) (E5, R70.5)"))
        H, pt = _f(kc[0].get("height_ft"), _f(w.get("height_ft"), 0)), _f(kp.get("panel_to_ft"), 0)
        if H is not None and pt is not None and H > pt + 0.01 and not re.search(r"3/8 \S+ x \S+ = [1-9]", txt):
            out.append((w.get("id"), "board_caras sin 3/8\" sobre el panel", "3/8\" gyp = L × (H − panel_to) (E6)"))
    for w in ctx.walls():
        if not es_k(w, ctx.canon) and any(isinstance(f, dict) and f.get("k_customer") for f in w.get("faces") or []):
            d = dets.get(str(w.get("detail") or "")) or {}
            if not (isinstance(d, dict) and d.get("k_panel")):
                out.append((w.get("id"), "k_customer en un wall no K", "k_customer solo en walls de detalle con k_panel"))
    return out


@chequeo("A_ANCHORS", "R72.8 / R71.15 / A13 / A14: expansion anchors solo bottom track de element wall exterior "
                      "(o anchorage con fuente); slab pins el resto; parte sin tag de un box en ninguno sin fuente")
def a_anchors(ctx):
    """Independiente de accesorios.py: recalcula qué walls van a anchors y lo compara con las salidas en disco Y
    con las que da la tool hoy (una salida vieja o editada a mano, o una tool que cambió la regla, falla)."""
    out = []
    if ctx.canon is None:
        return out
    recalcular(ctx)
    import chequeos_v72 as C72
    partes = C72.partes_de_box(ctx.walls())
    ext_lf = int_lf = 0.0
    n_ext = 0
    for w in ctx.walls():
        el = w.get("element", "wall")
        if el not in ("wall", "column_wrap", "pier", "chase"):
            continue
        try:
            lf = float(w["length_ft"]) * int(w.get("rows", 1) or 1)
        except (KeyError, TypeError, ValueError):
            continue
        an = w.get("anchorage")
        if isinstance(an, dict) and an.get("value") in ("anchors", "pins"):
            if not NHOJA.search(_txt(an.get("source"))):
                out.append((w.get("id"), "anchorage sin source N/HOJA", "anchorage {value, source N/HOJA} (R72.8)"))
            if an["value"] == "anchors":
                ext_lf += lf
                n_ext += 1
            else:
                int_lf += lf
            continue
        if w.get("id") in partes:
            continue
        if es_exterior(w) and el == "wall":
            ext_lf += lf
            n_ext += 1
        else:
            int_lf += lf
    for nombre, sal in (("disco", ctx.salidas.get("accesorios_out.json") or {}),
                        ("canon.py hoy", (ctx.fresco or {}).get("accesorios_out.json") or {})):
        for f in sal.get("filas") or []:
            op = str(f.get("operacion") or "")
            if f.get("familia") == "expansion anchors":
                m = re.search(r"exterior (?:sobre losa|bottom track on slab) ([\d.]+) LF \((\d+) walls\)", op)
                if m and (abs(float(m.group(1)) - ext_lf) > 0.05 or int(m.group(2)) != n_ext):
                    out.append((f"expansion anchors ({nombre})", f"{m.group(1)} LF / {m.group(2)} walls",
                                f"{round(ext_lf, 2)} LF / {n_ext} walls: solo element wall exterior (R72.8)"))
            if f.get("familia") == "slab pins":
                m = re.search(r"interior (?:sobre losa|bottom track on slab) ([\d.]+) LF", op)
                if m and abs(float(m.group(1)) - int_lf) > 0.05:
                    out.append((f"slab pins ({nombre})", f"{m.group(1)} LF", f"{round(int_lf, 2)} LF (A13, R72.8)"))
    return out


@chequeo("A_FRP", "D3 / D4 / A6 / R72.10 / 59.10: FRP neto = bruto − doors; inside + outside = vértices, inside − "
                  "outside = 4; > 4 vértices con jogs; PVC trim sin inside corner en jogs < 6\"")
def a_frp(ctx):
    out = []
    for t in (ctx.canon or {}).get("tile_base") or []:
        if not isinstance(t, dict):
            continue
        r = t.get("room")
        if t.get("perimeter_basis") == "net":
            g = _f(t.get("gross_ft"))
            dd = sum(_f(x.get("ro_w_ft"), 0) or 0 for x in t.get("doors") or [] if isinstance(x, dict))
            p = _f(t.get("perimeter_ft"))
            if g is None or p is None or abs(g - dd - p) > 0.05:
                out.append((f"tile_base {r}", f"perimeter {p}, gross {g}, doors {round(dd, 3)}",
                            "perimeter_ft (net) = gross_ft − Σ doors (D3, 59.10)"))
        ic = t.get("inside_corners")
        if isinstance(ic, dict):
            v, nv, oc = _f(ic.get("value")), _f(ic.get("vertices")), _f(ic.get("outside_corners"), 0)
            if v is not None and nv is not None and (abs(v + oc - nv) > 0.01 or abs(v - oc - 4) > 0.01):
                out.append((f"tile_base {r}", f"inside {v} outside {oc} vertices {nv}",
                            "inside + outside = vértices e inside − outside = 4 (D4)"))
            js = ic.get("jogs") if isinstance(ic.get("jogs"), list) else []
            chico = sum(2 if j.get("kind") == "notch" else 1 for j in js if isinstance(j, dict)
                        and _f(j.get("depth_in"), 99) < 6)
            for f in _sal(ctx, "accesorios_out.json").get("filas") or []:
                if f.get("familia") == "PVC trim at FRP" and re.search(rf"room {re.escape(str(r))}\b", str(f.get("operacion"))):
                    m = re.search(r"inside corners (\d+) x", str(f.get("operacion")))
                    if m and v is not None and int(m.group(1)) != int(v) - chico:
                        out.append((f"PVC trim room {r}", f"{m.group(1)} inside corners",
                                    f"{int(v) - chico} (inside corners − {chico} en jogs < 6\", R72.10)"))
    return out


@chequeo("A_MR", "D1 / 55.6 / 59.9 / 60.11: MR solo en la face del respaldo de un toilet / lavatory (wet_fixtures con "
                 "back_source)")
def a_mr(ctx):
    out = []
    wet = {}
    for x in (ctx.canon or {}).get("wet_fixtures") or []:
        if isinstance(x, dict):
            wet.setdefault(x.get("wall"), []).append(x)
            if not x.get("back_source"):
                out.append((x.get("ref"), "wet fixture sin back_source", "back_source N/HOJA (D1)"))
    for w in ctx.walls():
        for i, f in enumerate(w.get("faces") or []):
            if isinstance(f, dict) and re.search(r"\bMR\b|moisture", str(f.get("board") or ""), re.I):
                if w.get("id") not in wet and not f.get("mr_source"):
                    out.append((f"{w.get('id')}.faces[{i}]", f"MR sin fixture en ese wall", "MR solo detrás de toilet / "
                                                                                           "lavatory (wet_fixtures)"))
    return out


@chequeo("A_LAYERS_FACE", "60.2 / R71.13: layers POR FACE (nunca un 'layers' total del wall); cada face con layers")
def a_layers_face(ctx):
    out = []
    for w in ctx.walls():
        if "layers" in w and w.get("faces"):
            out.append((w.get("id"), f"layers {w.get('layers')} a nivel de wall", "layers por face (60.2)"))
        for i, f in enumerate(w.get("faces") or []):
            if isinstance(f, dict) and f.get("board") and not isinstance(f.get("layers"), (int, float)):
                d = ((ctx.canon or {}).get("details") or {}).get(str(w.get("detail") or "")) or {}
                if not (isinstance(d, dict) and d.get("layers_per_face") is not None):
                    out.append((f"{w.get('id')}.faces[{i}]", "sin layers", "layers por face o layers_per_face del detalle"))
    return out


@chequeo("A_WRAP_FOS", "E8 / R70.6 / 59.6 / PT7: column wrap / pier con parts o faces_ft + fos_chain completa con fuente")
def a_wrap_fos(ctx):
    out = []
    for w in ctx.walls(("column_wrap", "pier")):
        if w.get("part_of"):
            continue
        if not (w.get("parts") or w.get("faces_ft")):
            out.append((w.get("id"), "sin parts ni faces_ft", "parts o faces_ft + faces_ft_source (E8)"))
        fc = w.get("fos_chain")
        # v73 (P02 de la prueba 21B): canon.py exige la forma {value: [eslabones ft], source}; la lista suelta de
        # v72b también vale. Las dos tools aceptan ahora la misma forma.
        if isinstance(fc, dict):
            fc = fc.get("value") if str(fc.get("source") or "").strip() else None
        if not w.get("parts") and (not isinstance(fc, list) or not fc):
            out.append((w.get("id"), "sin fos_chain", "fos_chain con TODOS los eslabones impresos (E8)"))
    return out


@chequeo("A_CEILING_SYS", "E14 / R70.1 / D6: ceiling suspendido = furring; otro sistema solo con system_callout "
                          "{detail, on_plan: true, source}")
def a_ceiling_sys(ctx):
    out = []
    for c in (ctx.canon or {}).get("ceilings") or []:
        if not isinstance(c, dict):
            continue
        s = str(c.get("system") or "").lower()
        if s in ("", "furring", "suspension", "suspended", "act", "open", "none", "open_to_structure"):
            continue
        sc = c.get("system_callout")
        if not (isinstance(sc, dict) and sc.get("on_plan") is True and sc.get("detail") and sc.get("source")):
            out.append((c.get("id") or c.get("room"), f"system {s} sin system_callout on_plan",
                        "detalle LLAMADO sobre ese ceiling (R70.1) o furring + RFI"))
    return out


@chequeo("A_CANOPY", "E13 / 58.3: canopy soffit con wall_line_ft medida aparte (nunca L × 2)")
def a_canopy(ctx):
    out = []
    for r in ctx.walls("canopy_soffit") + [x for x in (ctx.canon or {}).get("runs") or [] if isinstance(x, dict)
                                          and x.get("element") == "canopy_soffit"]:
        wl = r.get("wall_line_ft")
        if not isinstance(wl, dict) or wl.get("value") is None or not wl.get("source"):
            out.append((r.get("id"), "sin wall_line_ft {value, source}", "fascia y wall medidas por separado (E13)"))
    return out


@chequeo("A_HANGER", "E4: hanger wire en PCS con altura por room (hanger_pcs en familias_out)")
def a_hanger(ctx):
    fo = _sal(ctx, "familias_out.json")
    if (_f((fo.get("familias") or {}).get("hanger wire"), 0) or 0) > 0 and not fo.get("hanger_pcs"):
        return _lista("hanger wire", "LF sin PCS por room", "PCS por room × altura (E4)")
    return []


@chequeo("A_TRANSICION", "E7 / E9 / D7: transición con kicker + vertical_studs {bod_ft, ceiling_ft} y edges; track "
                         "arriba y abajo (2 líneas)")
def a_transicion(ctx):
    out = []
    for w in ctx.walls("ceiling_transition"):
        vs = w.get("vertical_studs")
        if not (isinstance(vs, dict) and vs.get("bod_ft") is not None and vs.get("ceiling_ft") is not None):
            out.append((w.get("id"), "sin vertical_studs", "verticales a estructura además del kicker (E7, D7)"))
        if not w.get("edges"):
            out.append((w.get("id"), "sin edges", "edges contra otro ceiling (D7)"))
        tl = w.get("track_lines")
        if tl is not None and _f(tl) != 2:
            out.append((w.get("id"), f"track_lines {tl}", "2: arriba y abajo (E9)"))
    return out


@chequeo("A_FACADES", "E25 / 59.11 / 53.5: cada fachada con parapet_t_stud_ft y nailbase_top/bottom (o parapet_face_of)")
def a_facades(ctx):
    out = []
    fac = (ctx.canon or {}).get("facades")
    it = fac.items() if isinstance(fac, dict) else enumerate(fac or [])
    for k, f in it:
        if not isinstance(f, dict):
            continue
        if f.get("parapet_face_of"):
            continue
        if f.get("parapet_t_stud_ft") is None:
            out.append((f"facade {f.get('id', k)}", "sin parapet_t_stud_ft", "T/STUD del parapet con fuente (E25)"))
        if f.get("nailbase_top_ft") is None and f.get("nailbase_none_reason") is None:
            out.append((f"facade {f.get('id', k)}", "sin nailbase_top_ft", "nailbase T/BASE → T/STUD (E25, 59.11)"))
    return out


@chequeo("A_STOPS_BEAM", "E26: stops_under_beam con la marca W de ESA viga (peralte AISC) o motivo + RFI")
def a_stops_beam(ctx):
    out = []
    for w in ctx.walls():
        s = w.get("stops_under_beam")
        if s in (None, False):
            continue
        txt = _txt(s)
        if not re.search(r"\b(W|HSS|C|MC|S)\s?\d", txt) and not re.search(r"motivo|reason|RFI", txt, re.I):
            out.append((w.get("id"), txt[:80], "marca W de la viga (W14x22 …) o motivo + RFI (E26)"))
    return out


FAM_ACC = ("additional head track", "stud reinforcement MEP", "extra framing ATM", "2-1/2\" CFMF coiling grille head",
           "gypsum board jamb return", "PVC trim at FRP", "PAFs and duct sleeves", "access panel", "drywall screws",
           "framing screws", "slab pins", "expansion anchors", "joint compound", "joint tape")


@chequeo("A_ACCESORIOS", "§12 / 52.1: las familias de accesorios salen de accesorios.py, ninguna PENDIENTE; insulation "
                         "EXCLUDED")
def a_accesorios(ctx):
    out = []
    fo = _sal(ctx, "familias_out.json")
    if not fo:
        return out
    est = fo.get("estado_accesorios") or {}
    acc = fo.get("accesorios") or {}
    for f in FAM_ACC:
        if f not in acc and f not in (fo.get("familias") or {}):
            out.append((f, "falta", "familia de accesorios siempre incluida (§12)"))
        if str(est.get(f) or "").upper() == "PENDIENTE":
            out.append((f, "PENDIENTE", "resolver del plano (E23) o accesorios_rfi con delta"))
    if not any("insulation" in str(x).lower() for x in fo.get("excluidas") or []):
        out.append(("insulation", "no figura como EXCLUDED", "insulation EXCLUDED siempre (§7, §12)"))
    return out


@chequeo("A_CLADDING", "60.4 / §7: hat channel / Z-girt / lath del cladding EXCLUDED; nunca en furring channel")
def a_cladding(ctx):
    out = []
    for f in _sal(ctx, "familias_out.json").get("filas") or []:
        if str(f.get("familia")) == "furring channel" and re.search(r"hat channel|z-?girt|lath|cladding|EIFS|stucco",
                                                                    _txt(f.get("source"), f.get("operacion")), re.I):
            out.append((f.get("id") or "furring channel", _txt(f.get("source"))[:80], "EXCLUDED + RFI (60.4)"))
    return out


@chequeo("A_BLOCKING", "58.6 / D5 / 54.3: blocking en LF, material del rótulo (FRT solo rotulado), plies con fuente")
def a_blocking(ctx):
    out = []
    for b in (ctx.canon or {}).get("blocking_items") or []:
        if not isinstance(b, dict):
            continue
        u = str(b.get("unit") or "lf").lower()
        if u not in ("lf", ""):
            out.append((b.get("ref") or b.get("id"), f"unit {u}", "blocking en LF (58.6)"))
        lab = _txt(b.get("label"), b.get("material"))
        if not lab.strip():
            out.append((b.get("ref") or b.get("id"), "sin label / material", "rótulo literal del detalle (D5)"))
        elif re.search(r"\bFRT|FRTW|fire.?retardant", _txt(b.get("material")), re.I) and \
                not re.search(r"\bFRT|FRTW|fire.?retardant", _txt(b.get("label"), b.get("source")), re.I):
            out.append((b.get("ref") or b.get("id"), "FRT sin rótulo FRT", "FRT solo donde el detalle lo rotula (54.3)"))
    return out


PLY_MAT = ("nailbase_5_8_polyiso", "frt_5_8_ext", "frt_5_8", "frt_3_4")


@chequeo("A_PLYWOOD", "§4 v68 / 54.5: plywood_strips[].material de la lista cerrada")
def a_plywood(ctx):
    return [(p.get("ref") or p.get("id"), p.get("material"), "|".join(PLY_MAT))
            for p in (ctx.canon or {}).get("plywood_strips") or [] if isinstance(p, dict) and p.get("material") not in PLY_MAT]


@chequeo("A_CEMENT", "58.7 / §7: tile (cove base incluido) → cement board > 0")
def a_cement(ctx):
    if (ctx.canon or {}).get("tile_base") and (_f((_sal(ctx, "familias_out.json").get("familias") or {})
                                                  .get("cement board"), 0) or 0) <= 0:
        return _lista("cement board", 0, "cementitious backer board bajo todo tile (58.7)")
    return []


@chequeo("A_DECK_BREAKS", "D9 / 59.13: walls que cruzan un deck_break se parten en dos alturas (split)")
def a_deck_breaks(ctx):
    out = []
    c = ctx.canon or {}
    cb = (c.get("constants") or {}).get("deck_breaks")
    if cb is None and not str(c.get("deck_breaks_none_reason") or "").strip():
        return _lista("constants.deck_breaks", "no está", "deck_breaks {value, source} o deck_breaks_none_reason (D9)")
    if isinstance(cb, dict):
        if not HOJA.search(_txt(cb.get("source"))):
            out.append(("constants.deck_breaks", "sin source con hoja", "source N/HOJA (D9)"))
        for br in cb.get("value") or []:
            if isinstance(br, dict) and not br.get("sheet"):
                out.append((str(br.get("coord_pt")), "deck_break sin sheet", "deck_breaks por hoja (D9)"))
    return out


@chequeo("A_JAMB_RETURN", "A5 / R62: toda abertura exterior declara jamb_return true|false + fuente")
def a_jamb_return(ctx):
    out = []
    for wid, o in ctx.openings():
        ext = o.get("location") == "exterior" or str(o.get("kind")) in ("storefront", "window", "atm")
        if ext and (o.get("jamb_return") not in (True, False) or not o.get("jamb_return_source")):
            out.append((o.get("ref"), f"jamb_return {o.get('jamb_return')}", "jamb_return true|false + source (A5)"))
    return out


@chequeo("A_MEP_SWEEP", "A2 / R71.9: barrido MEP REGISTRADO (registro ≥ 25 coincidencias con fuente) por hoja")
def a_mep_sweep(ctx):
    out = []
    sw = (ctx.canon or {}).get("mep_sweep")
    if (ctx.canon or {}).get("walls") and not sw:
        return _lista("mep_sweep", "no está", "mep_sweep {sheets, registro, source} (A2)")
    regs = []
    if isinstance(sw, dict):
        regs = [sw.get("registro")] if sw.get("registro") else list((sw.get("registros") or {}).values()) \
            if isinstance(sw.get("registros"), dict) else (sw.get("registros") or [])
    for r in regs:
        if isinstance(r, dict) and (_f(r.get("matches"), 0) or 0) < 25:
            out.append(("mep_sweep.registro", r.get("matches"), "≥ 25 coincidencias por dirección (R71.9)"))
    return out


@chequeo("A_FO_MODULO", "60.10 / 59.8: hueco que cabe en el módulo (≤ stud bay) → sin framed opening ni header / sill")
def a_fo_modulo(ctx):
    out = []
    bay = 14.375
    for wid, o in ctx.openings():
        if str(o.get("kind")) in ("door", "storefront"):
            continue
        ro = _f(o.get("ro_w_ft"))
        if ro is not None and ro * 12 <= bay and (_f(o.get("pieces"), 0) or 0) > 0:
            out.append((o.get("ref"), f"R.O. {round(ro * 12, 2)}\" con header", "≤ stud bay → 0 + RFI (59.8, 60.10)"))
    return out


@chequeo("A_IDS", "NUMERACIÓN: ids de wall únicos en el canon (una sola numeración)")
def a_ids(ctx):
    vistos, out = set(), []
    for w in ctx.walls():
        i = w.get("id")
        if i in vistos:
            out.append((i, "id repetido", "un id por wall (numeración fija)"))
        vistos.add(i)
    return out


@chequeo("A_DUPLICADOS", "55.2 / 54.2b / 56.5: un elemento se cuenta una vez (misma hoja, mismo type, misma altura, banda "
                         "que se solapa ≥ 80 % = duplicado salvo part_of / same_box_ok)")
def a_duplicados(ctx):
    out = []
    ws = [w for w in ctx.walls() if isinstance(w.get("band"), dict) and not w.get("part_of") and not w.get("same_box_ok")]
    for i, a in enumerate(ws):
        ba = a["band"]
        for b in ws[i + 1:]:
            bb = b["band"]
            if (a.get("sheet"), a.get("tag"), str(ba.get("ori"))) != (b.get("sheet"), b.get("tag"), str(bb.get("ori"))):
                continue
            try:
                c = (min(float(ba["c1"]), float(bb["c1"])) - max(float(ba["c0"]), float(bb["c0"])))
                e = (min(float(ba["hi"]), float(bb["hi"])) - max(float(ba["lo"]), float(bb["lo"])))
                la = min(float(ba["hi"]) - float(ba["lo"]), float(bb["hi"]) - float(bb["lo"]))
            except (KeyError, TypeError, ValueError):
                continue
            if _f(a.get("height_ft")) != _f(b.get("height_ft")):
                continue                                   # filas apiladas de un detalle de sección (60.6)
            if c > 0 and la > 0 and e / la >= 0.8 and a.get("element") == b.get("element"):
                out.append((f"{a.get('id')} / {b.get('id')}", "misma banda", "un elemento se cuenta una vez (55.2)"))
    return out


@chequeo("A_PURLIN", "59.1: todo wall a deck declara purlin true|false con su fuente (dirección del flute)")
def a_purlin(ctx):
    out = []
    for w in ctx.walls("wall"):
        if w.get("to_deck") is True and (w.get("purlin") not in (True, False) or not w.get("purlin_source")):
            out.append((w.get("id"), f"purlin {w.get('purlin')}", "purlin true|false + purlin_source (59.1)"))
    return out


@chequeo("A_BRACING", "53.4 / §7: bracing rows = ROUNDUP(H × 12 / 48) − 1 en toda corrida vertical")
def a_bracing(ctx):
    out = []
    for r in _sal(ctx, "corridas_out.json").get("renglones") or []:
        if r.get("orientation") != "vertical":
            continue
        h, br = _f(r.get("height_ft")), _f(r.get("bracing_rows"))
        if h is None or br is None:
            continue
        esp = max(0, math.ceil(round(h * 12 / 48, 6)) - 1)
        if int(br) != esp and not r.get("bracing_source"):
            out.append((r.get("id"), f"bracing_rows {int(br)} (H {h})", f"{esp} (53.4)"))
    return out


@chequeo("A_PARAPET", "E3 / 59.4 / PT1 / D8: parapet en walls con line, ori, axis, from/to, length_basis stud_cl, facade")
def a_parapet(ctx):
    out = []
    for w in ctx.walls("parapet"):
        for k in ("line", "ori", "axis", "from_ft", "to_ft", "facade"):
            if w.get(k) in (None, ""):
                out.append((w.get("id"), f"sin {k}", "line, ori, axis, from_ft, to_ft, facade (E3, PT1)"))
        if w.get("length_basis") != "stud_cl":
            out.append((w.get("id"), f"length_basis {w.get('length_basis')}", "stud_cl (E3: largo en el eje del stud)"))
    for k in ("parapets", "runs"):
        for x in (ctx.canon or {}).get(k) or []:
            if isinstance(x, dict) and x.get("element") in ("parapet", "column_wrap", "pier", "ceiling_transition") \
                    and not x.get("counted_in_walls"):
                out.append((x.get("id"), f"{x.get('element')} en {k}", "parapet / wrap / pier / transición en walls (D8)"))
    return out


@chequeo("A_LEVEL5", "A15: cada item LEVEL 5 con base_ft {value, source}; tag sin item = PENDIENTE")
def a_level5(ctx):
    out = []
    for it in (ctx.canon or {}).get("level5_items") or []:
        if isinstance(it, dict) and not (isinstance(it.get("base_ft"), dict) and it["base_ft"].get("source")):
            out.append((it.get("ref") or it.get("wall"), "sin base_ft {value, source}", "base_ft con fuente (A15)"))
    return out


@chequeo("A_DESCARTADOS", "§Etapa 2 / R71.5: todo tag descartado con motivo de PLANO e imagen")
def a_descartados(ctx):
    out = []
    for t in (ctx.canon or {}).get("tags_descartados") or []:
        if isinstance(t, dict) and (not str(t.get("motivo") or "").strip() or not (t.get("imagen") or t.get("img"))):
            out.append((t.get("id"), "descarte sin motivo / imagen", "{id, motivo de plano, imagen}"))
    return out


@chequeo("A_CIEGO", "E24: SKILL.md y referencias/ sin valores de proyecto (test_ciego_v69)")
def a_ciego(ctx):
    import test_ciego_v69 as T
    sk = None
    for c in (os.path.join(AQUI, "..", "take-off-metal-frame-drywall"), os.path.join(AQUI, "take-off-metal-frame-drywall")):
        if os.path.isfile(os.path.join(c, "SKILL.md")):
            sk = c
    if sk is None:
        return []
    out = []
    for p in T.archivos([sk]):
        out += [(p, x, "sin valores de proyecto (E24)") for x in T.revisar_texto(open(p, encoding="utf-8").read(), p)]
    return out


def _bit_rows(ctx):
    if not ctx.bit_p:
        return None
    try:
        return [json.loads(l) for l in open(ctx.bit_p, encoding="utf-8") if l.strip()]
    except ValueError:
        return []


PRE_EXCEL = (("arranque.py", None), ("alcances.py", "escanear"), ("registro.py", "crear"), ("inventario_tags.py", None),
             ("canon.py", None), ("contar_tags.py", None), ("vigas_sobre_walls.py", None), ("framed_openings.py", None),
             ("alcance_filtro.py", None))


@chequeo("A_BITACORA", "57.1 / 60.0 / §2: las tools anteriores al Excel corrieron por bitacora.py run con rc 0")
def a_bitacora(ctx):
    rows = _bit_rows(ctx)
    if rows is None:
        return _lista("bitacora.jsonl", "no está", "toda tool por bitacora.py run (57.1); $BITACORA o work_dir")
    out = []
    for tool, sub in PRE_EXCEL:
        c = [r for r in rows if r.get("tool") == tool and (sub is None or (r.get("args") or [])[:1] == [sub])]
        if not c:
            out.append((f"{tool} {sub or ''}".strip(), "nunca corrió por bitacora.py run", "rc 0 en la bitácora"))
        elif c[-1].get("rc") != 0:
            out.append((f"{tool} {sub or ''}".strip(), f"rc {c[-1].get('rc')}", "rc 0 en su última corrida"))
    return out


@chequeo("A_ARCHIVOS", "Etapas 1–3 / 50.x: inventario.md, revision.md, resolucion.md y registro.json en disco")
def a_archivos(ctx):
    out = []
    raiz = ctx.d if os.path.isdir(os.path.join(ctx.d, "3_revision")) else os.path.join(ctx.d, "etapas")
    for rel in (("1_inventario", "inventario.md"), ("3_revision", "revision.md"), ("3_revision", "resolucion.md")):
        if not (os.path.isfile(os.path.join(raiz, *rel)) or os.path.isfile(os.path.join(ctx.d, *rel))):
            out.append(("/".join(rel), "no está", "salida de la etapa en disco (protocolo v69)"))
    if not _primero(os.path.join(ctx.d, "registro.json"), os.path.join(ctx.d, "..", "registro.json"),
                    os.path.join(raiz, "..", "registro.json")):
        out.append(("registro.json", "no está", "registro.py crear antes del primer número (50.1, 56.1)"))
    return out


@chequeo("A_AS_NEEDED", "E11: 'BOX HEADER AS NEEDED' SÍ se cuenta (piezas con fuente) y su tamaño no impreso va a RFI")
def a_as_needed(ctx):
    out = []
    rf = " ".join(_txt(r) for _, r in _rfis(ctx.canon))
    for wid, o in ctx.openings():
        if not re.search(r"AS NEEDED|AS REQ", _txt(o.get("header_type"), o.get("description")), re.I):
            continue
        if (_f(o.get("pieces"), 0) or 0) <= 0 and str(o.get("header") or "") != "steel_by_others":
            out.append((o.get("ref"), f"pieces {o.get('pieces')}", "'AS NEEDED' se cuenta: box = 4 piezas (E11)"))
        if not HOJA.search(_txt(o.get("pieces_source"), o.get("size_source"))) and str(o.get("ref")) not in rf:
            out.append((o.get("ref"), "tamaño sin fuente y sin RFI", "tamaño no impreso = RFI con delta (E11)"))
    return out


@chequeo("A_BANDAS", "SUPER REGLA: toda abertura (no ducto) declara head_ft con fuente y, si no es door, sill_ft con "
                     "fuente (las dos bandas)")
def a_bandas(ctx):
    out = []
    for wid, o in ctx.openings():
        k = str(o.get("kind"))
        if k in ("duct", "fire_damper"):
            continue
        if _f(o.get("head_ft")) is None or not HOJA.search(_txt(o.get("head_source"), o.get("source"))):
            out.append((o.get("ref"), f"head_ft {o.get('head_ft')}", "head_ft + head_source: banda de arriba "
                                                                     "(ceiling / deck − header)"))
        if k not in ("door",) and (_f(o.get("sill_ft")) is None or
                                   not HOJA.search(_txt(o.get("sill_source"), o.get("source")))):
            out.append((o.get("ref"), f"sill_ft {o.get('sill_ft')}", "sill_ft + sill_source: banda de abajo (sill − "
                                                                     "piso; 0 con fuente si va al piso)"))
    return out


def _buscar(ctx, nombre):
    cd = os.path.dirname(ctx.canon_p) if ctx.canon_p else ctx.d
    return _primero(os.path.join(ctx.d, nombre), os.path.join(ctx.d, "..", nombre), os.path.join(cd, nombre),
                    os.path.join(cd, "..", nombre), os.path.join(ctx.out_dir or cd, nombre),
                    os.path.join(os.path.dirname(ctx.excel or cd), nombre))


@chequeo("A_BITACORA_FINAL", "modo final (con el Excel): bitacora.py compuerta — tools obligatorias vigentes (closer, "
                             "planmark, registro entregable, alcances compuerta), registro antes del Excel, back-check")
def a_bitacora_final(ctx):
    if not ctx.excel:
        return []
    import bitacora
    tk, rg = _buscar(ctx, "takeoff.json"), _buscar(ctx, "registro.json")
    if not tk or not rg:
        return _lista("takeoff.json / registro.json", "no están junto al trabajo", "los dos para bitacora.py compuerta")
    if ctx.bit_p:
        os.environ.setdefault("BITACORA", ctx.bit_p)
    return [("bitacora", x, "bitacora.py compuerta sin fallas") for x in
            bitacora.compuerta(tk, rg, ctx.excel, _desde_reglas=True)]


@chequeo("A_CATALOGO", "44.3 / 51.1: toda regla del catálogo clasificada, todo chequeo existe, toda fusión apunta a algo")
def a_catalogo(ctx):
    cat, cls = catalogo_y_clasif()
    out = []
    for r in cat:
        c = cls.get(r["id"])
        if not c:
            out.append((r["id"], "sin clasificar", "tipo a / b / c"))
        elif c["tipo"] in ("a", "b"):
            for ch in c.get("chequeos") or []:
                if ch not in CHEQUEOS:
                    out.append((r["id"], ch, "chequeo implementado en reglas.py"))
        elif not str(c.get("fusion") or "").strip():
            out.append((r["id"], "fusión vacía", "regla vigente que la contiene o la anuló"))
    return out


@chequeo("A_EXCEL", "modo final: el Excel trae el sello de ESTA corrida de reglas y plantilla_excel.verificar_excel pasa")
def a_excel(ctx):
    if not ctx.excel:
        return []
    out = []
    if not os.path.isfile(ctx.excel):
        return _lista(ctx.excel, "no existe", "Excel armado por plantilla_excel.py armar")
    s = leer_sello(ctx.excel)
    if not s:
        out.append((os.path.basename(ctx.excel), "sin sello _reglas", "Excel armado por plantilla_excel.py (corre "
                                                                     "reglas.py y sella el libro)"))
    else:
        esp = huella(ctx)
        for k in ("canon", "familias_out", "evidencias"):
            if s.get(k) != esp.get(k):
                out.append((f"sello.{k}", str(s.get(k))[:12], f"{str(esp.get(k))[:12]}: el canon / las salidas / la "
                                                              f"evidencia cambió después de armar el Excel"))
    import plantilla_excel as P
    out += [("Excel", e, "plantilla_excel.verificar_excel sin errores") for e in
            P.verificar_excel(ctx.excel, [ctx.out_dir] if ctx.out_dir else None)]
    return out


# =============================================================================================== evidencia (b)
def _casos(ctx, trig):
    c = ctx.canon or {}
    W = ctx.walls
    if trig == "corrida":
        return [("corrida", {})]
    if trig == "walls":
        return [(w.get("id"), w) for w in W("wall")]
    if trig == "walls_medidos":
        return [(w.get("id"), w) for w in W("wall") if str(w.get("length_status") or w.get("status") or "").upper()
                != "CONFIRMADO"]
    if trig == "walls_ext":
        return [(w.get("id"), w) for w in W("wall") if es_exterior(w)]
    if trig == "walls_k":
        return [(w.get("id"), w) for w in W("wall") if es_k(w, c)]
    if trig == "walls_deck":
        return [(w.get("id"), w) for w in W("wall") if w.get("to_deck") is True]
    if trig == "walls_2row":
        return [(w.get("id"), w) for w in W() if (_f(w.get("rows"), 1) or 1) >= 2]
    if trig == "walls_cadena":
        rx = re.compile(r"string|cadena|\bsum\b|suma|tramos|segments|\d\s?\+\s?\d", re.I)
        return [(w.get("id"), w) for w in W("wall") if rx.search(_txt(w.get("length_source"), w.get("length_op")))]
    if trig == "vendor_ends":
        return [(x.get("id") or x.get("ref") or f"vendor_end[{i}]", x) for i, x in enumerate(c.get("vendor_ends") or [])
                if isinstance(x, dict)]
    if trig == "recessed":
        return [(x.get("tag") or f"recessed[{i}]", x) for i, x in enumerate(c.get("recessed_sweep") or [])
                if isinstance(x, dict)]
    if trig == "walls_sin_tag":
        return [(w.get("id"), w) for w in W("wall") if not [x for x in w.get("tag_ids") or [] if x]]
    if trig == "walls_alto":
        return [(w.get("id"), w) for w in W() if w.get("height_ft") is not None]
    if trig == "faces":
        return [(f"{w.get('id')}.faces[{i}]", dict(f, _wall=w)) for w in W() for i, f in enumerate(w.get("faces") or [])
                if isinstance(f, dict)]
    if trig == "openings":
        return [(o.get("ref"), o) for _, o in ctx.openings()]
    if trig == "openings_ext":
        return [(o.get("ref"), o) for _, o in ctx.openings() if o.get("location") == "exterior" or
                str(o.get("kind")) in ("storefront", "window", "atm")]
    if trig == "openings_headed":
        return [(o.get("ref"), o) for _, o in ctx.openings() if str(o.get("kind")) not in ("duct", "fire_damper")]
    if trig == "ceilings":
        return [(x.get("id") or x.get("room"), x) for x in c.get("ceilings") or [] if isinstance(x, dict)]
    if trig == "soffits":
        return [(w.get("id"), w) for w in W(("soffit_side", "soffit_drop", "soffit_bottom", "ceiling_bulkhead",
                                             "canopy_soffit", "fascia_box", "joist"))]
    if trig == "transiciones":
        return [(w.get("id"), w) for w in W("ceiling_transition")]
    if trig == "cajas":
        return [(w.get("id"), w) for w in W(("column_wrap", "pier", "chase"))]
    if trig == "parapets":
        return [(w.get("id"), w) for w in W("parapet")]
    if trig == "otros_framing":
        return [(w.get("id"), w) for w in W(("roof_step", "vestibule_head", "grille_head", "grille_box",
                                             "door_bulkhead", "storefront_cripples"))]
    if trig == "facades":
        f = c.get("facades")
        it = f.items() if isinstance(f, dict) else enumerate(f or [])
        return [(f"facade {x.get('id', k) if isinstance(x, dict) else k}", x if isinstance(x, dict) else {}) for k, x in it]
    if trig == "details":
        return [(k, v if isinstance(v, dict) else {}) for k, v in (c.get("details") or {}).items()]
    if trig == "blocking":
        return [(b.get("ref") or b.get("id") or f"blocking[{i}]", b) for i, b in enumerate(c.get("blocking_items") or [])
                if isinstance(b, dict)]
    if trig == "mep":
        return [(p.get("ref") or f"mep[{i}]", p) for i, p in enumerate(c.get("mep_penetrations") or []) if isinstance(p, dict)]
    if trig == "tile_base":
        return [(f"tile_base {t.get('room')}", t) for t in c.get("tile_base") or [] if isinstance(t, dict)]
    if trig == "wet":
        return [(x.get("ref"), x) for x in c.get("wet_fixtures") or [] if isinstance(x, dict)]
    if trig == "sealant":
        return [(x.get("ref") or x.get("id") or f"sealant[{i}]", x) for i, x in enumerate(c.get("sealant_items") or [])
                if isinstance(x, dict)]
    if trig == "trim":
        return [(x.get("ref") or x.get("id") or f"trim[{i}]", x) for i, x in enumerate(c.get("trim_items") or [])
                if isinstance(x, dict)]
    if trig == "rfis":
        return [(r.get("id") or r.get("familia") or f"{k}[{i}]", r) for i, (k, r) in enumerate(_rfis(c))]
    if trig == "pendientes":
        return [(w.get("id"), w) for w in W() if any(str(w.get(k) or "").upper() == "PENDIENTE"
                                                     for k in ("length_status", "height_status", "status"))]
    if trig == "excluidos":
        return [(x.get("ref") or x.get("id") or f"excl[{i}]", x) for i, x in enumerate(c.get("exclusions") or [])
                if isinstance(x, dict)]
    raise KeyError(f"trigger desconocido: {trig}")


def _campo(o, nombres):
    for n in nombres or ():
        v = o
        for parte in n.split("."):
            v = v.get(parte) if isinstance(v, dict) else None
        if v not in (None, "", [], {}):
            return v
    return None


def evidencia_registrada(ctx, regla, caso):
    for e in ctx.evidencias or []:
        if not isinstance(e, dict) or str(e.get("regla")) != regla:
            continue
        cs = e.get("caso")
        if cs == caso or (isinstance(e.get("casos"), list) and caso in e["casos"]):
            return e
    return None


def revisar_b(ctx, regla, spec):
    """-> [(caso, valor, esperado)] para una regla (b)."""
    out = []
    casos = _casos(ctx, spec["trigger"])
    cita = spec.get("cita", "hoja")
    for caso, obj in casos:
        e = evidencia_registrada(ctx, regla, caso)
        if e is not None:
            src, img, val, arch = e.get("source"), e.get("img"), e.get("valor"), e.get("archivo")
        elif spec["trigger"] == "corrida":
            out.append((caso, "sin evidencia registrada", f"evidencias_reglas.json {{regla: '{regla}', caso: 'corrida', "
                                                          f"archivo, valor}} ({spec.get('pide', 'cómo se cumplió')})"))
            continue
        else:
            src, img, val, arch = _campo(obj, spec.get("src")), _campo(obj, spec.get("img", ["img", "_wall.img"])), \
                _campo(obj, spec.get("valor")), None
        val_t = _txt(val)
        if NO_APLICA.match(val_t) and (cita == "archivo" or HOJA.search(_txt(src))):
            continue
        falta = []
        if cita == "archivo":
            if not arch or not (os.path.isfile(arch) or os.path.isfile(os.path.join(ctx.d, str(arch)))):
                falta.append("archivo que exista")
        else:
            rx = NHOJA if cita == "nhoja" else HOJA
            if not rx.search(_txt(src)):
                falta.append("N/HOJA" if cita == "nhoja" else "HOJA (o N/HOJA)")
            if not ctx.resolver_img(img if isinstance(img, str) else (img[0] if isinstance(img, list) and img else None)):
                falta.append(f"imagen que exista ({str(img)[:60] if img else 'sin img'})")
        if not val_t.strip():
            falta.append(f"valor ({spec.get('pide', '')})")
        if falta:
            out.append((caso, "falta " + ", ".join(falta), spec.get("pide", "evidencia por caso")))
            continue
        pro = spec.get("prohibido")
        if pro and re.search(pro, val_t + " " + _txt(src), re.I):
            out.append((caso, re.search(pro, val_t + " " + _txt(src), re.I).group(0), f"contradice la regla: "
                                                                                     f"{spec.get('pide', '')}"))
        req = spec.get("requerido")
        if req and not re.search(req, val_t, re.I):
            out.append((caso, val_t[:80], f"el valor tiene que decir {spec.get('pide', req)}"))
    return out


# =============================================================================================== compuerta
def huella(ctx):
    h = {"version": VERSION}
    h["canon"] = sha_archivo(ctx.canon_p) if ctx.canon_p else None
    fo = os.path.join(ctx.out_dir, "familias_out.json") if ctx.out_dir else None
    h["familias_out"] = sha_archivo(fo) if fo and os.path.isfile(fo) else None
    h["evidencias"] = sha_archivo(ctx.ev_p) if ctx.ev_p else None
    h["reglas"] = hashlib.sha256(b"".join(open(os.path.join(AQUI, f), "rb").read() for f in
                                          ("reglas.py", "reglas_clasificacion.py", "reglas_catalogo.py"))).hexdigest()
    return h


def catalogo_y_clasif():
    import reglas_catalogo as RC
    import reglas_clasificacion as CL
    cat = RC.extraer()
    return cat, CL.clasificar(cat)


def compuerta(d, excel=None, desde=None):
    """-> dict reglas_out. desde: 'etapas' (no vuelve a llamar a etapas.compuerta)."""
    ctx = Ctx(d, excel)
    cat, cls = catalogo_y_clasif()
    fallas, sin_clas = [], []
    a_usados = {}
    for r in cat:
        c = cls.get(r["id"])
        if not c:
            sin_clas.append(r["id"])
            continue
        if c["tipo"] in ("a", "b"):
            for ch in c.get("chequeos") or []:
                a_usados.setdefault(ch, []).append(r["id"])
    for rid in sin_clas:
        F(fallas, rid, "catálogo", "regla sin clasificar", "tipo a/b/c en reglas_clasificacion.py (ninguna regla sin chequeo)")
    resultados = {}
    for ch, reglas_ in sorted(a_usados.items()):
        if ch == "A_ETAPAS" and desde == "etapas":
            continue
        if ch == "A_EXCEL" and not excel:
            continue
        fn = CHEQUEOS.get(ch, (None,))[0]
        if fn is None:
            for rid in reglas_:
                F(fallas, rid, ch, "chequeo inexistente", "función en reglas.py")
            continue
        if ctx.canon is None and ch not in ("A_CANON", "A_ETAPAS", "A_CIEGO", "A_BITACORA", "A_ARCHIVOS"):
            continue
        try:
            res = fn(ctx)
        except Exception as e:                                     # noqa: BLE001 — un chequeo roto es una falla
            res = [(ch, f"{type(e).__name__}: {e}", "el chequeo corre sin error")]
        resultados[ch] = len(res)
        for elem, val, esp in res:
            F(fallas, reglas_[0], elem, val, esp)
            fallas[-1]["reglas"] = reglas_
            fallas[-1]["chequeo"] = ch
    if ctx.canon is not None:
        for r in cat:
            c = cls.get(r["id"])
            if c and c["tipo"] == "b":
                for caso, val, esp in revisar_b(ctx, r["id"], c["spec"]):
                    F(fallas, r["id"], caso, val, esp)
                    fallas[-1]["chequeo"] = "B:" + c["spec"]["nombre"]
    if ctx.err_ev:
        F(fallas, "§1.4", "evidencias_reglas.json", ctx.err_ev, "JSON válido")
    n = {"a": 0, "b": 0, "c": 0}
    for r in cat:
        if r["id"] in cls:
            n[cls[r["id"]]["tipo"]] += 1
    out = {"version": VERSION, "work_dir": ctx.d, "excel": ctx.excel, "ok": not fallas, "n_fallas": len(fallas),
           "reglas": len(cat), "por_tipo": n, "huella": huella(ctx), "fallas": fallas,
           "fallas_por_regla": _por_regla(fallas), "chequeos_corridos": resultados}
    return out


def _por_regla(fallas):
    c = {}
    for f in fallas:
        c[f["regla"]] = c.get(f["regla"], 0) + 1
    return dict(sorted(c.items(), key=lambda x: -x[1]))


def escribir(out, dest):
    os.makedirs(dest, exist_ok=True)
    json.dump(out, open(os.path.join(dest, "reglas_out.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    L = [f"# REGLAS_FALLAS — reglas.py compuerta ({VERSION})", "",
         f"Carpeta: `{out['work_dir']}` · reglas: {out['reglas']} (a {out['por_tipo']['a']} · b {out['por_tipo']['b']} · "
         f"c {out['por_tipo']['c']}) · fallas: **{out['n_fallas']}** · " + ("**OK**" if out["ok"] else "**NO OK**"), ""]
    if out["fallas"]:
        L += ["## Fallas por regla", "", "| Regla | Fallas |", "|---|---|"]
        L += [f"| {k} | {v} |" for k, v in out["fallas_por_regla"].items()]
        L += ["", "## Detalle", "", "| Regla | Wall / renglón | Valor | Qué esperaba |", "|---|---|---|---|"]
        for f in out["fallas"]:
            L.append("| " + " | ".join(str(f[k]).replace("|", "/").replace("\n", " ") for k in
                                       ("regla", "elemento", "valor", "esperado")) + " |")
    open(os.path.join(dest, "REGLAS_FALLAS.md"), "w", encoding="utf-8").write("\n".join(L) + "\n")


def exigir(d, excel=None):
    """Para las tools que entregan: corre la compuerta, escribe sus salidas y lanza ReglasNoOK si no da OK."""
    out = compuerta(d, excel)
    try:
        escribir(out, d)
    except OSError:
        pass
    if not out["ok"]:
        raise ReglasNoOK(f"reglas.py compuerta: {out['n_fallas']} falla(s) — ver {os.path.join(d, 'REGLAS_FALLAS.md')}")
    return out


SELLO_HOJA = "_reglas"


def sellar(xlsx_wb, out):
    ws = xlsx_wb.create_sheet(SELLO_HOJA)
    ws["A1"] = "reglas.py compuerta OK"
    ws["A2"] = json.dumps(out["huella"], sort_keys=True)
    ws.sheet_state = "hidden"


def leer_sello(xlsx):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx, read_only=True)
        if SELLO_HOJA not in wb.sheetnames:
            return None
        ws = wb[SELLO_HOJA]
        vals = [r[0] for r in ws.iter_rows(min_row=1, max_row=2, values_only=True)]
        return json.loads(vals[1]) if vals and vals[0] == "reglas.py compuerta OK" else None
    except Exception:                                              # noqa: BLE001
        return None


def faltantes(d):
    ctx = Ctx(d)
    cat, cls = catalogo_y_clasif()
    pl = []
    for r in cat:
        c = cls.get(r["id"])
        if c and c["tipo"] == "b" and ctx.canon is not None:
            for caso, val, esp in revisar_b(ctx, r["id"], c["spec"]):
                pl.append({"regla": r["id"], "caso": caso, "source": "", "img": "", "valor": "", "archivo": "",
                           "falta": val, "pide": esp})
    return pl


def tabla(fmt="md"):
    cat, cls = catalogo_y_clasif()
    import reglas_clasificacion as CL
    L = ["| Regla | Dónde | Tipo | Chequeo | Test |", "|---|---|---|---|---|"]
    for r in cat:
        c = cls.get(r["id"], {})
        t = c.get("tipo", "?")
        if t == "a":
            ch = ", ".join(c["chequeos"])
        elif t == "b":
            s = c["spec"]
            ch = f"evidencia por caso ({s['nombre']}: {s['trigger']}, cita {s.get('cita', 'hoja')})"
        else:
            ch = c.get("fusion", "")
        L.append(f"| {r['id']} — {r['titulo'][:70]} | {r['archivo']}:{r['linea']} | {t} | {ch.replace('|', '/')} | "
                 f"{CL.test_de(r['id'], c)} |")
    return "\n".join(L)


def main(argv):
    a = list(argv[1:])
    if a[:1] == ["compuerta"] and len(a) >= 2:
        excel = dest = None
        if "--excel" in a:
            i = a.index("--excel")
            excel = a[i + 1]
            del a[i:i + 2]
        if "--out" in a:
            i = a.index("--out")
            dest = a[i + 1]
            del a[i:i + 2]
        out = compuerta(a[1], excel)
        escribir(out, dest or a[1])
        for f in out["fallas"][:200]:
            print(f"FALLA {f['regla']} | {f['elemento']} | {f['valor']} | {f['esperado']}")
        if len(out["fallas"]) > 200:
            print(f"… {len(out['fallas']) - 200} más en REGLAS_FALLAS.md")
        print(f"reglas {out['reglas']} (a {out['por_tipo']['a']} · b {out['por_tipo']['b']} · c {out['por_tipo']['c']})"
              f" · fallas {out['n_fallas']} -> " + ("REGLAS OK" if out["ok"] else "NO OK (NO SE ARMA / NO ENTREGABLE)"))
        return 0 if out["ok"] else 1
    if a[:1] == ["faltantes"] and len(a) == 2:
        json.dump(faltantes(a[1]), sys.stdout, ensure_ascii=False, indent=1)
        return 0
    if a[:1] == ["tabla"]:
        print(tabla())
        return 0
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv))
