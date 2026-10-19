# -*- coding: utf-8 -*-
"""test_bitacora.py -- la 4a compuerta se ve FALLAR en cada hueco (regla 44.3)."""
import os, subprocess, sys, tempfile, json

AQUI = os.path.dirname(os.path.abspath(__file__))
B = os.path.join(AQUI, "bitacora.py")
FALSAS = ["arranque.py", "alcances.py", "registro.py", "plantilla_excel.py", "canon.py", "framed_openings.py", "alcance_filtro.py",
          "inventario.py", "contar_tags.py", "vigas_sobre_walls.py", "test_ciego.py", "reglas.py", "inventario_tags.py"]
sys.path.insert(0, AQUI)


def excel_bueno(ruta):
    """v72b: Excel REAL armado por plantilla_excel.armar sobre el proyecto sintético de fixture_reglas (reglas.py
    compuerta OK en la carpeta) con su familias_out.json: sale sellado."""
    import plantilla_excel as P
    import test_plantilla_v67 as TP
    d = os.path.dirname(ruta)
    fo = json.load(open(os.path.join(d, "out", "familias_out.json")))
    P.armar(P.fusionar(TP.takeoff(), fo), ruta, d)


def preparar(d):
    import fixture_reglas
    fixture_reglas.crear(d)                     # v72b: carpeta de trabajo real (canon, salidas, evidencia)
    for t in FALSAS:  # tools falsas: salen con el rc de $RC_FALSO (0 por defecto)
        open(os.path.join(d, t), "w").write(
            "import os,sys\nsys.exit(int(os.environ.get('RC_FALSO','0')))\n")
    open(os.path.join(d, "cerrar_muros.py"), "w").write("import json\nfrom closer import emit\n")
    open(os.path.join(d, "closer.py"), "w").write("def emit(): pass\n")
    open(os.path.join(d, "marcar.py"), "w").write("import os, planmark\n")
    open(os.path.join(d, "planmark.py"), "w").write("")
    excel_bueno(os.path.join(d, "final.xlsx"))
    os.remove(os.path.join(d, "bitacora.jsonl"))  # v72b: la bitácora de la entrega la escribe el flujo de cada caso
    for f, txt in (("takeoff.json", "{}"), ("registro.json", "{}"), ("fo.json", "{}"),
                   ("plan.pdf", "p"), ("esc.json", "{}"), ("bc.md", "B" * 300)):
        open(os.path.join(d, f), "w").write(txt)


def b(d, *args, rc=None):
    env = dict(os.environ, BITACORA=os.path.join(d, "bitacora.jsonl"))
    if rc is not None:
        env["RC_FALSO"] = str(rc)
    return subprocess.run([sys.executable, B] + list(args), cwd=d, env=env,
                          capture_output=True, text=True)


def flujo(d, saltar=(), orden_tarde=False):
    pasos = [("arranque.py", "plan.pdf"), ("alcances.py", "escanear", "plan.pdf", "esc.json"),
             ("registro.py", "crear", "registro.json", "esc.json"), ("cerrar_muros.py",),
             ("marcar.py", "plan.pdf"), ("inventario_tags.py", "plan.pdf"), ("inventario.py", "esc.json"), ("canon.py", "esc.json"), ("contar_tags.py", "plan.pdf"), ("vigas_sobre_walls.py", "plan.pdf"), ("test_ciego.py", "skill"), ("framed_openings.py", "esc.json"), ("alcance_filtro.py", "esc.json"),
             ("plantilla_excel.py", "validar", "takeoff.json"),
             ("plantilla_excel.py", "armar", "takeoff.json", "final.xlsx", "--familias", "fo.json"),
             ("registro.py", "entregable", "registro.json"),
             ("alcances.py", "compuerta", "registro.json", "esc.json", "final.xlsx"),
             ("reglas.py", "compuerta", ".", "--excel", "final.xlsx")]
    if orden_tarde:
        pasos.append(pasos.pop(2))
    for p in pasos:
        if p[:2] in saltar or p[:1] in saltar:
            continue
        b(d, "run", *p)
    b(d, "backcheck", "bc.md", "final.xlsx")


def gate(d):
    return b(d, "compuerta", "takeoff.json", "registro.json", "final.xlsx")


def caso(nombre, fn, espera_ok, frase=None):
    with tempfile.TemporaryDirectory() as d:
        preparar(d)
        fn(d)
        r = gate(d)
        ok = (r.returncode == 0) == espera_ok and (frase is None or frase in r.stdout)
        print(("PASA " if ok else "FALLA"), nombre)
        if not ok:
            print(r.stdout, r.stderr)
        return ok


def editar_a_mano(d):
    flujo(d, saltar=(("marcar.py",),))  # sin planmark, y luego se "inventa" a mano
    p = os.path.join(d, "bitacora.jsonl")
    rows = open(p).read().splitlines()
    r = json.loads(rows[0]); r["tool"] = "planmark.py"; rows[0] = json.dumps(r)
    open(p, "w").write("\n".join(rows) + "\n")


def excel_cambiado(d):
    flujo(d)
    open(os.path.join(d, "final.xlsx"), "w").write("x2")


def ultima_fallo(d):
    flujo(d)
    b(d, "run", "registro.py", "entregable", "registro.json", rc=1)


def fallo_luego_paso(d):
    b(d, "run", "registro.py", "entregable", "registro.json", rc=1)
    flujo(d)


def bc_viejo(d):
    flujo(d)
    open(os.path.join(d, "final.xlsx"), "w").write("x2")
    b(d, "run", "plantilla_excel.py", "armar", "takeoff.json", "final.xlsx")
    b(d, "run", "alcances.py", "compuerta", "registro.json", "esc.json", "final.xlsx")


def sin_bc(d):
    flujo(d)
    p = os.path.join(d, "bitacora.jsonl")
    rows = open(p).read().splitlines()[:-1]
    open(p, "w").write("\n".join(rows) + "\n")


def bc_vacio(d):
    flujo(d)
    open(os.path.join(d, "bc.md"), "w").write("")


def fuera_de_bitacora(d):
    flujo(d, saltar=(("marcar.py",),))
    subprocess.run([sys.executable, "marcar.py"], cwd=d)


def pipe_head(d):
    """v63: `bitacora.py run tool | head -1` ya no corta la tool: queda anotada con su rc real y el log completo."""
    flujo(d)
    open(os.path.join(d, "ruidosa.py"), "w").write("for i in range(200000): print('x' * 40)\n")
    env = dict(os.environ, BITACORA=os.path.join(d, "bitacora.jsonl"))
    subprocess.run("%s %s run ruidosa.py | head -1" % (sys.executable, B), shell=True, cwd=d, env=env,
                   capture_output=True)
    rows = [json.loads(l) for l in open(os.path.join(d, "bitacora.jsonl")) if l.strip()]
    r = [x for x in rows if x["tool"] == "ruidosa.py"][-1]
    assert r["rc"] == 0 and os.path.getsize(r["log"]) > 8000000, r


def excel_sin_familia(d):
    import openpyxl
    x = os.path.join(d, "final.xlsx")
    wb = openpyxl.load_workbook(x)
    for row in wb["Material Summary"].iter_rows(min_row=4):
        if row[0].value == "Framing screws, pan head self-drilling":
            row[0].value = "screws"
    wb.save(x)
    flujo(d)


def excel_type_viejo(d):
    import openpyxl
    x = os.path.join(d, "final.xlsx")
    wb = openpyxl.load_workbook(x)
    wb["Partition Takeoff"].append(["TYPE PARAPET"])
    wb.save(x)
    flujo(d)


def armar_sin_familias(d):
    flujo(d, saltar=(("plantilla_excel.py", "armar"), ("alcances.py", "compuerta")))
    b(d, "run", "plantilla_excel.py", "armar", "takeoff.json", "final.xlsx")
    b(d, "run", "alcances.py", "compuerta", "registro.json", "esc.json", "final.xlsx")
    b(d, "backcheck", "bc.md", "final.xlsx")


def sin_reglas(d):
    """v72b: la corrida de reglas.py no está en la bitácora -> rechaza."""
    flujo(d, saltar=(("reglas.py", "compuerta"),))


def excel_sin_sello(d):
    """v72b: un Excel escrito por fuera de armar (sello de test) -> reglas --excel lo rechaza."""
    import plantilla_excel as P
    import test_plantilla_v67 as TP
    fo = json.load(open(os.path.join(d, "out", "familias_out.json")))
    P._escribir(P.fusionar(TP.takeoff(), fo), os.path.join(d, "final.xlsx"), {"huella": {"version": "test"}})
    flujo(d)


def canon_cambiado(d):
    """v72b: el canon cambia después de armar el Excel -> sello vencido (y salidas vencidas) -> rechaza."""
    flujo(d)
    p = os.path.join(d, "4_canon", "takeoff_canon.json")
    c = json.load(open(p))
    c["walls"][0]["length_ft"] = 11.0
    json.dump(c, open(p, "w"))


def regla_violada(d):
    """v72b: una door con sill en el canon (regla de doors) -> reglas.py falla -> bitácora rechaza."""
    flujo(d)
    p = os.path.join(d, "4_canon", "takeoff_canon.json")
    c = json.load(open(p))
    c["walls"][0]["openings"] = [{"ref": "D1", "kind": "door", "ro_w_ft": 3.0, "with_sill": True, "sill_pieces": 2,
                                  "head_ft": 7.0, "head_source": "Z1.01", "source": "Z1.01", "header_type": "track",
                                  "pieces": 1, "pieces_source": "1/Z5.01"}]
    json.dump(c, open(p, "w"))


res = [
    caso("25 v72b sin reglas.py en la bitácora -> rechaza", sin_reglas, False, "reglas.py"),
    caso("26 v72b Excel sin sello de reglas -> rechaza", excel_sin_sello, False, "sello"),
    caso("27 v72b canon cambiado después del Excel -> rechaza", canon_cambiado, False, "REGLAS"),
    caso("28 v72b regla violada (door con sill) -> rechaza", regla_violada, False, "REGLAS"),
    caso("21 v67 sin test_ciego -> rechaza", lambda d: flujo(d, saltar=(("test_ciego.py",),)), False, "test_ciego.py"),
    caso("22 v67 Excel sin familia de accesorios -> rechaza", excel_sin_familia, False, "framing screws"),
    caso("23 v67 type fuera de convencion -> rechaza", excel_type_viejo, False, "PARAPET"),
    caso("24 v67 armar sin --familias -> rechaza", armar_sin_familias, False, "--familias"),
    caso("16 sin inventario -> rechaza", lambda d: flujo(d, saltar=(("inventario.py",),)), False, "inventario.py"),
    caso("19 sin contar_tags (v65) -> rechaza", lambda d: flujo(d, saltar=(("contar_tags.py",),)), False, "contar_tags.py"),
    caso("20 sin vigas_sobre_walls (v66, E26) -> rechaza", lambda d: flujo(d, saltar=(("vigas_sobre_walls.py",),)), False,
         "vigas_sobre_walls.py"),
    caso("17 run | head no corta la tool", pipe_head, True, "BITACORA OK"),
    caso("13 sin canon -> rechaza", lambda d: flujo(d, saltar=(("canon.py",),)), False, "canon.py"),
    caso("14 sin framed_openings -> rechaza", lambda d: flujo(d, saltar=(("framed_openings.py",),)), False, "framed_openings.py"),
    caso("15 sin alcance_filtro -> rechaza", lambda d: flujo(d, saltar=(("alcance_filtro.py",),)), False, "alcance_filtro.py"),
    caso("0 sin script que use closer -> rechaza", lambda d: flujo(d, saltar=(("cerrar_muros.py",),)), False, "importa closer"),
    caso("1 flujo completo -> OK", flujo, True, "BITACORA OK"),
    caso("2 sin planmark -> rechaza", lambda d: flujo(d, saltar=(("marcar.py",),)), False, "importa planmark"),
    caso("3 sin alcances compuerta -> rechaza", lambda d: flujo(d, saltar=(("alcances.py", "compuerta"),)), False, "alcances.py compuerta"),
    caso("4 Excel cambiado despues -> rechaza", excel_cambiado, False, "otra version de final.xlsx"),
    caso("5 bitacora editada a mano -> rechaza", editar_a_mano, False, "cadena rota en el renglon 2"),
    caso("6 registro creado tarde -> rechaza", lambda d: flujo(d, orden_tarde=True), False, "56.1"),
    caso("7 ultima corrida rc=1 -> rechaza", ultima_fallo, False, "rc=1"),
    caso("8 fallo y luego paso -> OK", fallo_luego_paso, True),
    caso("9 back-check sobre Excel viejo -> rechaza", bc_viejo, False, "back-check se hizo sobre otra version"),
    caso("10 sin back-check -> rechaza", sin_bc, False, "NO HAY BACK-CHECK"),
    caso("11 informe de back-check vaciado -> rechaza", bc_vacio, False, "informe del back-check"),
    caso("12 tool corrida por fuera no cuenta", fuera_de_bitacora, False, "importa planmark"),
]
print("%d/%d" % (sum(res), len(res)))
sys.exit(0 if all(res) else 1)
