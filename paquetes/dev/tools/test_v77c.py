#!/usr/bin/env python3
"""test_v77c.py — decisiones del usuario U1-U11 (v77c). Cada caso falla con el dato malo y pasa con el bueno.
Datos SINTÉTICOS (types / hojas X / Z, valores inventados). U1 está en test_v77.py (reemplaza a D1)."""
import json
import os
import re
import sys
import tempfile

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
F, N = [], [0]


def caso(nombre):
    def d(fn):
        N[0] += 1
        try:
            fn()
            print("OK  ", nombre)
        except Exception as e:                       # noqa: BLE001
            print("FALLA", nombre, type(e).__name__, e)
            F.append(nombre)
        return fn
    return d


@caso("U7 studs @ O.C. + 1 + extra en esquina L y en intersección T (studs_nodos)")
def _():
    import studs_nodos as S
    w = lambda i: {"id": i, "tag": "Z1", "stud_size": "3-5/8\"", "length_ft": 10.0, "oc_in": 16, "rows": 1,
                   "height_ft": 10.0}
    libre = S.contar({"walls": [w("A")], "nodes": [{"id": "n1", "kind": "free", "wall": "A"},
                                                    {"id": "n2", "kind": "free", "wall": "A"}]})
    pcs = lambda r, i: next(x for x in r["walls"] if x["id"] == i)["pcs_full"]
    assert pcs(libre, "A") == 1 + 8, libre                  # 1 + ROUNDUP(10 / 1.333) = 9
    lt = S.contar({"walls": [w("A"), w("B"), w("C")],
                   "nodes": [{"id": "n1", "kind": "L", "walls": ["A", "B"]},
                             {"id": "n2", "kind": "T", "branch": "C", "through": "A"},
                             {"id": "n3", "kind": "free", "wall": "A"}, {"id": "n4", "kind": "free", "wall": "B"},
                             {"id": "n5", "kind": "free", "wall": "C"}]})
    assert pcs(lt, "A") == 10 and pcs(lt, "C") == 10 and pcs(lt, "B") == 9, lt   # +1 L, +1 T al branch


@caso("U8 grader modo excluir: sin renglón = EXCLUIDA OK; con renglón de acoustical sealant = MAL")
def _():
    import openpyxl
    import eval_takeoff as E
    t = tempfile.mkdtemp()

    def libro(nombre, filas):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Material Summary"
        ws.append(["x"]); ws.append(["Material", "Unit", "Net"])
        for r in filas:
            ws.append(r)
        p = os.path.join(t, nombre); wb.save(p); return p
    base = [["3-5/8\" metal stud (LF of material)", "LF", 100.0],
            ["Sealant at base of gypsum (exterior, not acoustical)", "LF", 20.0]]
    ref = libro("r.xlsx", base)
    bueno = libro("b.xlsx", base)
    malo = libro("m.xlsx", base + [["Acoustical sealant", "LF", 50.0]])
    rg = {"acoustical sealant": {"modo": "excluir"}}
    fb = {x["familia"]: x for x in E.evaluate(ref, bueno, rangos=rg)["familias"]}
    fm = {x["familia"]: x for x in E.evaluate(ref, malo, rangos=rg)["familias"]}
    assert fb["acoustical sealant"]["estado"] == "EXCLUIDA OK" and fb["acoustical sealant"]["puntos"] == 1.0, fb
    assert fm["acoustical sealant"]["estado"].startswith("MAL") and fm["acoustical sealant"]["puntos"] == 0.0, fm
    assert fb["sealant at base of gypsum (exterior)"]["estado"] == "IGUAL"      # el exterior no se toca


@caso("U8 alcances.py: ALC-19 cerrado CUANTIFICADO o con renglón en el Excel = falla; 'EXCLUIDO: ...' pasa")
def _():
    import openpyxl
    import alcances as A
    t = tempfile.mkdtemp()

    def libro(filas):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Material Summary"
        ws.append(["Material", "Unit", "Net"])
        for r in filas:
            ws.append(r)
        p = os.path.join(t, f"x{len(os.listdir(t))}.xlsx"); wb.save(p); return p
    esc = {"alcances": {aid: {"nombre": n, "hallazgos": []} for aid, n, *_ in A.ALCANCES}}

    def reg(nota):
        rr = [{"id": aid, "estado": "verificado", "nota": "NO APLICA: X2.01 sin ese alcance"} for aid, *_ in A.ALCANCES]
        next(r for r in rr if r["id"] == "ALC-19")["nota"] = nota
        p = os.path.join(t, "reg.json")
        json.dump({"renglones": rr}, open(p, "w"))
        return p
    import registro as R
    orig = R.leer
    R.leer = lambda p: json.load(open(p))
    try:
        ok = A.compuerta(reg(A.EXCLUIDOS["ALC-19"][1]), esc, libro([["Sealant at base of gypsum (exterior)", "LF", 5]]))[0]
        assert not [x for x in ok if "ALC-19" in x], ok
        m1 = A.compuerta(reg("CUANTIFICADO: 1/X12.01"), esc, libro([]))[0]
        assert any("ALC-19" in x and "U8" in x for x in m1), m1
        m2 = A.compuerta(reg(A.EXCLUIDOS["ALC-19"][1]), esc, libro([["Acoustical sealant", "LF", 40]]))[0]
        assert any("ALC-19" in x and "U8" in x for x in m2), m2
    finally:
        R.leer = orig


@caso("U8 fusionar: un alcance de acoustical sealant que la corrida cerró CUANTIFICADO sale EXCLUDED")
def _():
    import plantilla_excel as P
    import test_plantilla_v68 as T8
    d = T8.takeoff()
    d["alcances"] = list(d["alcances"]) + [{"id": "ALC-19", "nombre": "Acoustical sealant / sound attenuation",
                                            "estado": "CUANTIFICADO", "nota": "1/X12.01"}]
    out = P.fusionar(d, T8.fo_base())
    a = [x for x in out["alcances"] if "acoustical sealant" in x["nombre"].lower()]
    assert a and all(x["nota"].startswith("EXCLUDED") and x["estado"] == "NO APLICA" for x in a), a


@caso("U9 / U10 PARA REVISAR: floor plan por type A-Z, luego detalles en el orden del plano; n/N con total")
def _():
    import plantilla_excel as P
    lst = [{"fase": "detalle", "orden_plano": 2, "hoja_detalle": "5/X12.01", "punto": "d2"},
           {"fase": "floor_plan", "type": "G1", "punto": "g1", "hoja_detalle": "X2.01"},
           {"fase": "detalle", "orden_plano": 1, "hoja_detalle": "3/X12.01", "punto": "d1"},
           {"fase": "floor_plan", "type": "A10", "punto": "a10", "hoja_detalle": "X2.01"},
           {"fase": "floor_plan", "type": "A2", "punto": "a2", "hoja_detalle": "X2.01"}]
    assert P.revisar_err(lst), "desordenado y sin numerar debía fallar"
    o = P.ordenar_revisar(lst)
    assert [x["punto"] for x in o] == ["a2", "a10", "g1", "d1", "d2"], o
    assert [x["n"] for x in o] == [1, 2, 3, 4, 5] and all(x["total"] == 5 for x in o)
    assert not P.revisar_err(o), P.revisar_err(o)
    assert any("fase" in e for e in P.revisar_err([{"punto": "x", "n": 1, "total": 1}]))
    assert any("orden_plano" in e for e in P.revisar_err([{"fase": "detalle", "hoja_detalle": "X12.01", "n": 1,
                                                           "total": 1}]))


@caso("U9 / U10 verificar_excel: Review Items sin total, n/N roto o detalles antes del floor plan = error")
def _():
    import openpyxl
    import plantilla_excel as P
    t = tempfile.mkdtemp()

    def libro(filas, total=True):
        wb = openpyxl.Workbook(); wb.active.title = "Material Summary"
        ws = wb.create_sheet(P.HOJA_REVISAR)
        ws.append([f"{len(filas)} questions in total, asked ONE at a time" if total else "Review"])
        ws.append(["#", "Phase", "Type / order", "Item", "Sheet + detail", "Image"])
        for r in filas:
            ws.append(r)
        ws = wb.create_sheet(P.HOJA_ALCANCES)
        ws.append(["ALC-19", "Acoustical sealant", "NOT APPLICABLE", "EXCLUDED in all projects"])
        p = os.path.join(t, f"r{len(os.listdir(t))}.xlsx"); wb.save(p); return wb
    bueno = [["1/3", "Floor plan", "A1", "x"], ["2/3", "Floor plan", "B", "x"], ["3/3", "Detail", 1, "x"]]
    assert not P.verificar_v77c(libro(bueno)), P.verificar_v77c(libro(bueno))
    assert any("U10" in e for e in P.verificar_v77c(libro([bueno[2], bueno[0], bueno[1]][:0] +
                                                          [["1/3", "Detail", 1, "x"], ["2/3", "Floor plan", "A1", "x"],
                                                           ["3/3", "Floor plan", "B", "x"]])))
    assert any("U9" in e for e in P.verificar_v77c(libro([["1/3", "Floor plan", "A1", "x"], ["3/3", "Floor plan", "B", "x"],
                                                           ["3/3", "Detail", 1, "x"]])))
    assert any("U9" in e for e in P.verificar_v77c(libro(bueno, total=False)))


@caso("U9 / U10 / U11 / U1 / U8: el SKILL escribe las reglas (sin números de un proyecto)")
def _():
    sk = None
    for c in (os.path.join(AQUI, "..", "SKILL.md"), os.path.join(AQUI, "..", "take-off-metal-frame-drywall", "SKILL.md")):
        if os.path.isfile(c):
            sk = open(c, encoding="utf-8").read()
            break
    assert sk, "SKILL.md no encontrado"
    for r in ("R77.11", "R77.12", "R77.13", "R77.14", "R77.15", "R77.16"):
        assert re.search(r"\| " + re.escape(r) + r" \|", sk), r
    sec = sk[sk.index("## 25."):]
    sec = sec[:sec.index("\n# ", 1)] if "\n# " in sec[1:] else sec
    assert re.search(r"preguntas en total", sec) and re.search(r"ALFAB", sec) and re.search(r"FLOOR PLAN", sec)
    assert not re.search(r"\b\d+\.\d{2}\b", sec), "números de un proyecto en §25"


print(f"{N[0] - len(F)} OK, {len(F)} falla(s)")
sys.exit(1 if F else 0)
