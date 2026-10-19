#!/usr/bin/env python3
"""Tests sintéticos v67 de plantilla_excel.py (E24: números inventados): fusión de accesorios en Material Summary,
hoja Partition Takeoff con la convención de types y verificar_excel (lo que exige bitacora.py compuerta)."""
import copy
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl  # noqa: E402
import accesorios as A  # noqa: E402
import partition_types as PT  # noqa: E402
import plantilla_excel as P
# v72b: estos tests revisan el FORMATO del libro con familias sintéticas; armar() exige reglas.py compuerta OK
# sobre una carpeta de trabajo real. El libro de aquí lleva un sello de test, que bitacora / reglas --excel rechazan.
SELLO_TEST = {"huella": {"version": "test"}}  # noqa: E402
import test_accesorios as TA  # noqa: E402

OK = []


def chk(n, c, m=""):
    if not c:
        raise AssertionError(f"{n}: {m}")
    OK.append(n)


def familias_out(canon=None):
    c = canon or TA.canon_base()
    c["constants"]["level5_tags"] = {"value": ["WC-9"], "source": "X0.40"}
    c.setdefault("accesorios_none", [{"familia": "extra framing ATM", "why": "sin ATM", "source": "X2.01"}])
    fb = TA.fam_base()
    acc = A.calcular(fb, c, TA.STUDS)
    fo = dict(fb, filas=list(fb.get("filas") or []) + list(acc.get("filas") or []),   # como canon.py
              accesorios={k: acc[k] for k in ("familias", "estado", "unidades", "waste", "nombres_excel")},
              rfis_accesorios=acc["rfis"], excluidas=acc["excluidas"])
    for w in c["walls"]:
        w.setdefault("detail", "1/X12.01")
    fo["partition_types"] = PT.partition_types(c)
    return fo


def takeoff():
    d = copy.deepcopy(P.EJEMPLO)
    d["materials"].append({"material": "Level 5 finish (a mano)", "unit": "SF", "net": 1.0, "metal": False,
                           "source": "x"})
    return d


def test_fusion_material_summary():
    fo = familias_out()
    d = P.fusionar(takeoff(), fo)
    nombres = [m["material"] for m in d["materials"]]
    chk("renglón a mano reemplazado", "Level 5 finish (a mano)" not in nombres)
    for f, n in fo["accesorios"]["nombres_excel"].items():
        chk(f"renglón {f}", nombres.count(n) == 1, nombres)
    m = {x["material"]: x for x in d["materials"]}
    chk("yield sin waste", m["Drywall screws 1-1/4\" bugle head"].get("waste") == "none")
    chk("metal con waste", m["Additional head track at openings over 4'-0\" wide"]["metal"] is True)
    chk("insulation EXCLUDED en Alcances", any("insulation" in a["nombre"] for a in d["alcances"]))
    chk("RFIs de accesorios", any(r["id"].startswith("RFI-ACC") for r in d["rfis"]))
    chk("validar OK", P.validar(d) == [], P.validar(d))


def test_net_del_canon_vigente():
    fo = familias_out()
    fo["familias"]["6\" stud (todos los gauges)"] = 321.0
    d = takeoff()
    d["materials"].append({"material": "6\" metal stud", "unit": "LF", "net": 1.0, "metal": True, "stock": 20,
                           "cut_ft": 10, "source": "familias_out.json (canon.py): '6\" stud (todos los gauges)'; 1/X12.01"})
    m = {x["material"]: x for x in P.fusionar(d, fo)["materials"]}
    chk("NET del canon vigente", m["6\" metal stud"]["net"] == 321.0, m["6\" metal stud"])


def test_pendiente_no_se_arma():
    c = TA.canon_base()
    del c["mep_sweep"]
    try:
        P.fusionar(takeoff(), familias_out(c))
        chk("PENDIENTE bloquea", False)
    except ValueError as e:
        chk("PENDIENTE bloquea", "PENDIENTE" in str(e), str(e))


def test_excel_convencion():
    d = P.fusionar(takeoff(), familias_out())
    with tempfile.TemporaryDirectory() as t:
        x = os.path.join(t, "t.xlsx")
        P._escribir(d, x, SELLO_TEST)
        chk("verificar_excel OK", P.verificar_excel(x) == [], P.verificar_excel(x))
        wb = openpyxl.load_workbook(x)
        ws = wb["Partition Takeoff"]
        filas = [r for r in ws.iter_rows(values_only=True)]
        tot = [r for r in filas if "TOTAL - PARTITIONS" in [str(v) for v in r]]
        chk("TOTAL con col A vacía", tot and tot[0][0] is None and tot[0][1] == "TOTAL - PARTITIONS")
        tipos = [str(r[0])[5:] for r in filas if str(r[0] or "").startswith("TYPE ")]
        chk("EXT-6CFMF / EXT-PARAPET / BULKHEAD / OTHER", {"EXT-6CFMF", "EXT-PARAPET", P.OTHER_TYPE} <= set(tipos)
            and any(t_.startswith("BULKHEAD-") for t_ in tipos), tipos)
        chk("OTHER FRAMING después del TOTAL", tipos[-1] == P.OTHER_TYPE)
        ms = wb["Material Summary"]
        r = [row for row in ms.iter_rows(values_only=True) if row[0] == "Joint tape, paper"][0]
        chk("waste none = 0", r[3] == 0, r)
        # Excel con type viejo y TOTAL en la col. A -> verificar lo rechaza
        ws.append(["TYPE PARAPET"])
        ws.append(["TOTAL - X", None, 1])
        y = os.path.join(t, "malo.xlsx")
        wb.save(y)
        e = P.verificar_excel(y)
        chk("type viejo rechazado", any("PARAPET" in x_ for x_ in e), e)
        chk("TOTAL en col A rechazado", any("TOTAL" in x_ for x_ in e), e)
        wb = openpyxl.load_workbook(x)
        ms = wb["Material Summary"]
        for row in ms.iter_rows(min_row=4):
            if row[0].value == "Joint tape, paper":
                row[0].value = "tape"
        wb.save(y)
        chk("familia faltante rechazada", any("joint tape" in x_ for x_ in P.verificar_excel(y)))


if __name__ == "__main__":
    for n, f in list(globals().items()):
        if n.startswith("test_"):
            f()
    print(f"test_plantilla_v67: {len(OK)} checks OK")
