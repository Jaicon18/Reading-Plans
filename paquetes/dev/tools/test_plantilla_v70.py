#!/usr/bin/env python3
"""Tests v70 de plantilla_excel.py verificar con un workbook y salidas de tool SINTÉTICOS (ningún proyecto real).
Correr: python3 test_plantilla_v70.py

Renglón del Material Summary que cita '<tool>_out.json ... <familia>' con Net 0 mientras la salida da > 0 -> FALLA.
"""
import json
import os
import sys
import tempfile

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import openpyxl  # noqa: E402
import plantilla_excel as P  # noqa: E402


def libro(ruta, filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Material Summary"
    ws.append(["MATERIAL SUMMARY"])
    ws.append(["nota"])
    ws.append(["Material", "Unit", "Net", "Waste %", "Waste", "Purchase", "Selling unit", "Pieces to buy",
               "Where it comes from", "Yield check (live formula)"])
    for nombre, net, src in filas:
        ws.append([nombre, "LF", net, 0.15, None, None, "stick 10 ft", None, src, None])
    pt = wb.create_sheet("Partition Takeoff")
    for _ in range(4):
        pt.append([None])
    pt.append(["TYPE X1", "row", 1])
    inp = wb.create_sheet("Inputs")
    inp.append(["Waste metal", 0.15])
    wb.save(ruta)


def salidas(d, fo_lf, fam_lf, acc=0.0):
    json.dump({"totales_lf": {"framed openings MEP": fo_lf}, "renglones": []},
              open(os.path.join(d, "framed_openings_out.json"), "w"))
    json.dump({"familias": {"sill track": fam_lf}, "accesorios": {"familias": {"slab pins": acc}}},
              open(os.path.join(d, "familias_out.json"), "w"))
    json.dump({"familias": {"slab pins": acc}}, open(os.path.join(d, "accesorios_out.json"), "w"))


FILAS = [("Framed openings MEP", 0, "framed_openings_out.json totales_lf 'framed openings MEP'; Z0.01; Z1.01"),
         ("Sill track", 0, "familias_out.json 'sill track'; 2/Z5.20"),
         ("Slab pins", 0, "accesorios_out.json 'slab pins'; 7/Z5.20")]


def test_renglon_en_cero_con_salida_positiva_falla():
    with tempfile.TemporaryDirectory() as d:
        x = os.path.join(d, "t.xlsx")
        libro(x, FILAS)
        salidas(d, 123.4, 12.0, 40.0)
        err = P.verificar_excel(x)
        m = [e for e in err if "Net 0 pero" in e]
        assert any("framed_openings_out.json" in e and "framed openings MEP" in e for e in m), err
        assert any("familias_out.json" in e and "sill track" in e for e in m), err
        assert any("accesorios_out.json" in e and "slab pins" in e for e in m), err


def test_salida_en_cero_o_renglon_con_valor_no_falla():
    with tempfile.TemporaryDirectory() as d:
        x = os.path.join(d, "t.xlsx")
        libro(x, FILAS)
        salidas(d, 0.0, 0.0, 0.0)
        assert not [e for e in P.verificar_excel(x) if "Net 0 pero" in e]
        libro(x, [("Framed openings MEP", 123.4, FILAS[0][2])])
        salidas(d, 123.4, 0.0)
        assert not [e for e in P.verificar_excel(x) if "Net 0 pero" in e]


def test_salidas_en_otra_carpeta():
    with tempfile.TemporaryDirectory() as d:
        os.makedirs(os.path.join(d, "xl"))
        os.makedirs(os.path.join(d, "out"))
        x = os.path.join(d, "xl", "t.xlsx")
        libro(x, FILAS[:1])
        salidas(os.path.join(d, "out"), 100.0, 0.0)
        assert not [e for e in P.verificar_excel(x) if "Net 0 pero" in e]          # no la encuentra: no inventa
        assert [e for e in P.verificar_excel(x, [os.path.join(d, "out")]) if "Net 0 pero" in e]


def libro_yields(ruta, net_screws, j_screws="=Inputs!$B$2*1000", net_tape=370.0):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Material Summary"
    ws.append(["MATERIAL SUMMARY"])
    ws.append(["nota"])
    ws.append(["Material", "Unit", "Net", "Waste %", "Waste", "Purchase", "Selling unit", "Pieces to buy",
               "Where it comes from", "Yield check (live formula)"])
    ws.append(["Drywall screws 1-1/4\"", "PCS", net_screws, None, None, None, "box", None, "yield", j_screws])
    ws.append(["Joint tape, paper", "LF", net_tape, None, None, None, "roll", None, "yield", "=Inputs!$B$3*1000"])
    pt = wb.create_sheet("Partition Takeoff")
    for _ in range(4):
        pt.append([None])
    pt.append(["TYPE X1", "row", 1])
    inp = wb.create_sheet("Inputs")
    inp.append(["Waste metal", 0.15])
    inp.append(["Drywall screws per SF of board", 1.2])
    inp.append(["Joint tape (LF per SF)", 0.37])
    wb.save(ruta)


def test_r70_8_yield_desde_inputs():
    with tempfile.TemporaryDirectory() as d:
        x = os.path.join(d, "t.xlsx")
        libro_yields(x, 1200.0)                                   # = 1.2 x 1000: pasa
        assert not [e for e in P.verificar_excel(x) if "R70.8" in e]
        libro_yields(x, "=Inputs!$B$2*1000")                      # fórmula viva: pasa
        assert not [e for e in P.verificar_excel(x) if "R70.8" in e]
        libro_yields(x, 1875.0)                                   # yield escrito a mano (no el de Inputs): falla
        e = [e for e in P.verificar_excel(x) if "R70.8" in e]
        assert e and "yield must come from Inputs" in e[0] and "Drywall screws" in e[0], e
        libro_yields(x, 1200.0, j_screws=None)                    # sin celda de Inputs citada: falla
        assert [e for e in P.verificar_excel(x) if "R70.8" in e]
        libro_yields(x, 1200.0, net_tape=380.0)                   # tape fuera de 0.5 %: falla  [ciego-ok]
        assert [e for e in P.verificar_excel(x) if "R70.8" in e and "tape" in e]


if __name__ == "__main__":
    fallas = 0
    for k, f in sorted(globals().items()):
        if k.startswith("test_") and callable(f):
            try:
                f()
                print("OK  ", k)
            except Exception as e:  # noqa: BLE001
                fallas += 1
                print("FALLA", k, type(e).__name__, str(e)[:300])
    print(f"{'TODOS OK' if not fallas else str(fallas) + ' FALLA(S)'}")
    sys.exit(1 if fallas else 0)
