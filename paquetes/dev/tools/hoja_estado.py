# -*- coding: utf-8 -*-
"""hoja_estado.py -- genera BOA_ESTADO.md DESDE EL LIBRO, no a mano.

Por que existe: el 28-sep BOA_ESTADO.md seguia diciendo "73 elementos" y "W.R. 100.80 SF"
un dia despues de que el libro dijera 77 y <n>. Un documento de estado escrito a mano
se queda viejo por la misma razon que el Excel: nadie lo vuelve a escribir cuando cambia
un numero. La regla de vigencia vale tambien para el resumen que el usuario lee.

Aqui la tabla de compra se LEE del .xlsx recien generado. Si el libro cambia, este
documento cambia solo.
"""
import os as _os_rutas  # rutas de trabajo: TAKEOFF_WORK / TAKEOFF_OUT (antes fijas en /home/claude/boa)
_TW = _os_rutas.environ.get("TAKEOFF_WORK", "/home/claude/boa")
_TO = _os_rutas.environ.get("TAKEOFF_OUT", "/mnt/user-data/outputs")
import os, sys, subprocess, datetime
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import verif_vigencia as VV
import glosario as G

B = f'{_TW}'
DEST = [f'{_TO}/entrega/BOA_ESTADO.md']

CAB = """# BOA Broomfield - ESTADO DEL TAKE-OFF - %(fecha)s

> Este documento se GENERA del Excel (`tools/hoja_estado.py`). No se escribe a mano: un
> resumen escrito a mano se queda viejo igual que el libro.

## %(sellos)s

%(bloques)s

## Resumen de compra - leido del libro, no tecleado

%(compra)s

Waste: **metal 15%%, board 6%%**, sobre el neto, por tipo, en su propia celda de
Inputs. El FRT wood blocking va en 0%% por decision declarada en Pending.

## Vigencia del entregable

%(vigencia)s

El `.xlsx` lleva un manifiesto con el `sha256` de sus %(n_insumos)d insumos -- los cierres
sellados **y** el codigo que los convierte en libro. `tools/entregar.py` es la unica
puerta de salida: sellos -> regenera -> exige vigencia -> y solo entonces copia. No se
puede entregar un libro viejo porque no hay otra ruta para sacarlo.
"""


def _titulo(sellos):
    """El renglon TOTAL trae 'TOTAL  <n elementos>  <n> sellos OK / <n> rotos'. Se parsea
    con regex por CONTENIDO -- partirlo por espacios repetia el conteo."""
    import re
    if not sellos:
        return 'ESTADO NO VERIFICADO'
    m = re.search(r'(\d+)\s+(\d+)\s+sellos OK\s*/\s*(\d+)\s+rotos', sellos[-1])
    if not m:
        return 'ESTADO NO VERIFICADO'
    n, ok, rotos = m.groups()
    estado = 'TERMINADO' if rotos == '0' else '*** %s SELLO(S) ROTO(S)' % rotos
    return '%s. %s elementos, %s sellos, %s rotos.' % (estado, n, ok, rotos)


def tabla_compra(xlsx):
    ws = openpyxl.load_workbook(xlsx, data_only=True)['Material Summary']
    out = ['| material | unidad | neto | compra | comprar |', '|---|---|---|---|---|']
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        if not r or not r[0]:
            continue
        n = str(r[0]).strip()
        if len(r) > 5 and isinstance(r[2], (int, float)) and isinstance(r[5], (int, float)):
            pz = r[7] if len(r) > 7 and isinstance(r[7], (int, float)) else ''
            un = r[6] if len(r) > 6 and r[6] else ''
            out.append('| %s | %s | %.2f | **%.2f** | %s |'
                       % (n, r[1], r[2], r[5], ('%s %s' % (pz, un)).strip()))
        elif n.isupper() and len(n) < 40:
            out.append('| **%s** | | | | |' % n)
    return '\n'.join(out)


def main(xlsx=VV.LIBRO):
    r = subprocess.run(['python3', B + '/tools/resumen_total.py'],
                       capture_output=True, text=True)
    lin = [l.rstrip() for l in (r.stdout or '').splitlines()]
    sellos = [l for l in lin if 'sellos OK' in l]
    bloques = [l for l in lin if ' OK / ' in l and 'sellos OK' not in l]
    v = VV.verificar(xlsx)
    txt = CAB % {
        'fecha': datetime.date.today().strftime('%d-%b-%Y'),
        'sellos': _titulo(sellos),
        'bloques': '```\n' + '\n'.join(bloques) + '\n```' if bloques else '',
        'compra': tabla_compra(xlsx),
        'vigencia': ('**OK** - el Excel refleja el estado sellado de hoy, probado por hash.'
                     if not v else '**DESACTUALIZADO:**\n\n- ' + '\n- '.join(v)),
        'n_insumos': len(VV.fuentes()),
    }
    # El nombre de material va en INGLES. Se pasa por el glosario antes de escribir, no
    # despues: asi una palabra en espanol que se cuele en la plantilla se corrige sola en
    # vez de quedarse esperando a que alguien la vea.
    txt = G.traducir(txt)
    for d in DEST:
        os.makedirs(os.path.dirname(d), exist_ok=True)
        open(d, 'w', encoding='utf-8').write(txt)
    print('BOA_ESTADO.md regenerado desde el libro ->', DEST[0])
    return DEST[0]


if __name__ == '__main__':
    main()
