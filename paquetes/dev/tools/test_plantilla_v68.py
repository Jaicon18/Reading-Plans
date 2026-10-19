#!/usr/bin/env python3
"""Tests sintéticos v68 de plantilla_excel.py (arbitraje #16, E24: números inventados): compra y presentación.
1 FRT plywood por material / espesor · 2 unidades de compra en celdas de Inputs · 3 desperdicio una vez ·
4 yields con celda real de Inputs y fórmula viva · 5 sealant exterior aparte · 6 Excel íntegro en inglés ·
7 Pending Items sin duplicados / cerrados y con delta · 8 NET de familia solo en su unidad (hanger wire PCS + LF)."""
import copy
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl  # noqa: E402
import plantilla_excel as P
# v72b: estos tests revisan el FORMATO del libro con familias sintéticas; armar() exige reglas.py compuerta OK
# sobre una carpeta de trabajo real. El libro de aquí lleva un sello de test, que bitacora / reglas --excel rechazan.
SELLO_TEST = {"huella": {"version": "test"}}  # noqa: E402
import test_plantilla_v67 as T7  # noqa: E402

OK = []
FRT = "FRT plywood (SF)"


def chk(n, c, m=""):
    if not c:
        raise AssertionError(f"{n}: {m}")
    OK.append(n)


def raises(n, f, txt):
    try:
        f()
    except ValueError as e:
        chk(n, txt in str(e), str(e))
        return
    chk(n, False, "no levantó ValueError")


def fo_base():
    fo = T7.familias_out()
    fo = copy.deepcopy(fo)
    fo["familias"][FRT] = 100.0 + 40.0 + 25.0 + 10.0
    fo["familias"]["acoustical sealant"] = 90.0
    fo["familias"]["hanger wire"] = 33.0
    fo["hanger_pcs"] = {"total_pcs": 11}
    fo["filas"] += [
        {"familia": FRT, "ref": "roof side 5/8\" FRT plywood bonded to polyiso", "source": "1/X9.21",
         "operacion": "strip 10 x 10 = 100 SF"},
        {"familia": FRT, "ref": "FB-1", "source": "2/X9.26", "operacion": "fascia_box sheathing 20 x 2 = 40 SF "
                                                                       "(5/8\" ext. grade FRT plywood, 2/X9.26)"},
        {"familia": FRT, "ref": "electrical room 3/4\" FRT plywood backboard", "source": "X11.03",
         "operacion": "strip 5 x 5 = 25 SF"},
        {"familia": FRT, "ref": "pier 3/4\" FRT plywood", "valor": 10.0, "source": "1/X9.25", "operacion": "x"},
        {"familia": "acoustical sealant", "ref": "sealant at base of gypsum (exterior)", "source": "5/X9.26",
         "operacion": "12.5 LF (not acoustical)"}]
    return fo


def takeoff():
    d = T7.takeoff()
    d["materials"] += [
        {"material": "FRT plywood (nailbase / fascia / strips)", "unit": "SF", "net": 1.0, "metal": False,
         "selling_unit": "4x8 sheet", "source": f"familias_out.json '{FRT}'; 1/X9.21"},
        {"material": "Acoustical sealant", "unit": "LF", "net": 1.0, "metal": False, "selling_unit": "LF",
         "source": "familias_out.json 'acoustical sealant'; 1/X12.01"},
        {"material": "Hanger wire 12 ga", "unit": "LF", "net": 1.0, "metal": True, "selling_unit": "roll",
         "source": "familias_out.json 'hanger wire'; 3/X9.33"},
        {"material": "Bracing 1-1/2\" C.R. channel @ 48\" O.C.", "unit": "LF", "net": 100.0, "metal": True,
         "selling_unit": "piece 16 ft", "source": "F/1/X5.20"},
        {"material": "Furring channel 7/8\" hat", "unit": "LF", "net": 100.0, "metal": True, "selling_unit": "12 ft",
         "source": "X4.01"},
        {"material": "Metal trim - corner bead", "unit": "LF", "net": 100.0, "metal": True, "selling_unit": "piece",
         "source": "8/X12.01"},
        {"material": "Bottom track 3-5/8\"", "unit": "LF", "net": 100.0, "metal": True, "selling_unit": "piece 10 ft",
         "source": "7/X5.20"},
        # v71 R71.8: track por tamaño de stud (los tamaños que trae el Partition Takeoff del fixture)
        *[{"material": f"{c} {t}", "unit": "LF", "net": 10.0, "metal": True, "selling_unit": "piece 10 ft",
           "source": "7/X5.20"} for c, t in [("Bottom track", '1-5/8"'), ("Bottom track", '2-1/2"'),
                                             ("Bottom track", '6"'), ("Top / deflection track", '1-5/8"'),
                                             ("Top / deflection track", '2-1/2"'), ("Top / deflection track", '3-5/8"'),
                                             ("Top / deflection track", '6"')]]]
    return d


def armado(d=None, fo=None):
    d = P.fusionar(d or takeoff(), fo or fo_base())
    t = tempfile.mkdtemp()
    x = os.path.join(t, "t.xlsx")
    P._escribir(d, x, SELLO_TEST)
    return d, x


def filas_ms(x, data_only=True):
    ws = openpyxl.load_workbook(x, data_only=data_only)["Material Summary"]
    return {r[0]: r for r in ws.iter_rows(min_row=4, values_only=True) if r[0]}


def test_frt_por_material():
    d, x = armado()
    n = [m["material"] for m in d["materials"] if "plywood" in m["material"].lower()]
    chk("FRT: renglón mezclado reemplazado", "FRT plywood (nailbase / fascia / strips)" not in n, n)
    chk("FRT: nailbase 5/8\"", '5/8" FRT plywood nailbase bonded to polyiso (exterior sheathing)' in n, n)
    chk("FRT: canopy 5/8\"", '5/8" FRT plywood sheathing (canopy, fascia and soffits)' in n, n)
    chk("FRT: backboard 3/4\"", '3/4" FRT plywood backboards' in n, n)
    chk("FRT: piers 3/4\"", '3/4" FRT plywood at piers' in n, n)
    chk("FRT: suma = familia", abs(sum(m["net"] for m in d["materials"] if "plywood" in m["material"].lower())
                                   - 175.0) < 1e-6)
    chk("FRT: Excel pasa verificar", P.verificar_excel(x) == [], P.verificar_excel(x))
    fo = fo_base(); fo["filas"].append({"familia": FRT, "ref": "FRT plywood strip", "valor": 5.0, "source": "9/X1"})
    fo["familias"][FRT] += 5.0
    raises("FRT: sin espesor = no se arma", lambda: P.fusionar(takeoff(), fo), "sin espesor")
    fo = fo_base(); fo["familias"][FRT] += 50.0
    raises("FRT: desglose que no suma", lambda: P.fusionar(takeoff(), fo), "suma")
    fo = fo_base(); fo["frt_plywood_by_material"] = {'5/8" FRT plywood substrate': 175.0}
    d2 = P.fusionar(takeoff(), fo)
    chk("FRT: frt_plywood_by_material manda", [m["net"] for m in d2["materials"] if "plywood" in m["material"]] == [175.0])
    fo["frt_plywood_by_material"] = {"FRT plywood substrate": 175.0}
    raises("FRT: dict sin espesor", lambda: P.fusionar(takeoff(), fo), "sin espesor")
    chk("validar: plywood sin espesor", any("sin espesor" in e for e in P.validar(dict(
        T7.takeoff(), materials=[{"material": "FRT plywood", "unit": "SF", "net": 1, "metal": False, "source": "s"}]))))
    chk("plywood_err: mezcla", P.plywood_err('5/8" FRT plywood (nailbase / fascia / soffit)') != [])
    chk("plywood_err: canopy column wraps = un material", P.plywood_err('5/8" FRT plywood at canopy column wraps') == [])


def test_compra_inputs_y_desperdicio_una_vez():
    d, x = armado()
    f = filas_ms(x, data_only=False)
    v = filas_ms(x)
    chk("C.R. channel / Inputs B16", f["Bracing 1-1/2\" C.R. channel @ 48\" O.C."][7] == "=ROUNDUP(F{}/Inputs!$B$16,0)".format(
        list(f).index("Bracing 1-1/2\" C.R. channel @ 48\" O.C.") + 4))
    chk("C.R. channel piezas", v["Bracing 1-1/2\" C.R. channel @ 48\" O.C."][7] == 8, v["Bracing 1-1/2\" C.R. channel @ 48\" O.C."])   # 115/16
    chk("hat channel / B17", "Inputs!$B$17" in f["Furring channel 7/8\" hat"][7] and v["Furring channel 7/8\" hat"][7] == 10)
    chk("trim stick / B18", "Inputs!$B$18" in f["Metal trim - corner bead"][7] and v["Metal trim - corner bead"][7] == 12)
    chk("track run / B11", "Inputs!$B$11" in f["Bottom track 3-5/8\""][7] and v["Bottom track 3-5/8\""][7] == 12)
    chk("sheet / B9", any("Inputs!$B$9" in str(r[7]) for k, r in f.items() if "plywood" in k))
    s = [k for k in f if k.startswith("Acoustical sealant")][0]
    chk("sealant en tubes / B15", "Inputs!$B$15" in f[s][7] and "tube" in f[s][6], f[s])
    chk("hanger wire LF continuous coil", "coil" in str(f["Hanger wire 12 ga"][6]), f["Hanger wire 12 ga"])
    # desperdicio una vez: stud de 100 LF, 15 % -> ROUNDUP(115 / 20) = 6 (antes 1 stick por corte de 13 ft = 9)
    stud = {"material": "6\" metal stud 600S137-33", "unit": "LF", "net": 100.0, "metal": True, "stock": 20,
            "cut_ft": 13, "selling_unit": "stick 20 ft", "source": "1/X12.01"}
    chk("stud: ROUNDUP(Purchase / Inputs B10)", P.formula_piezas(stud, 9) == "=ROUNDUP(F9/Inputs!$B$10,0)",
        P.formula_piezas(stud, 9))
    d2 = dict(T7.takeoff(), materials=[stud])
    t = tempfile.mkdtemp(); x2 = os.path.join(t, "s.xlsx"); P._escribir(d2, x2, SELLO_TEST)
    chk("stud: 6 sticks", filas_ms(x2)[stud["material"]][7] == 6, filas_ms(x2)[stud["material"]])
    wb = openpyxl.load_workbook(x2); ws = wb["Material Summary"]
    ws["H4"] = "=ROUNDUP(ROUNDUP(F4/13,0)/MAX(1,INT(20/13)),0)"; y = os.path.join(t, "m.xlsx"); wb.save(y)
    chk("verificar rechaza desperdicio doble", any("dos veces" in e for e in P.verificar_excel(y)), P.verificar_excel(y))


def test_yields_inputs_reales():
    d, x = armado()
    wb = openpyxl.load_workbook(x, data_only=True)
    inp = wb["Inputs"]
    chk("Inputs B20 drywall screws", inp["A20"].value.startswith("Drywall screws per SF") and inp["B20"].value == 1.2)
    chk("Inputs B23 joint compound = 0.072", inp["A23"].value.startswith("Joint compound") and inp["B23"].value == 0.072)
    chk("Inputs B15 sealant", inp["A15"].value.startswith("Sealant coverage") and inp["B15"].value == 25)
    chk("waste en B5 / B6", inp["B5"].value == 0.15 and inp["B6"].value == 0.06)
    v = filas_ms(x)
    f = filas_ms(x, data_only=False)
    ds = 'Drywall screws 1-1/4" bugle head'
    chk("yield: fórmula viva", f[ds][9].startswith("=Inputs!$B$20*"), f[ds])
    chk("yield: fórmula = NET", abs(v[ds][9] - v[ds][2]) < 0.6, v[ds])
    chk("yield: source cita la celda real", "Inputs!B20" in v[ds][8] and "criterio" not in v[ds][8], v[ds][8])
    chk("framing screws B21", f["Framing screws, pan head self-drilling"][9].startswith("=Inputs!$B$21*"))
    sp = [k for k in f if "slab pins" in k][0]
    chk("slab pins = LF / B22", f[sp][9].endswith("/Inputs!$B$22") and abs(v[sp][9] - v[sp][2]) < 0.6, f[sp])
    # cita a una celda de Inputs que no existe -> rechazo
    wb = openpyxl.load_workbook(x); wb["Material Summary"]["I4"] = "yield per Inputs B31"
    y = os.path.join(os.path.dirname(x), "c.xlsx"); wb.save(y)
    chk("verificar: cita Inputs inexistente", any("B31" in e for e in P.verificar_excel(y)), P.verificar_excel(y))
    fo = fo_base()
    for r in fo["filas"]:
        if r.get("familia") == "joint compound":
            r["operacion"] = r["operacion"].replace("0.072 x", "0.138 x")
    raises("yield que no reproduce el NET", lambda: P.fusionar(takeoff(), fo), "da")


def test_sealant_exterior_aparte():
    d, x = armado()
    m = {r["material"]: r for r in d["materials"]}
    chk("sealant exterior en su renglón", m[P.SEAL_EXT_NOMBRE]["net"] == 12.5, m.get(P.SEAL_EXT_NOMBRE))
    chk("acoustical sin el exterior", m["Acoustical sealant"]["net"] == 77.5, m["Acoustical sealant"])
    d3 = takeoff()
    for m3 in d3["materials"]:
        if m3["material"] == "Acoustical sealant":
            m3["material"] = "Acoustical sealant (partitions) and sealant at base exterior"
    m3 = {r["material"]: r for r in P.fusionar(d3, fo_base())["materials"]}
    chk("renglón mezclado se separa y se renombra", m3.get("Acoustical sealant", {}).get("net") == 77.5
        and m3[P.SEAL_EXT_NOMBRE]["net"] == 12.5, list(m3))
    wb = openpyxl.load_workbook(x); ws = wb["Material Summary"]
    ws.append(["Acoustical sealant and sealant at base of gypsum exterior", "LF", 1, 0, 0, 1, "tube", "=ROUNDUP(F1/Inputs!$B$15,0)"])
    y = os.path.join(os.path.dirname(x), "s.xlsx"); wb.save(y)
    chk("verificar: sealant mezclado", any("mezcla sealant" in e for e in P.verificar_excel(y)))


def test_ingles():
    d, x = armado()
    wb = openpyxl.load_workbook(x)
    chk("hojas en inglés", P.HOJA_REVISAR in wb.sheetnames and P.HOJA_ALCANCES in wb.sheetnames
        and "PARA REVISAR EN LA MANANA" not in wb.sheetnames, wb.sheetnames)
    chk("verificar OK en inglés", P.verificar_excel(x) == [], P.verificar_excel(x))
    chk("en(): textos de tools", P.en("ancho del jamb return no acotado") == "jamb return width not dimensioned")
    chk("en(): criterio", P.en("criterio estimador (X)") == "estimator criterion (X)")
    chk("en(): decisión", P.en("decisión #7") == "user decision #7" and P.en("user decision #2") == "user decision #2")
    ws = wb[P.HOJA_REVISAR]; ws.append([9, "Altura de la cara oeste sin cota", "1/X1", "-"])
    ws = wb["Material Summary"]; ws["I4"] = "criterio"
    y = os.path.join(os.path.dirname(x), "es.xlsx"); wb.save(y)
    e = P.verificar_excel(y)
    chk("verificar: frase en español", any(P.HOJA_REVISAR in x_ and "español" in x_ for x_ in e), e)
    chk("verificar: palabra suelta", any("Material Summary!I4" in x_ for x_ in e), e)


def test_rfis_sin_duplicados():
    d = takeoff()
    d["rfis"] = [{"id": "RFI-01", "pregunta": "Wall height not dimensioned", "delta": "+10 SF"},
                 {"id": "RFI-02", "pregunta": "Wall height not dimensioned", "delta": "+10 SF"},
                 {"id": "RFI-03", "pregunta": "Old question", "delta": "0", "estado": "CERRADO"},
                 {"id": "RFI-ACC1", "pregunta": "hand copy of a tool RFI", "delta": "1"}]
    d2 = P.fusionar(d, fo_base())
    ids = [r["id"] for r in d2["rfis"]]
    chk("RFI duplicado por tema fuera", "RFI-02" not in ids and "RFI-01" in ids, ids)
    chk("RFI cerrado fuera", "RFI-03" not in ids, ids)
    chk("RFI-ACC a mano reemplazado por el de la tool", "RFI-ACC1" not in ids
        and any(i.startswith("RFI-ACC-") for i in ids), ids)
    chk("ids estables por familia", len(ids) == len(set(ids)), ids)
    chk("delta dict legible", all(not str(r["delta"]).startswith("{") for r in d2["rfis"]), d2["rfis"])
    chk("validar: RFI sin delta", any("sin delta" in e for e in P.validar(dict(
        T7.takeoff(), rfis=[{"id": "R", "pregunta": "q", "delta": ""}]))))
    chk("fmt_delta", P.fmt_delta({"delta_16in": -1.5, "delta_24in": -3}, "pcs") == "16in: -1.5 pcs; 24in: -3 pcs")


def test_unidad_de_familia():
    d = takeoff()
    d["materials"].append({"material": "Hanger wire pieces", "unit": "PCS", "net": 1, "metal": True,
                           "source": "familias_out.json 'hanger wire'; 3/X9.33"})
    raises("NET LF en renglón PCS = no se arma", lambda: P.fusionar(d, fo_base()), "misma unidad")
    d2 = P.fusionar(takeoff(), fo_base())
    hw = [m for m in d2["materials"] if "hanger" in m["material"].lower()]
    chk("hanger wire LF + PCS", {m["unit"] for m in hw} == {"LF", "PCS"}, hw)
    chk("hanger PCS = conteo", [m["net"] for m in hw if m["unit"] == "PCS"] == [11.0], hw)
    chk("hanger LF = familia", [m["net"] for m in hw if m["unit"] == "LF"] == [33.0], hw)


def test_yield_operacion_en_ingles():
    """Integración v68: accesorios.py escribe la operación de framing screws en inglés ("stud pieces")."""
    fo = {"filas": [{"familia": "framing screws", "operacion": "4 x 60 stud pieces [W1: 60]"}]}
    chk("framing screws 'stud pieces'", P.yield_de(fo, "framing screws", 240.0) == (4.0, 60.0))
    fo = {"filas": [{"familia": "framing screws", "operacion": "4 x 60 piezas"}]}
    chk("framing screws 'piezas' (v67)", P.yield_de(fo, "framing screws", 240.0) == (4.0, 60.0))


def test_partition_pendiente_no_se_entrega():
    """C-P5 v68: un renglón de Partition Takeoff PENDING (partition_types sin dato) bloquea el Excel."""
    d, x = armado()
    chk("Excel base sin PENDING", P.verificar_excel(x) == [], P.verificar_excel(x))
    wb = openpyxl.load_workbook(x)
    ws = wb["Partition Takeoff"]
    fila = next(r for r in range(5, ws.max_row + 1) if str(ws.cell(r, 1).value or "").strip()
                and not str(ws.cell(r, 1).value).startswith("TYPE "))
    ws.cell(fila, 11).value = "PENDING: runs without ori / axis"
    y = os.path.join(os.path.dirname(x), "pend.xlsx"); wb.save(y)
    chk("PENDING rechazado", any("PENDING" in e for e in P.verificar_excel(y)), P.verificar_excel(y))


if __name__ == "__main__":
    for n, f in list(globals().items()):
        if n.startswith("test_"):
            f()
    print(f"test_plantilla_v68: {len(OK)} checks OK")
