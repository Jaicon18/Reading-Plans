# -*- coding: utf-8 -*-
"""hoja_rfi.py -- genera la carta de RFIs LEYENDO la hoja Pending Items del libro.

No se teclea en paralelo: si un pendiente se resuelve y sale del Excel, sale solo de la
carta. Si aparece uno nuevo, entra solo. Un documento que se escribe a mano al lado de su
fuente siempre termina diciendo otra cosa.

Entra a la carta todo pendiente cuya columna "Ask" apunte a alguien FUERA de la oficina
-- arquitecto, plomeria, campo. Los que dependen de una decision del estimador NO son RFI
y se listan aparte, para que se vea que no se olvidaron.
"""
import os as _os_rutas  # rutas de trabajo: TAKEOFF_WORK / TAKEOFF_OUT (antes fijas en /home/claude/boa)
_TW = _os_rutas.environ.get("TAKEOFF_WORK", "/home/claude/boa")
_TO = _os_rutas.environ.get("TAKEOFF_OUT", "/mnt/user-data/outputs")
import os, sys, datetime, re

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import glosario as G
from openpyxl import load_workbook

LIBRO = f'{_TW}/out/BOA_Broomfield_Takeoff.xlsx'
DEST = f'{_TO}/entrega/BOA_RFI_al_arquitecto.md'
FUERA = ('architect', 'plumbing', 'field')       # a quien se le pregunta


def leer_pendientes(xlsx=LIBRO):
    ws = load_workbook(xlsx, data_only=True)['Pending Items']
    out = []
    for r in ws.iter_rows(values_only=True):
        if not r or not r[0] or not str(r[0]).strip().isdigit():
            continue
        n, item, area, det, quien, imp = (list(r) + [None] * 6)[:6]
        out.append({'n': int(n), 'item': item, 'area': area, 'det': det,
                    'quien': quien or '', 'imp': imp or ''})
    return out


ENCABEZADO = """# REQUEST FOR INFORMATION — Metal Framing & Drywall

**Proyecto:** Bank of America · Broomfield, CO — NWC Hwy 7 and Palisade Pkwy · CO1-111
**Set:** Permit / Bid Set, 95 hojas · **Arquitecto:** Nelson Architecture, Inc.
**De:** JC Hunter · **Fecha:** %(fecha)s · **Alcance:** CSI Div. 05 Metals y Div. 09 Finishes

---

Estos %(n)d puntos **no se pueden resolver con el set tal como está**. En cada uno se dice
qué se buscó, qué sí dice el plano y qué falta. Ninguna de estas cantidades se inventó: las
que dependen de la respuesta van marcadas **PENDING** en el take-off, nunca como cero.

Se agradece la respuesta antes de la fecha de entrega de propuestas, porque %(n_precio)d de
estos puntos afectan precio.

---

"""

CIERRE = """
---

# Decisiones que NO son RFI

Estos puntos no dependen del arquitecto sino de un criterio nuestro. Se listan para que
quede constancia de que están identificados y de cómo se están llevando hoy:

%(dec)s

---

# Cómo se está cotizando mientras tanto

Cada cantidad que depende de una de estas respuestas sale marcada **PENDING** en la hoja
`Pending Items` del take-off, con su motivo y a quién se le pidió. **Ninguna se cotizó en
cero y ninguna se estimó.** En cuanto llegue la respuesta, entra al Excel y el resumen de
compra se recalcula solo.
"""


def main(xlsx=LIBRO, dest=DEST):
    P = leer_pendientes(xlsx)
    # Un pendiente ya RESUELTO no es RFI ni decision: se queda en el Excel como registro
    # de que se cerro y por que, pero no se le pregunta nada a nadie.
    cerrado = [p for p in P if 'resolved' in (str(p['item']) + str(p['quien'])).lower()
               or 'closed' in str(p['quien']).lower()]
    abiertos = [p for p in P if p not in cerrado]
    rfi = [p for p in abiertos if any(k in p['quien'].lower() for k in FUERA)]
    dec = [p for p in abiertos if p not in rfi]
    n_precio = sum(1 for p in rfi if 'affect' in p['imp'].lower() or 'price' in p['imp'].lower())

    cuerpo = []
    for i, p in enumerate(rfi, 1):
        cuerpo.append('## RFI %02d — %s\n' % (i, p['item']))
        cuerpo.append('**Área:** %s  ·  **Se le pregunta a:** %s  ·  **Impacto:** %s\n'
                      % (p['area'], p['quien'], p['imp']))
        cuerpo.append('%s\n' % p['det'])
        cuerpo.append('---\n')

    lin = []
    for p in dec:
        lin.append('- **%s** (%s) — %s  *Impacto: %s*'
                   % (p['item'], p['area'], (p['det'] or '')[:220].rstrip() + '…', p['imp']))

    txt = (ENCABEZADO % {'fecha': datetime.date.today().strftime('%d-%b-%Y'),
                         'n': len(rfi), 'n_precio': n_precio}
           + '\n'.join(cuerpo)
           + CIERRE % {'dec': '\n'.join(lin)})
    txt = G.traducir(txt)                    # el material, en ingles
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    open(dest, 'w', encoding='utf-8').write(txt)
    print('carta de RFIs: %d preguntas al exterior, %d decisiones internas, %d ya cerrados'
          ' -> %s' % (len(rfi), len(dec), len(cerrado), dest))
    return dest


if __name__ == '__main__':
    main()
