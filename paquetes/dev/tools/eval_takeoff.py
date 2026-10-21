#!/usr/bin/env python3
"""eval_takeoff.py — calificador (grader) de un take-off contra un take-off de referencia.

Idea (Anthropic, "Demystifying evals for AI agents"): calificar el RESULTADO, no el camino;
grader de código (determinista) para cantidades; crédito parcial; tolerancia para que
120.0 vs 120 no cuente como error; y reportar cada diferencia para leerla una por una.

Compara dos niveles:
  1) Material Summary: cada material, su unidad y su cantidad NET (sin waste).
  2) Partition Takeoff: por partition type, cuántos walls y cuántos LF.

Uso:
  eval_takeoff.py <referencia.xlsx> <candidato.xlsx> [--out reporte.md] [--json resultado.json]
                  [--tol 0.02] [--close 0.10] [--rangos rangos.json]

Calificación de cada renglón:
  IGUAL    diferencia <= tol (2% por defecto)      -> 1.0 punto
  CERCA    diferencia <= close (10% por defecto)   -> 0.5 punto
  MAL      diferencia mayor                        -> 0
  FALTA    está en la referencia y no en el candidato -> 0
  SOBRA    está en el candidato y no en la referencia -> se reporta, no suma ni resta
Con --rangos (15-oct-2026): las familias de CRITERIO de la referencia definitiva tienen un rango
[min, max]; dentro del rango = IGUAL, fuera se mide contra el borde más cercano.
Con modo "excluir" (v77c, U8) la familia está EXCLUIDA: sin renglón = EXCLUIDA OK (1 punto); con renglón = MAL.
Exit code: 0 si la calificación total >= --pass (95% por defecto), 1 si no.
"""
import argparse, difflib, json, re, sys
import openpyxl
import os as _os
sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from familias import *  # noqa: F401,F403  (num, norm_unit, find_sheet, read_materials, FAMILIAS, familia, por_familia)

SIZES = ['6"', '3-5/8"', '2-1/2"', '1-5/8"', '7/8"', '5/8"', '1/2"', '3/8"', '3/4"', '1-1/2"', '1-1/4"']
KEYWORDS = ["stud", "track", "bottom", "top", "deflection", "parapet", "ceiling", "soffit", "canopy",
            "bracing", "channel", "header", "sill", "opening", "reinforcement", "purlin", "furring",
            "carrying", "hanger", "wire", "extra", "gypsum", "board", "moisture", "exterior", "sheathing",
            "frp", "cement", "plywood", "glass", "mat", "jamb", "ballistic", "frt", "level", "trim",
            "metal", "pvc", "fastener", "blocking", "sealant", "access", "panel", "screw", "framing",
            "pin", "compound", "tape", "csr", "infill", "head", "mep", "regular"]


def signature(name):
    """Firma comparable de un material: medidas + palabras clave, sin depender de la redacción."""
    s = str(name).lower().replace("”", '"').replace("″", '"')
    sizes = sorted({z for z in SIZES if z in s})
    words = sorted({k for k in KEYWORDS if re.search(r"\b" + re.escape(k) + r"s?\b", s)})
    return sizes, words


def read_partitions(wb):
    cands = [w for w in wb.worksheets if "partition" in w.title.lower()] + \
            [w for w in wb.worksheets if w.title.strip().lower() == "takeoff"]
    ws = rows = hdr_i = None
    for w in cands:
        rr = list(w.iter_rows(values_only=True))
        hi = next((i for i, r in enumerate(rr)
                   if r and any("length" in str(c or "").lower() for c in r)
                   and any(("wall" in str(c or "").lower() or "partition" in str(c or "").lower()) for c in r)), None)
        if hi is not None:
            ws, rows, hdr_i = w, rr, hi
            break
    if ws is None:
        return None
    wb._hoja_partitions = ws.title
    hdr = [str(c or "").lower() for c in rows[hdr_i]]
    cl = next(i for i, h in enumerate(hdr) if "length" in h)
    cs = next((i for i, h in enumerate(hdr) if "partition" in h), None)
    types, cur = {}, None
    for r in rows[hdr_i + 1:]:
        first = str(r[0] or "").strip()
        scope = str(r[cs] or "").strip() if cs is not None and cs != 0 else ""
        if cs == 0 and cur and first and not first.upper().startswith(("TYPE", "TOTAL", "SUBTOTAL")):
            lf0 = num(r[cl]) if cl < len(r) else None
            if lf0 is not None:
                types.setdefault(cur, {"walls": 0, "lf": 0.0})
                types[cur]["walls"] += 1
                types[cur]["lf"] += lf0
            continue
        if first.upper().startswith("TYPE ") or first.upper().startswith("TYPE/"):
            cur = (first.split(":", 1)[1] if ":" in first else first[5:]).strip().upper().split()[0]
            types.setdefault(cur, {"walls": 0, "lf": 0.0})
            continue
        lf = num(r[cl]) if cl < len(r) else None
        if scope:                                   # formato "Wall # | A1 - 3-5/8\" MST ..."
            if scope.upper().startswith("TOTAL"):
                continue
            t = scope.split(" - ")[0].split()[0].upper()
        else:
            if lf is None or not first or first.upper().startswith("TOTAL"):
                continue
            t = cur or first.split()[0].upper()
            if first.split()[0].upper() != t and "#" in first:      # renglones fuera de un bloque TYPE
                t = first.split()[0].upper()
        if lf is None:
            continue
        types.setdefault(t, {"walls": 0, "lf": 0.0})
        types[t]["walls"] += 1
        types[t]["lf"] += lf
    return types


# Familias de material: comparan el RESULTADO aunque cada take-off reparta los renglones distinto
# (p.ej. uno separa stud interior y exterior y el otro los junta). Orden = prioridad.
def otra_hoja(wb, t):
    """Nombre de la hoja donde aparece el type como renglón con cantidad, fuera de la hoja de partitions."""
    pat = re.compile(r"^\s*" + re.escape(t) + r"(\s|-|$)", re.I)
    raiz = t.split("-")[0] if "-" in t and len(t.split("-")[0]) >= 5 else None   # BULKHEAD-G -> "bulkhead"
    for ws in wb.worksheets:
        tl = ws.title.lower()
        if ws.title == getattr(wb, "_hoja_partitions", None) or "revisar" in tl or "rfi" in tl:
            continue
        for row in ws.iter_rows(values_only=True):
            if any(isinstance(c, str) and (pat.match(c) or (raiz and raiz.lower() in c.lower())) for c in row) and \
               any(isinstance(c, (int, float)) and c > 0 for c in row):
                return ws.title
    return None


def grade(ref, cand, tol, close):
    if ref == 0 and cand == 0:
        return "IGUAL", 0.0, 1.0
    d = abs(cand - ref) / max(abs(ref), 1e-9)
    if d <= tol:
        return "IGUAL", d, 1.0
    if d <= close:
        return "CERCA", d, 0.5
    return "MAL", d, 0.0


def grade_rango(lo, hi, cand, tol, close):
    """Familia de CRITERIO: cualquier valor dentro de [lo, hi] es IGUAL; fuera, se mide contra el borde."""
    if lo * (1 - tol) <= cand <= hi * (1 + tol):
        return "IGUAL (rango)", 0.0, 1.0
    borde = lo if cand < lo else hi
    d = abs(cand - borde) / max(abs(borde), 1e-9)
    return ("CERCA (rango)", d, 0.5) if d <= close else ("MAL (rango)", d, 0.0)


def match_materials(ref, cand):
    """Empareja cada material de la referencia con el mejor del candidato (misma unidad, misma firma)."""
    used, pairs = set(), []
    for i, r in enumerate(ref):
        rs = signature(r["name"])
        best, best_score = None, 0.0
        for j, c in enumerate(cand):
            if j in used or c["unit"] != r["unit"]:
                continue
            cs = signature(c["name"])
            if rs[0] != cs[0]:                        # las medidas deben coincidir exactamente
                continue
            inter = len(set(rs[1]) & set(cs[1]))
            union = len(set(rs[1]) | set(cs[1])) or 1
            score = inter / union * 0.7 + difflib.SequenceMatcher(
                None, r["name"].lower(), c["name"].lower()).ratio() * 0.3
            if score > best_score:
                best, best_score = j, score
        if best is not None and best_score >= 0.45:
            used.add(best)
            pairs.append((i, best, best_score))
        else:
            pairs.append((i, None, 0.0))
    extra = [j for j in range(len(cand)) if j not in used]
    return pairs, extra


def evaluate(ref_path, cand_path, tol=0.02, close=0.10, rangos=None):
    """rangos: {familia: {"min": n, "max": n}} de la referencia definitiva; esas familias se califican por rango."""
    rangos = rangos or {}
    rwb = openpyxl.load_workbook(ref_path, data_only=True)
    cwb = openpyxl.load_workbook(cand_path, data_only=True)
    res = {"referencia": ref_path, "candidato": cand_path, "tol": tol, "close": close,
           "materiales": [], "materiales_sobran": [], "types": [], "types_sobran": [], "errores": []}

    rm, cm = read_materials(rwb), read_materials(cwb)
    if rm is None:
        res["errores"].append("La referencia no tiene hoja Material Summary legible")
    if cm is None:
        res["errores"].append("El candidato no tiene hoja Material Summary legible")
    if rm and cm is not None:
        pairs, extra = match_materials(rm, cm)
        for i, j, sc in pairs:
            r = rm[i]
            if j is None:
                res["materiales"].append({"material": r["name"], "unidad": r["unit"], "ref": r["net"],
                                          "cand": None, "estado": "FALTA", "dif": None, "puntos": 0.0})
                continue
            c = cm[j]
            est, d, pts = grade(r["net"], c["net"], tol, close)
            res["materiales"].append({"material": r["name"], "cand_nombre": c["name"], "unidad": r["unit"],
                                      "ref": r["net"], "cand": c["net"], "estado": est, "dif": d,
                                      "puntos": pts, "emparejado": round(sc, 2)})
        res["materiales_sobran"] = [{"material": cm[j]["name"], "unidad": cm[j]["unit"], "cand": cm[j]["net"]}
                                    for j in extra]
        rf, cf = por_familia(rm), por_familia(cm)
        res["familias"] = []
        for f in [x[0] for x in FAMILIAS]:
            rg = rangos.get(f) or {}
            if str(rg.get("modo") or "") == "excluir":           # v77c (U8): familia EXCLUIDA en todos los proyectos
                c_ = cf.get(f)
                malo = c_ is not None and abs(c_["net"]) > 1e-9
                res["familias"].append({"familia": f, "ref": 0.0, "cand": c_["net"] if c_ else None,
                                        "estado": "MAL (excluida, trae renglón)" if malo else "EXCLUIDA OK",
                                        "dif": None, "puntos": 0.0 if malo else 1.0,
                                        "cand_renglones": c_["renglones"] if c_ else []})
                continue
            if f not in rf and f not in cf:
                continue
            r_, c_ = rf.get(f), cf.get(f)
            if r_ is None:
                res["familias"].append({"familia": f, "ref": None, "cand": c_["net"], "estado": "SOBRA",
                                        "dif": None, "puntos": None, "cand_renglones": c_["renglones"]})
                continue
            if c_ is None:
                res["familias"].append({"familia": f, "ref": r_["net"], "cand": None, "estado": "FALTA",
                                        "dif": None, "puntos": 0.0, "ref_renglones": r_["renglones"]})
                continue
            rg = rangos.get(f) or {}
            if str(rg.get("modo") or "") == "consistencia":     # v73: no califica contra un número viejo de la ref
                res["familias"].append({"familia": f, "ref": r_["net"], "cand": c_["net"], "estado": "CONSISTENCIA",
                                        "dif": None, "puntos": None, "ref_renglones": r_["renglones"],
                                        "cand_renglones": c_["renglones"]})
                continue
            # v73 (arbitraje #21): min == max es una decisión de valor fijo y manda sobre el número de la ref
            if rg.get("min") is not None and rg.get("max") is not None and rg["max"] >= rg["min"]:
                est, d, pts = grade_rango(rg["min"], rg["max"], c_["net"], tol, close)
            else:
                est, d, pts = grade(r_["net"], c_["net"], tol, close)
            res["familias"].append({"familia": f, "ref": r_["net"], "cand": c_["net"], "estado": est, "dif": d,
                                    "puntos": pts, "ref_renglones": r_["renglones"], "cand_renglones": c_["renglones"]})

    rp, cp = read_partitions(rwb), read_partitions(cwb)
    if rp is None:
        res["errores"].append("La referencia no tiene hoja Partition Takeoff legible")
    if cp is None:
        res["errores"].append("El candidato no tiene hoja Partition Takeoff legible")
    ext_aparte = any(w.title.strip().lower().startswith("exterior") for w in cwb.worksheets)
    res["types_no_comparados"] = []
    if rp and cp is not None:
        for t in sorted(rp):
            if ext_aparte and t.startswith("EXT") and t not in cp:
                res["types_no_comparados"].append(t)
                continue
            r = rp[t]
            if t not in cp:
                hoja = otra_hoja(cwb, t)
                if hoja:      # el candidato lo tiene, pero en otra hoja (p.ej. "Opening Bands"): no es FALTA
                    res["types"].append({"type": t, "ref_walls": r["walls"], "ref_lf": r["lf"], "cand_walls": None,
                                         "cand_lf": None, "estado": f"EN OTRA HOJA ({hoja})", "dif": None,
                                         "puntos": None})
                    continue
                res["types"].append({"type": t, "ref_walls": r["walls"], "ref_lf": r["lf"], "cand_walls": None,
                                     "cand_lf": None, "estado": "FALTA", "dif": None, "puntos": 0.0})
                continue
            c = cp[t]
            # v73: los rangos de types (_partition_types) mandan sobre la hoja de la ref (LF y número de walls)
            tr = (rangos.get("_partition_types") or {}).get(t) or {}
            lo, hi = tr.get("lf_min", tr.get("lf")), tr.get("lf_max", tr.get("lf"))
            wlo = tr.get("walls_min", tr.get("walls", r["walls"]))
            whi = tr.get("walls_max", tr.get("walls", r["walls"]))
            if lo is not None and hi is not None:
                est, d, pts = grade_rango(lo, hi, c["lf"], tol, close)
                est = est.replace(" (rango)", "")
            else:
                est, d, pts = grade(r["lf"], c["lf"], tol, close)
            if not (wlo <= c["walls"] <= whi) and est == "IGUAL":
                est, pts = "CERCA", 0.5                  # mismos LF pero distinto número de walls: revisar
            res["types"].append({"type": t, "ref_walls": r["walls"], "ref_lf": r["lf"], "cand_walls": c["walls"],
                                 "cand_lf": c["lf"], "estado": est, "dif": d, "puntos": pts})
        res["types_sobran"] = [{"type": t, **cp[t]} for t in sorted(cp) if t not in rp and not t.startswith("OTHER")]

    def pct(items):
        return round(100 * sum(x["puntos"] for x in items) / len(items), 1) if items else None
    res["calificacion_renglones"] = pct(res["materiales"])          # renglón por renglón (depende del formato)
    res["calificacion_materiales"] = pct([f for f in res.get("familias", []) if f["puntos"] is not None]) \
        if res.get("familias") else res["calificacion_renglones"]
    res["calificacion_types"] = pct([t for t in res["types"] if t["puntos"] is not None])
    parts = [p for p in (res["calificacion_materiales"], res["calificacion_types"]) if p is not None]
    res["calificacion_total"] = round(sum(parts) / len(parts), 1) if parts and not res["errores"] else 0.0
    return res


def fmt(v):
    return "—" if v is None else (f"{v:,.2f}" if isinstance(v, float) else str(v))


def pc(d):
    return "—" if d is None else f"{d:.1%}"


def report_md(res):
    L = [f"# Calificación del take-off: {res['calificacion_total']}%", "",
         f"- Materiales por familia (cantidad NET): **{fmt(res['calificacion_materiales'])}%**",
         f"- Partition types (LF y número de walls): **{fmt(res['calificacion_types'])}%**",
         f"- Tolerancia: IGUAL ≤ {res['tol']:.0%}, CERCA ≤ {res['close']:.0%}", ""]
    for e in res["errores"]:
        L.append(f"- ⚠️ {e}")
    fams = res.get("familias", [])
    if fams:
        L += ["", "## Materiales por familia (calificación principal)", "",
              "| Familia | Referencia | Candidato | Dif | Estado |", "|---|---|---|---|---|"]
        for f in fams:
            L.append(f"| {f['familia']} | {fmt(f['ref'])} | {fmt(f['cand'])} | {pc(f['dif'])} | {f['estado']} |")
    L += ["", f"Renglón por renglón (depende de cómo cada take-off nombra y separa los renglones): "
          f"{fmt(res.get('calificacion_renglones'))}%"]
    bad = [m for m in res["materiales"] if m["estado"] != "IGUAL"]
    L += ["", f"## Materiales que no salieron iguales ({len(bad)} de {len(res['materiales'])})", "",
          "| Material | Unit | Referencia | Candidato | Dif | Estado |", "|---|---|---|---|---|---|"]
    for m in bad:
        L.append(f"| {m['material'][:60]} | {m['unidad']} | {fmt(m['ref'])} | {fmt(m['cand'])} | "
                 f"{pc(m['dif'])} | {m['estado']} |")
    if res["materiales_sobran"]:
        L += ["", "## Materiales que solo están en el candidato (SOBRA)", ""]
        L += [f"- {m['material'][:70]} ({m['unidad']}): {fmt(m['cand'])}" for m in res["materiales_sobran"]]
    badt = [t for t in res["types"] if t["estado"] != "IGUAL"]
    L += ["", f"## Partition types que no salieron iguales ({len(badt)} de {len(res['types'])})", "",
          "| Type | Walls ref | Walls cand | LF ref | LF cand | Dif | Estado |", "|---|---|---|---|---|---|---|"]
    for t in badt:
        L.append(f"| {t['type']} | {t['ref_walls']} | {fmt(t['cand_walls'])} | {fmt(t['ref_lf'])} | "
                 f"{fmt(t['cand_lf'])} | {pc(t['dif'])} | {t['estado']} |")
    if res.get("types_no_comparados"):
        L += ["", "Types exteriores no comparados aquí (el candidato los tiene en su hoja Exterior; "
              "sus materiales sí cuentan en la tabla por familia): " + ", ".join(res["types_no_comparados"])]
    if res["types_sobran"]:
        L += ["", "## Types que solo están en el candidato (SOBRA)", ""]
        L += [f"- {t['type']}: {t['walls']} walls, {t['lf']:.2f} LF" for t in res["types_sobran"]]
    L += ["", "Antes de dar un MAL por error de Claude, abre el renglón: a veces la referencia o el "
          "emparejamiento de nombres es el que está mal (regla del artículo: leer cada resultado)."]
    return "\n".join(L) + "\n"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("referencia"); ap.add_argument("candidato")
    ap.add_argument("--out"); ap.add_argument("--json")
    ap.add_argument("--tol", type=float, default=0.02); ap.add_argument("--close", type=float, default=0.10)
    ap.add_argument("--pass", dest="passing", type=float, default=95.0)
    ap.add_argument("--rangos", help="rangos.json de la referencia definitiva (familias de criterio)")
    a = ap.parse_args()
    rangos = json.load(open(a.rangos, encoding="utf-8")) if a.rangos else None
    res = evaluate(a.referencia, a.candidato, a.tol, a.close, rangos)
    md = report_md(res)
    if a.out:
        open(a.out, "w", encoding="utf-8").write(md)
    if a.json:
        json.dump(res, open(a.json, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(md.split("\n\n")[0].replace("# ", ""))
    print(f"materiales {fmt(res['calificacion_materiales'])}% · types {fmt(res['calificacion_types'])}%")
    sys.exit(0 if res["calificacion_total"] >= a.passing else 1)


if __name__ == "__main__":
    main()
